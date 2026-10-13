"""1.8.0 regression tests (feedback from a real S01-S25 run + Codex 1.9-1.11 borrowings).

Called from tools/selftest.py; kept separate so the main self test stays navigable.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

NUMERIC_CSL = """<?xml version="1.0" encoding="utf-8"?>
<style xmlns="http://purl.org/net/xbiblio/csl" class="in-text" version="1.0">
  <info><title>Selftest numeric</title><id>selftest-numeric</id>
    <category citation-format="numeric"/><updated>2024-01-01T00:00:00+00:00</updated></info>
  <macro name="author"><names variable="author"><name initialize-with="" name-as-sort-order="all"
    sort-separator=" " delimiter=", "/></names></macro>
  <citation><sort><key variable="citation-number"/></sort>
    <layout prefix="[" suffix="]" delimiter=", "><text variable="citation-number"/></layout></citation>
  <bibliography second-field-align="flush"><layout suffix=".">
    <text variable="citation-number" suffix=". "/><text macro="author" suffix=". "/>
    <text variable="title" suffix=". "/><text variable="container-title" form="short"/>
    <date variable="issued" prefix=" "><date-part name="year"/></date>
    <text variable="volume" prefix=";"/><text variable="page" prefix=":"/></layout></bibliography>
</style>
"""

STYLE = {
    "font_family": "Times New Roman", "body_font_pt": 12, "title_font_pt": 14,
    "section_heading_font_pt": 12, "subsection_heading_font_pt": 12,
    "line_spacing": 2.0, "paragraph_spacing_after_pt": 0, "margins_in": 1.0,
    "paper_size": "A4", "all_text_black": True, "external_hyperlinks": False,
    "source": "Fixture journal is silent; documented medical-manuscript fallback.",
    "fallbacks": ["font_family"],
}

LEGENDS = """# Figure legends

## Figure 1.
Study flow and adjusted associations. (A) Participant flow from the source registry. (B) Adjusted
hazard ratios with 95% confidence intervals for each exposure category.

## Figure S1.
Sensitivity analysis layout. (A) Restricted cohort. (B) Alternative exposure definition used for
the supplementary comparison.
"""


def _manuscript(declarations: str, discussion: str = "Exposure was associated with outcome.") -> str:
    return f"""# Exposure and Clinical Outcome in a Multicentre Cohort Study

# Abstract

## Background
We evaluated 1284 adults in 2020.

## Results
The hazard ratio was 1.87.

Keywords: cohort studies, exposure, prognosis

# Introduction

Prior work examined this question [@alpha2020; @beta2021; @gamma2022].

# Methods

We studied 1284 participants (Supplementary Methods; Table S1).

# Results

Results are shown in Figure 1 and Figure S1.

# Discussion

{discussion}

# Declarations and Statements

{declarations}

# References

::: {{#refs}}
:::

# Figure legends

## Figure 1.
Study flow and adjusted associations. (A) Participant flow from the source registry. (B) Adjusted
hazard ratios with 95% confidence intervals for each exposure category.
"""


BIB = """@article{alpha2020, author = {{de Souza}, Ana}, title = {Adverse events of the da {Vinci} system},
  journal = {Journal of Robotic Surgery}, shortjournal = {J Robot Surg}, year = {2020}, volume = {12}}
@article{beta2021, author = {Lee, Bo}, title = {A second cohort}, journal = {Cohort Journal},
  shortjournal = {Cohort J}, year = {2021}, volume = {3}, pages = {101--109}}
@article{gamma2022, author = {Kim, Cy}, title = {A third cohort}, journal = {Cohort Journal},
  shortjournal = {Cohort J}, year = {2022}, volume = {4}, pages = {e12345}}
"""


def _paper(base: Path, target: dict | None) -> tuple[Path, dict]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from openpyxl import Workbook
    from wfcore import registry
    from wfcore.state import State
    p = base
    for sub in ("01_protocol", "05_figures/out", "04_tables/main", "07_manuscript",
                "08_submission/bundle", "06_refs", "00_input"):
        (p / sub).mkdir(parents=True, exist_ok=True)
    pipe = registry.load()
    State(p, ".wf").create(pipe.meta["name"], pipe.meta["version"], pipe.first().id)
    plan = {"main_figures": [{"id": "Figure 1", "file": "05_figures/out/Figure1.png"}],
            "supp_figures": [{"id": "Figure S1", "file": "05_figures/out/FigureS1.png"}],
            "main_tables": [{"id": "Table 1", "file": "04_tables/main/Table1.xlsx"}],
            "supp_tables": [{"id": "Table S1", "file": "04_tables/supplementary/TableS.xlsx"}]}
    (p / "01_protocol/artifact_plan.json").write_text(json.dumps(plan), encoding="utf-8")
    fig = plt.figure(figsize=(2, 1.5))
    plt.plot([0, 1], [0, 1])
    fig.savefig(p / "05_figures/out/Figure1.png", dpi=50)
    plt.close(fig)
    wb = Workbook()
    ws = wb.active
    ws.append(["Characteristic", "Group A", "Group B"])
    ws.append(["Age, years", "62.1", "61.7"])
    wb.save(p / "04_tables/main/Table1.xlsx")
    (p / "04_tables/table_captions.md").write_text(
        "# Table captions\n\n## Table 1.\nBaseline characteristics of the cohort.\n", encoding="utf-8")
    (p / "05_figures/legends.md").write_text(LEGENDS, encoding="utf-8")
    (p / "07_manuscript/title.md").write_text(
        "# Exposure and Clinical Outcome in a Multicentre Cohort Study\n", encoding="utf-8")
    (p / "06_refs/refs.bib").write_text(BIB, encoding="utf-8")
    (p / "06_refs/numeric.csl").write_text(NUMERIC_CSL, encoding="utf-8")
    (p / "08_submission/docx_style.json").write_text(json.dumps(STYLE), encoding="utf-8")
    (p / "00_input/author_info.json").write_text(json.dumps(
        {"authors": [{"name": "Ana Author"}, {"name": "Bo Author"}],
         "corresponding": {"name": "Ana Author", "email": "ana@example.org"}}), encoding="utf-8")
    if target is not None:
        (p / "08_submission/target_journal.json").write_text(json.dumps(target), encoding="utf-8")
    env = {**os.environ, "MEDPAPER_PROJECT": str(p), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    return p, env


def _build(p: Path, env: dict, kind: str, src: Path, out: str, *extra) -> subprocess.CompletedProcess:
    cmd = [sys.executable, str(ROOT / "tools/manuscript/build_docx.py"), "build", "--kind", kind,
           "--input", str(src), "--output", str(p / "08_submission/bundle" / out),
           "--style-config", str(p / "08_submission/docx_style.json"), *map(str, extra)]
    return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                          errors="replace", env=env)


def _xml(path: Path, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as zf:
        if part == "footers":
            return "".join(zf.read(n).decode("utf-8") for n in zf.namelist()
                           if n.startswith("word/footer"))
        return zf.read(part).decode("utf-8")


def run_word(tmp: Path, record, section) -> None:
    from manuscript import build_docx as bd
    from wfcore import docxaudit, journalspec
    from xml.etree import ElementTree as ET
    section("1.8.0 Word build: page numbers, references, legends, declarations, SI")
    target = {"journal": "Fixture Journal", "embed_figures": True, "word_tables": True,
              "legend_style": {"prefix": "Fig.", "bold": True, "panel_label_case": "lower",
                               "end_punctuation": "any"},
              "supplementary_aliases": {"Online Resource 1": ["Supplementary Methods", "Table S1"]}}
    p, env = _paper(tmp / "word", target)
    good_decl = ("## Funding\nNone.\n\n## Competing interests\nNone.\n\n"
                 "## Ethics approval\nApproved by the board.\n\n## Data availability\nPublic data.")
    src = p / "08_submission/full_manuscript.md"
    src.write_text(_manuscript(good_decl), encoding="utf-8")
    proc = _build(p, env, "manuscript", src, "manuscript.docx", "--bibliography",
                  p / "06_refs/refs.bib", "--csl", p / "06_refs/numeric.csl")
    out = p / "08_submission/bundle/manuscript.docx"
    record("manuscript builds with journal rules applied", proc.returncode == 0,
           (proc.stdout + proc.stderr).strip()[-300:])
    if proc.returncode != 0:
        return
    doc = _xml(out)
    text = bd.docx_text(out)
    record("every DOCX footer carries an automatic PAGE field",
           "PAGE" in _xml(out, "footers") and "footerReference" in doc)
    root = ET.fromstring(doc.encode("utf-8"))
    record("reference entries are number + TAB with a hanging indent",
           not docxaudit.reference_problems(root) and "w:hanging" in doc and "<w:tab/>" in doc,
           str(docxaudit.reference_problems(root)))
    record("adjacent citations collapse to a range through the CSL patch",
           "[1\u20133]" in text or "[1-3]" in text, [ln for ln in text.splitlines() if "[1" in ln][:2].__str__())
    record("legend label/panel style follows legend_style (Fig. 1, lower-case panels)",
           "Fig. 1" in text and "(a)" in text and "(A)" not in text.split("Figure legends")[-1])
    record("embedded figures and Word tables follow the journal switches",
           "<w:drawing" in doc and "<w:tbl>" in doc.replace("<w:tbl ", "<w:tbl>"))
    no_footer = p / "08_submission/bundle/no_footer.docx"
    import docx as _docx
    d = _docx.Document()
    d.add_paragraph("Plain")
    d.save(no_footer)
    with zipfile.ZipFile(no_footer) as zf:
        xmls = {n: zf.read(n) for n in zf.namelist() if n.startswith("word/")}
    record("audit rejects a DOCX without automatic page numbers",
           bool(docxaudit.footer_page_problems(xmls)))
    no_footer.unlink()
    bad_ref = ET.fromstring(
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        '<w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>1. Doe J. Title. '
        'J Med. 2020;12:.</w:t></w:r></w:p></w:body></w:document>')
    issues = docxaudit.reference_problems(bad_ref)
    record("audit rejects space-separated numbers, no hanging indent and dangling colons",
           any("hanging" in i for i in issues) and any("TAB" in i for i in issues)
           and any("dangling" in i for i in issues), str(issues))
    # Supplementary legends travel with the supplementary file, never the main text.
    with_supp = _manuscript(good_decl) + "\n## Figure S1.\nSensitivity analysis layout for the supplementary comparison only.\n"
    root2 = ET.fromstring(doc.encode("utf-8"))
    record("main-text legend audit covers main figures only",
           not docxaudit.legend_problems(root2, ["Figure 1"], 180, supplementary=["Figure S1"]))
    from wfcore.checks.manuscript import _markdown_legend_problems
    record("a supplementary legend in the main-text Figure legends is rejected",
           any("supplementary legend" in x for x in _markdown_legend_problems(with_supp, ["Figure 1"], 180)))
    from manuscript.assemble import main_legends
    record("assembly keeps main legends and drops Figure S* blocks",
           "Figure S1" not in main_legends(LEGENDS) and "Figure 1." in main_legends(LEGENDS))
    # Declarations order (default Springer Nature order) and Discussion prose.
    order = journalspec.declarations_order({})
    record("default declarations order is the Springer Nature order",
           order[:3] == ["Funding", "Competing interests", "Ethics approval"])
    record("out-of-order declarations are rejected; synonyms are recognised",
           bool(journalspec.declaration_order_problems(["Data availability", "Conflicts of interest"],
                                                      order, "Declarations"))
           and not journalspec.declaration_order_problems(["Funding", "Conflicts of interest"],
                                                          order, "Declarations"))
    from wfcore import registry
    from wfcore.checks import Ctx
    from wfcore.checks.manuscript import declaration_and_discussion_problems
    from wfcore.state import State
    pipe = registry.load()
    ctx = Ctx(pipeline=pipe, state=State(p, ".wf").load(), project=p,
              stage=pipe.stage("S22_polish"), spec={"check": "manuscript_structure"})
    sub = _manuscript(good_decl, "## Principal findings\nText.\n\n## Limitations\nText.")
    record("Discussion subheadings are rejected unless the journal requires them",
           any("Discussion has subheadings" in x for x in declaration_and_discussion_problems(ctx, sub))
           and not declaration_and_discussion_problems(ctx, _manuscript(good_decl)))
    # Supplementary: info page, supplementary legends, pipeline code stripped, monospace code.
    supp = p / "08_submission/supplementary_methods.md"
    supp.write_text("# Supplementary Methods\n\nCoding rules are listed below.\n\n```python\n"
                    "from wfcore.paths import project_dir\nRULE = r'^(?:abc|def)[0-9]{3}$'\n"
                    "DATA = 'F:\\\\data\\\\raw\\\\events.csv'\n```\n", encoding="utf-8")
    proc = _build(p, env, "supplementary", supp, "supplementary.docx")
    sdoc = p / "08_submission/bundle/supplementary.docx"
    stext = bd.docx_text(sdoc) if sdoc.is_file() else ""
    record("SI export: info page, supplementary legends, no pipeline code or local paths",
           proc.returncode == 0 and "Journal: Fixture Journal" in stext
           and "Corresponding author: Ana Author" in stext and "Figure S1" in stext
           and "wfcore" not in stext and "F:" not in stext and "events.csv" in stext,
           (proc.stdout + proc.stderr).strip()[-200:])
    record("SI code is small monospace and single spaced",
           sdoc.is_file() and "Courier New" in _xml(sdoc) and 'w:line="240"' in _xml(sdoc))
    record("audit flags pipeline-internal code in a submission file",
           bool(docxaudit.pipeline_code_problems("import wfcore; MEDPAPER_PROJECT")))
    # Measured checklist: generated from the files, stale or hand-edited copies fail.
    from wfcore import submissionqc
    (p / "08_submission/manuscript").mkdir(parents=True, exist_ok=True)
    (p / "08_submission/manuscript/full_manuscript.md").write_text(_manuscript(good_decl), encoding="utf-8")
    (p / "08_submission/bundle/manifest.json").write_text(json.dumps({"items": [
        {"role": "manuscript", "file": "08_submission/bundle/manuscript.docx", "required_by": "guide"},
        {"role": "supplementary", "file": "08_submission/bundle/supplementary.docx", "required_by": "guide"},
        {"role": "checklist", "file": submissionqc.CHECKLIST_REL, "required_by": "guide"}]}),
        encoding="utf-8")
    wrote = subprocess.run([sys.executable, str(ROOT / "tools/manuscript/submission_qc.py"), "write"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
    ctx = Ctx(pipeline=pipe, state=State(p, ".wf").load(), project=p, stage=pipe.stage("S23_package"),
              spec={"check": "submission_checklist_measured"})
    after_write = submissionqc.verify(ctx)
    cl = p / submissionqc.CHECKLIST_REL
    text_cl = cl.read_text(encoding="utf-8") if cl.is_file() else ""
    record("submission_qc.py writes a measured checklist (page numbers, references, legends ...)",
           "Automatic page numbers in every DOCX footer | PASS" in text_cl
           and "References: number + TAB" in text_cl
           and not any("not the current measurement" in x or "stale" in x for x in after_write),
           (wrote.stdout + wrote.stderr)[-300:] + str(after_write)[:200])
    cl.write_text(text_cl.replace("| FAIL |", "| PASS |") + "\n- [x] page numbers added\n", encoding="utf-8")
    record("a hand-ticked or stale checklist fails submission_checklist_measured",
           any("not the current measurement" in x for x in submissionqc.verify(ctx)))


def run_facts(tmp: Path, record, section) -> None:
    from pubmed import build_library as bl
    from wfcore import journalspec, scifreeze
    from wfcore.checks.refs import bib_format_problems
    section("1.8.0 fact preservation, aliases and refs.bib formatting")
    a = scifreeze._facts("Shown before [@a2020] [@b2021] and in Table S1.")
    b = scifreeze._facts("Shown before [@a2020; @b2021] and in Online Resource 1.",
                         {"Online Resource 1": ["Table S1"]})
    record("grouped citations and journal aliases preserve the same facts",
           sorted(a["citekeys"]) == sorted(b["citekeys"]) == ["a2020", "b2021"]
           and a["artifacts"] == b["artifacts"], f"{a} vs {b}")
    record("supplementary legend blocks moving out of the main text are not 'dropped' facts",
           scifreeze._facts("# Figure legends\n\n## Figure S1.\nSome legend (n = 12).\n")["numbers"] == [])
    record("alias expansion maps one resource onto several canonical ids",
           journalspec.expand_aliases("see Online Resource 2", {"Online Resource 2": ["Table S1", "Table S2"]})
           == "see Table S1, Table S2")
    record("titles brace proper nouns and acronyms",
           bl.protect_title("Adverse events of the da Vinci system reported to the FDA and MAUDE")
           == "Adverse events of the da {Vinci} system reported to the {FDA} and {MAUDE}")
    record("page ranges are expanded consistently",
           bl.expand_pages("123-9") == "123--129" and bl.expand_pages("e12345") == "e12345")
    record("surname particles are brace-protected",
           bl.bib_name({"last": "van der Berg", "first": "Jan"}) == "{van der Berg}, Jan")
    bib = bl.to_bibtex([{"citekey": "x2020", "title": "Use of the da Vinci robot",
                         "authors": [{"last": "de Souza", "first": "Ana"}], "journal": "Journal X",
                         "journal_abbrev": "J X", "year": "2020", "volume": "4", "pages": "",
                         "elocator": "e101"}])
    record("exported refs.bib passes the format audit (ISO abbreviation, article number)",
           not bib_format_problems(bib) and "shortjournal = {J X}" in bib and "pages = {e101}" in bib,
           str(bib_format_problems(bib)))
    record("format audit flags unprotected names and abbreviated ranges",
           len(bib_format_problems("@article{k,\n  author = {de Souza, Ana},\n  title = {The da Vinci "
                                   "robot},\n  journal = {J},\n  pages = {12-9}\n}\n")) >= 3)


def _gate(p: Path, name: str, stage: str, **spec):
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State
    load_all()
    pipe = registry.load()
    return get(name)(Ctx(pipeline=pipe, state=State(p, ".wf").load(), project=p,
                         stage=pipe.stage(stage), spec={"check": name, **spec}))


def run_counts(tmp: Path, record, section) -> None:
    from wfcore import journalspec
    section("1.8.0 typed journal limits, Word-compatible counts, drafting targets")
    record("Word-compatible count includes numbers and counts a citation group once",
           journalspec.word_count("We studied 1,284 adults (95% CI 1.2-2.3) [@a; @b].") == 8)
    caps, problems = journalspec.limits({"limits": {"reference_limit": 45, "display_items": 6,
                                                   "source": "45 references, no more than 6 figures/tables"}})
    record("typed limits keep 45 references and 6 display items apart (no prose parse)",
           not problems and caps["reference_limit"] == 45 and caps["display_items"] == 6)
    record("a missing limits object or a non-integer cap is a schema error",
           journalspec.limits({})[1] != [] and journalspec.limits(
               {"limits": {"main_text_words": "4000", "source": "x"}})[1] != [])
    p, _ = _paper(tmp / "counts", {"journal": "J", "limits": {
        "main_text_words": 40, "abstract_words": 250, "reference_limit": 2,
        "display_items": 6, "source": "guide URL; quoted rule"}})
    body = " ".join(["word"] * 38)
    (p / "08_submission/manuscript").mkdir(parents=True, exist_ok=True)
    (p / "07_manuscript/full_manuscript.md").write_text(
        f"# T\n\n# Abstract\n\nShort abstract.\n\nKeywords: a, b, c\n\n# Introduction\n\n{body} "
        "[@a; @b; @c]\n\n# References\n", encoding="utf-8")
    r = _gate(p, "journal_limits_met", "S22_polish")
    record("38+1 words against a 40-word cap fails the 97% safety margin and the 2-reference cap",
           not r.ok and "margin 38" in r.detail and "references 3/2" in r.detail, r.detail)
    target = json.loads((p / "08_submission/target_journal.json").read_text(encoding="utf-8"))
    target["limits"].update(main_text_words=100, reference_limit=5)
    (p / "08_submission/target_journal.json").write_text(json.dumps(target), encoding="utf-8")
    r = _gate(p, "journal_limits_met", "S22_polish")
    record("within typed caps passes", r.ok, r.detail)
    long_body = " ".join(["word"] * 4100)
    (p / "07_manuscript/full_manuscript.md").write_text(
        f"# T\n\n# Abstract\n\nShort.\n\n# Introduction\n\n{long_body}\n\n# References\n",
        encoding="utf-8")
    r = _gate(p, "draft_length_targets", "S17_assemble")
    record("the assembled draft is held to the default 4000-word body target", not r.ok, r.detail)
    # The final DOCX is what S23/S25 measure.
    from docx import Document
    d = Document()
    for text, style in (("A Title", None), ("Abstract", "Heading 1"), ("Background", "Heading 2"), ("One two three.", None),
                        ("Results", "Heading 2"), ("Four five.", None), ("Keywords: a, b", None),
                        ("Introduction", "Heading 1"), ("Six seven 8 nine.", None),
                        ("References", "Heading 1"), ("1. Ref one.", None), ("2. Ref two.", None)):
        para = d.add_paragraph(text)
        if style:
            para.style = d.styles["Heading 1"]
    d.save(p / "08_submission/bundle/m.docx")
    from manuscript.build_docx import load_style, normalize_docx
    normalize_docx(p / "08_submission/bundle/m.docx", load_style(p / "08_submission/docx_style.json"),
                   "manuscript")
    c = journalspec.docx_counts(p / "08_submission/bundle/m.docx", journalspec.word_count_scope({}))
    record("final-DOCX counts keep structured-abstract subheadings inside the abstract",
           c["abstract"] == 7 and c["main"] == 4 and c["references"] == 2, str(c))


def run_language(tmp: Path, record, section) -> None:
    from text import polish
    from wfcore.checks.text import backstage_hits
    section("1.8.0 backstage language and per-scope abbreviations")
    record("pipeline history phrases are detected",
           len(backstage_hits("This analysis was added after independent review. An earlier draft "
                              "wrongly kept these cases. At the reviewer's request we added it.")) >= 4)
    record("ordinary methods wording is not flagged",
           not backstage_hits("Two reviewers independently screened records; the previous version "
                              "of the guideline was used for coding."))
    docs = [("08_submission/manuscript/full_manuscript.md",
             "# T\n\n# Abstract\n\nAn adverse event (AE) was rare; AE rates fell; AE was coded.\n\n"
             "# Introduction\n\nEach AE was examined. Robotic surgery (RSX) was used once.\n")]
    found = polish.lint_abbreviations(docs)
    kinds = {(x["kind"], x["file"]) for x in found}
    record("an abbreviation defined only in the abstract is undefined in the body",
           ("abbrev_undefined", "(body)") in kinds and
           ("abbrev_undefined", "(abstract)") not in kinds, str(kinds))
    record("defined-but-unused is blocking within its own scope",
           ("abbrev_defined_unused", "(body)") in kinds, str(kinds))
    disp = polish.lint_display_abbreviations([("Table 1", "Baseline values.\nAbbreviations: BMI, "
                                                          "body mass index; SBP, systolic blood pressure.\nBMI 27")])
    record("a display item may not define abbreviations it does not use",
           any("SBP" in x["detail"] for x in disp) and not any("BMI" in x["detail"] for x in disp))


def _rw_paper(base: Path, at: str) -> Path:
    from wfcore import registry
    from wfcore.state import State
    pipe = registry.load()
    base.mkdir(parents=True, exist_ok=True)
    st = State(base, ".wf").create(pipe.meta["name"], pipe.meta["version"], at)
    for s in pipe.stages:
        if s.index < pipe.stage(at).index:
            st.stage_info(s.id)["status"] = "done"
    st.stage_info(at)["status"] = "active"
    st.record_decision("manuscript_human_reviewed", "NO_FURTHER_REVIEW",
                       "The fixture user approved the scientific manuscript at S19.")
    st.save()
    for rel, text in {"05_figures/code/fig1.R": "# fig1\n", "05_figures/code/fig2.R": "# fig2\n",
                      "05_figures/code/show.py": "print('rendered')\n",
                      "08_submission/cover_letter.md": "# Cover letter\n\nDear Editor.\n",
                      "08_submission/manuscript/statements.md": "# Declarations\n\n## Funding\nNone.\n",
                      "01_protocol/artifact_plan.json": json.dumps(
                          {"main_figures": [], "supp_figures": [], "main_tables": [],
                           "supp_tables": []})}.items():
        (base / rel).parent.mkdir(parents=True, exist_ok=True)
        (base / rel).write_text(text, encoding="utf-8")
    return base


def run_rework(tmp: Path, record, section) -> None:
    from wfcore import revision as rv, scifreeze
    section("1.8.0 rework: merge, priority/suspend, resume progress, S23 layout interception")
    acc = ["The requested change is visible in the source", "Nothing unrelated changed"]

    def plan(name: str, feedback: str, kind: str, src: str) -> Path:
        path = tmp / f"{name}.json"
        path.write_text(json.dumps({
            "feedback_verbatim": feedback,
            "interpretation": "Apply the user's request to the named sources only and keep every "
                              "other approved artifact unchanged.",
            "ambiguous_or_requires_user_decision": [],
            "items": [{"kind": kind, "request": feedback, "affected_sources": [src],
                       "acceptance_criteria": acc}]}), encoding="utf-8")
        return path

    def rw(P: Path, *args):
        env = {**os.environ, "MEDPAPER_PROJECT": str(P), "MEDPAPER_ROOT": str(ROOT),
               "PYTHONIOENCODING": "utf-8"}
        return subprocess.run([sys.executable, str(ROOT / "tools/rework.py"), *map(str, args),
                               "--project", str(P)], capture_output=True, text=True,
                              encoding="utf-8", errors="replace", env=env)

    def state(P: Path) -> dict:
        return json.loads((P / ".wf/state.json").read_text(encoding="utf-8"))

    P = _rw_paper(tmp / "rw", "S24_package_human_review")
    (P / "08_submission/manuscript/title_page.md").write_text("# Title page\n\nA. Author\n",
                                                             encoding="utf-8")
    stm = "08_submission/manuscript/statements.md"
    a = rw(P, "batch", "--sealed", "--why", "The user wants the cover letter fixed now.", "--plan",
           plan("p1", "Fix the greeting line of the cover letter.",
                "cover-letter-or-package-structure", "08_submission/cover_letter.md"))
    b = rw(P, "batch", "--plan", plan("p2", "Also reword the funding statement.",
                                      "title-page-or-statements", stm))
    record("new feedback on a sealed round names --merge / --priority instead of a dead end",
           a.returncode == 0 and b.returncode != 0 and "--merge" in b.stderr and "--priority" in b.stderr,
           (a.stdout + a.stderr + b.stderr)[-300:])
    c = rw(P, "batch", "--priority", "--plan", plan("p3", "Reword the funding statement first.",
                                                    "title-page-or-statements", stm))
    s = state(P)
    record("--priority parks the open round and opens the new feedback first",
           c.returncode == 0 and s.get("active_revision_round") == "R002"
           and s.get("suspended_revision_rounds") == ["R001"], (c.stdout + c.stderr)[-200:])
    rw(P, "seal", "--why", "The user said apply the funding statement fix now.")
    m = rw(P, "batch", "--merge", "--plan", plan("p4", "And correct the author initials.",
                                                 "title-page-or-statements",
                                                 "08_submission/manuscript/title_page.md"))
    r2 = json.loads((P / ".wf/revisions/R002.json").read_text(encoding="utf-8"))
    record("--merge folds new feedback into the sealed round and re-plans its scope",
           m.returncode == 0 and len(r2["items"]) == 2 and r2["status"] == "active"
           and "08_submission/manuscript/title_page.md" in r2["allowed_files"],
           (m.stdout + m.stderr)[-200:])
    (P / stm).write_text("# Declarations\n\n## Funding\nNo specific funding.\n", encoding="utf-8")
    g = rw(P, "progress", "--item", "R002-01", "--summary", "funding reworded; COI pending",
           "--changed-file", stm)
    v = rw(P, "resume")
    record("an interrupted round shows per-item status, notes and changed files",
           g.returncode == 0 and "[in_progress] R002-01" in v.stdout
           and "funding reworded" in v.stdout and stm in v.stdout and "R001" in v.stdout,
           (v.stdout + g.stderr)[-300:])
    blocked = rw(P, "resume", "--round", "R001")
    rw(P, "suspend", "--why", "back to the cover letter round for a moment")
    back = rw(P, "resume", "--round", "R001")
    r1 = json.loads((P / ".wf/revisions/R001.json").read_text(encoding="utf-8"))
    record("resuming a parked round rebases files other rounds changed meanwhile",
           blocked.returncode != 0 and back.returncode == 0 and r1["status"] == "active"
           and not rv.scope_problems(P, r1) and stm in r1["suspensions"][-1]["rebased"],
           (back.stdout + back.stderr)[-200:])
    record("snapshots skip the declared agent notes directory",
           (P / "08_submission/notes").mkdir(parents=True) or
           ((P / "08_submission/notes/plan.md").write_text("x", encoding="utf-8") and
            "08_submission/notes/plan.md" not in rv.snapshot(P)))
    # ---- S23: certified source edited outside a round ---------------------------------
    Q = _rw_paper(tmp / "s23", "S23_package")
    rdir = Q / ".wf/revisions"
    rdir.mkdir(parents=True)
    certified = rv.sha256(Q / "05_figures/code/fig2.R")
    (rdir / "R001.json").write_text(json.dumps({
        "schema_version": 2, "round_id": "R001", "status": "complete", "items": [],
        "validation": {"ok": True, "artifact_hashes": {"05_figures/code/fig2.R": certified}}}),
        encoding="utf-8")
    (Q / "05_figures/code/fig2.R").write_text("# fig2 with the box moved\n", encoding="utf-8")
    env = {**os.environ, "MEDPAPER_PROJECT": str(Q), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    run = subprocess.run([sys.executable, str(ROOT / "tools/wf.py"), "run-script",
                          "05_figures/code/show.py"], capture_output=True, text=True,
                         encoding="utf-8", errors="replace", env=env)
    record("rendering after an out-of-round edit to a certified script is refused at once",
           run.returncode == 2 and "05_figures/code/fig2.R" in run.stdout and "layout" in run.stdout,
           run.stdout[-300:])
    old = os.environ.get("MEDPAPER_PROJECT")
    os.environ["MEDPAPER_PROJECT"] = str(Q)
    try:
        sys.path.insert(0, str(ROOT / "tools/hooks"))
        import frozen_guard
        payload = json.dumps({"tool_input": {"path": str(Q / "05_figures/code/fig2.R")}})
        hit = frozen_guard.blocked(payload)
        lay = rw(Q, "layout", "--source", "05_figures/code/fig2.R", "--why",
                 "Move the legend box of Figure 2 to the right margin.")
        after = frozen_guard.blocked(payload)
    finally:
        if old is None:
            os.environ.pop("MEDPAPER_PROJECT", None)
        else:
            os.environ["MEDPAPER_PROJECT"] = old
    run2 = subprocess.run([sys.executable, str(ROOT / "tools/wf.py"), "run-script",
                           "05_figures/code/show.py"], capture_output=True, text=True,
                          encoding="utf-8", errors="replace", env=env)
    rq = json.loads((rdir / "R002.json").read_text(encoding="utf-8")) if (rdir / "R002.json").is_file() else {}
    record("the write hook blocks a certified source and names the layout command",
           hit is not None and "rework.py layout" in hit and after is None, str(hit))
    record("rework.py layout opens a sealed layout round at S23 whose baseline is the certified file",
           lay.returncode == 0 and rq.get("status") == "active" and rq.get("auto_layout")
           and rq["baseline"]["05_figures/code/fig2.R"] == certified
           and "05_figures/code/fig2.R" in rq.get("allowed_files", []) and run2.returncode == 0,
           (lay.stdout + lay.stderr + run2.stdout)[-300:])
    # ---- freeze: identical content after a newer approval re-stamps, no --force -----------
    F = _rw_paper(tmp / "frz", "S19_human_review")
    for rel, text in {"07_manuscript/full_manuscript.md": "# T\n", "06_refs/refs.bib": "% bib\n",
                      "06_refs/library.json": "{}", "03_analysis/results/a.json": "{}"}.items():
        (F / rel).parent.mkdir(parents=True, exist_ok=True)
        (F / rel).write_text(text, encoding="utf-8")
    scifreeze.write_freeze(F)
    first = scifreeze.load_freeze(F)["frozen_at"]
    st = json.loads((F / ".wf/state.json").read_text(encoding="utf-8"))
    st["decisions"]["manuscript_human_reviewed"]["at"] = "2999-01-01T00:00:00Z"
    (F / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    scifreeze.write_freeze(F)
    fz = scifreeze.load_freeze(F)
    record("re-capturing identical science after a newer approval re-stamps without --force",
           fz.get("restamps") and fz["restamps"][-1]["from"] == first, str(fz.get("restamps")))


CARD_CMD_RE = __import__("re").compile(
    r"^(?:\.\\\.venv\\Scripts\\python\.exe|python)\s+(tools[/\\][\w/\\.]+\.py)\s*(.*)$")


def card_commands() -> list[tuple[str, str, str]]:
    """(card, tool, argument string) for every engine command shown in a stage card."""
    import re
    out = []
    for card in sorted((ROOT / "pipeline/stages").glob("*.md")):
        text = card.read_text(encoding="utf-8")
        lines = [ln.strip() for block in re.findall(r"```[^\n]*\n(.*?)```", text, re.S)
                 for ln in block.splitlines()]
        lines += re.findall(r"`(\.\\\.venv\\Scripts\\python\.exe tools/[^`]+)`", text)
        for line in lines:
            m = CARD_CMD_RE.match(line)
            if m:
                args = re.sub(r"\s+#\s.*$", "", m.group(2))  # trailing shell comment
                out.append((card.name, m.group(1).replace("\\", "/"), args))
    return out


def parse_card_command(tool: str, argstr: str) -> str | None:
    """Parse the arguments with the tool's real argparse definition; None when valid."""
    import argparse
    import contextlib
    import io
    import re
    import runpy
    import shlex
    tokens = [re.sub(r"<[^>]*>", "X", t) for t in shlex.split(argstr.replace("\\", "/"))]

    class Parsed(Exception):
        pass
    original = argparse.ArgumentParser.parse_args

    def fake(self, args=None, namespace=None):
        original(self, tokens, namespace)
        raise Parsed
    err = io.StringIO()
    argparse.ArgumentParser.parse_args = fake
    try:
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
            if tool == "tools/wf.py":
                from wfcore import cli
                cli.main(tokens)
            else:
                runpy.run_path(str(ROOT / tool), run_name="cardcheck")["main"]()
    except Parsed:
        return None
    except SystemExit:
        return err.getvalue().strip().splitlines()[-1] if err.getvalue().strip() else "exit"
    except Exception as exc:  # noqa: BLE001
        return f"{type(exc).__name__}: {exc}"
    finally:
        argparse.ArgumentParser.parse_args = original
    return "main() returned without parsing arguments"


def run_runtime(tmp: Path, record, section) -> None:
    section("1.8.0 runtime: stage-card commands match the CLI, UTF-8, --log")
    cmds = card_commands()
    bad = []
    for card, tool, argstr in cmds:
        if not (ROOT / tool).is_file():
            bad.append(f"{card}: {tool} does not exist")
            continue
        problem = parse_card_command(tool, argstr)
        if problem:
            bad.append(f"{card}: {tool} {argstr[:60]} -> {problem[:120]}")
    record(f"all {len(cmds)} engine commands in the stage cards parse with the real CLI",
           len(cmds) > 40 and not bad, "; ".join(bad[:6]))
    record("the S25 card returns to S24 with `wf loop --to` (no positional stage)",
           "loop --to S24_package_human_review" in
           (ROOT / "pipeline/stages/S25_submission_audit.md").read_text(encoding="utf-8"))
    P = _rw_paper(tmp / "rt", "S03_protocol")
    env = {**os.environ, "MEDPAPER_PROJECT": str(P), "MEDPAPER_ROOT": str(ROOT)}
    env.pop("PYTHONIOENCODING", None)
    why = "用户选择该期刊：范围匹配，近期发表过同类公开数据库研究，审稿周期可接受，费用在预算内。" * 2
    log = P / "temp/check.log"
    proc = subprocess.run([sys.executable, str(ROOT / "tools/wf.py"), "decide", "journal_note", "X",
                           "--why", why], capture_output=True, env=env)
    chk = subprocess.run([sys.executable, str(ROOT / "tools/wf.py"), "check", "--log", str(log)],
                         capture_output=True, env=env)
    handoff = (P / ".wf/state.json").read_text(encoding="utf-8")
    record("Chinese --why text is stored and printed as UTF-8 even without PYTHONIOENCODING",
           proc.returncode == 0 and why in handoff and
           (why in proc.stdout.decode("utf-8", "strict") or proc.stdout.decode("utf-8", "strict") != ""),
           proc.stdout[-200:].decode("utf-8", "replace"))
    record("wf check --log writes a UTF-8 log ending with a completion marker",
           log.is_file() and "[wf] finished rc=" in log.read_text(encoding="utf-8"),
           chk.stdout[-200:].decode("utf-8", "replace"))


def run_review(tmp: Path, record, section) -> None:
    from wfcore import readerreview as rr
    section("1.8.0 review: hash-bound text copies, delta re-review, design precedent")
    p, _ = _paper(tmp / "review", None)
    plan = json.loads((p / "01_protocol/artifact_plan.json").read_text(encoding="utf-8"))
    plan["supp_figures"], plan["supp_tables"] = [], []
    (p / "01_protocol/artifact_plan.json").write_text(json.dumps(plan), encoding="utf-8")
    (p / "07_manuscript/full_manuscript.md").write_text("# T\n\nText.\n", encoding="utf-8")
    sections = "".join(f"\n## {h}\n" + "The reviewer read the text, tables and figures carefully. " * 3 + "\n"
                       for h in rr.SECTIONS)

    def report(block: str) -> None:
        (p / rr.REVIEW_REL).write_text("# Review\n\n## Verdict\nREADY\n" + sections + "\n" + block + "\n",
                                       encoding="utf-8")
    report(rr.evidence_block(p))
    full_ok = not rr.verify(p)
    copies = rr.write_text_copies(p, rr.reader_inputs(p))
    txt = (p / copies[0]).read_text(encoding="utf-8") if copies else ""
    record("reviewers get hash-bound text copies of workbooks",
           copies and "# sha256: " in txt and "Characteristic" in txt, str(copies))
    from openpyxl import load_workbook
    wb = load_workbook(p / "04_tables/main/Table1.xlsx")
    wb.active["B2"] = "62.4"
    wb.save(p / "04_tables/main/Table1.xlsx")
    stale = rr.verify(p)
    delta = rr.delta_inputs(p)
    report(rr.evidence_block(p, delta=True))
    delta_ok = not rr.verify(p)
    record("a small change re-reviews only the changed file; unchanged coverage is kept",
           full_ok and stale and "04_tables/main/Table1.xlsx" in delta
           and len(delta) < len(rr.reader_inputs(p)) and delta_ok,
           f"{stale} {list(delta)}")
    wb.active["B2"] = "62.9"
    wb.save(p / "04_tables/main/Table1.xlsx")
    record("a delta review that misses a changed file still fails", bool(rr.verify(p)))
    rr.write_text_copies(p, [])
    record("stale text copies are removed when regenerated", not list((p / rr.TEXT_DIR_REL).glob("*.txt")))
    proto = p / "01_protocol/protocol_v1.md"
    proto.write_text("# Design precedent\n- Comparator: precedent yes - PMID 31234567\n"
                     "- Denominator: precedent no - justification: device-level reports have no exposure "
                     "denominator, so proportions of reports are used\n", encoding="utf-8")
    bad = _gate(p, "design_precedent_recorded", "S03_protocol", path="01_protocol/protocol_v1.md")
    proto.write_text(proto.read_text(encoding="utf-8") +
                     "- Primary outcome: precedent yes - [@alpha2020]\n", encoding="utf-8")
    good = _gate(p, "design_precedent_recorded", "S03_protocol", path="01_protocol/protocol_v1.md")
    record("design precedent: comparator, denominator and primary outcome each need precedent or defence",
           not bad.ok and "primary outcome" in bad.detail.lower() and good.ok,
           bad.detail + " | " + good.detail)


def run(tmp: Path, record, section) -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    run_review(tmp, record, section)
    run_runtime(tmp, record, section)
    run_rework(tmp, record, section)
    run_word(tmp, record, section)
    run_facts(tmp, record, section)
    run_counts(tmp, record, section)
    run_language(tmp, record, section)

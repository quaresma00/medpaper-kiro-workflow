"""Three-line (booktabs) xlsx writer.

Layout produced, which is exactly what the `tables_threeline` gate verifies:

    row 1            title, unruled
    row 2            header, rule above and rule below
    rows 3..n        data, unruled
    row n            bottom rule
    rows n+1..       footnotes, unruled

No vertical rules, no interior horizontal rules. Use this rather than styling cells
by hand, so every table in the paper is identical.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

_TOOLS = Path(__file__).resolve().parents[1]
if str(_TOOLS) not in sys.path:
    sys.path.insert(0, str(_TOOLS))
from wfcore import paths as wfpaths  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required for table building.\n"
        "  uv pip install --python .venv/Scripts/python.exe openpyxl"
    ) from exc

RULE = Side(style="thin", color="FF000000")
BASE_FONT = "Arial"
TITLE_PT = 10
BODY_PT = 9
FOOT_PT = 8
MAX_CELL_CHARS = 300
MAX_FOOTNOTE_CHARS = 1500


class TableSpec(dict):
    """{sheet, title, header, rows, footnotes, align, col_widths}"""


def _check(spec: dict) -> None:
    for key in ("sheet", "title", "header", "rows"):
        if not spec.get(key):
            raise ValueError(f"table spec missing '{key}'")
    if not spec.get("footnotes"):
        raise ValueError(
            f"'{spec['sheet']}': footnotes are required. A three-line table needs at least "
            "an abbreviations or data-format note."
        )
    ncol = len(spec["header"])
    for i, row in enumerate(spec["rows"], 1):
        if len(row) != ncol:
            raise ValueError(
                f"'{spec['sheet']}' row {i} has {len(row)} cells but the header has {ncol}"
            )
        for cell in row:
            if isinstance(cell, str) and len(cell) > MAX_CELL_CHARS:
                raise ValueError(
                    f"'{spec['sheet']}' row {i}: cell of {len(cell)} chars. A table is not a "
                    "place for prose - move it to the footnote or the manuscript."
                )
            raw = RAW_STATUS_RE.search(str(cell)) if cell is not None else None
            if raw:
                raise ValueError(
                    f"'{spec['sheet']}' row {i}: raw status '{raw.group(0)}' - use "
                    "not_estimable(lower) or MISSING (defined once in the footnote)")
    # 1.11.0 (#2d): footnotes are limited by total words, never per row; a footnote row is a
    # whole note (do not split sentences over rows to fit a cell limit).
    total = sum(len(str(f)) for f in spec["footnotes"])
    if total > MAX_FOOTNOTE_CHARS:
        raise ValueError(
            f"'{spec['sheet']}': footnotes total {total} chars (limit {MAX_FOOTNOTE_CHARS}). "
            "That is an analysis report, not a table footnote."
        )
    words = sum(len(str(f).split()) for f in spec["footnotes"])
    if words > FOOTNOTE_WORDS_MAX:
        raise ValueError(
            f"'{spec['sheet']}': footnotes total {words} words (limit {FOOTNOTE_WORDS_MAX}): keep "
            "units, abbreviations and the estimation method in one sentence; move the rest to "
            "Methods and diagnostics to a supplementary table.")
    if not re.match(r"^\s*(table|supplementary table)\s+S?\d+[.:]", spec["title"], re.I):
        raise ValueError(
            f"'{spec['sheet']}': title must start with e.g. 'Table 1.' or 'Table S1.' - got "
            f"{spec['title'][:40]!r}"
        )


def _autowidth(header: list, rows: list, explicit: list | None) -> list[float]:
    if explicit:
        return list(explicit)
    widths = []
    for i, head in enumerate(header):
        longest = max([len(str(head))] + [len(str(r[i])) for r in rows] or [0])
        widths.append(min(48.0, max(10.0, longest * 1.15 + 2)))
    return widths


def _styled(spec: dict) -> dict:
    """1.9.0: journal number style and spelling applied at build time (wfcore/numfmt.py)."""
    from wfcore import numfmt
    st = numfmt.style()
    if all(v is None for v in st.values()):
        return spec

    def f(v):
        return numfmt.apply_text(v, st) if isinstance(v, str) else v
    out = dict(spec)
    out["title"] = f(spec["title"])
    out["header"] = [f(h) for h in spec["header"]]
    out["rows"] = [[f(v) for v in row] for row in spec["rows"]]
    out["footnotes"] = [f(str(n)) for n in spec["footnotes"]]
    return out


def _write_sheet(ws, spec: dict) -> None:
    spec = _styled(spec)
    header = spec["header"]
    rows = spec["rows"]
    footnotes = spec["footnotes"]
    ncol = len(header)
    align = spec.get("align") or ["left"] + ["center"] * (ncol - 1)
    last_col = get_column_letter(ncol)

    # --- title (row 1, unruled) ---------------------------------------
    ws.cell(row=1, column=1, value=spec["title"])
    ws.cell(row=1, column=1).font = Font(name=BASE_FONT, size=TITLE_PT, bold=True)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    if ncol > 1:
        ws.merge_cells(f"A1:{last_col}1")

    # --- header (row 2, ruled above and below) -------------------------
    for c, text in enumerate(header, 1):
        cell = ws.cell(row=2, column=c, value=text)
        cell.font = Font(name=BASE_FONT, size=BODY_PT, bold=True)
        cell.alignment = Alignment(horizontal=align[c - 1], vertical="bottom", wrap_text=True)
        cell.border = Border(top=RULE, bottom=RULE)

    # --- data ----------------------------------------------------------
    first_data = 3
    for r, row in enumerate(rows, first_data):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=BASE_FONT, size=BODY_PT)
            cell.alignment = Alignment(horizontal=align[c - 1], vertical="top", wrap_text=False)
    last_data = first_data + len(rows) - 1

    # --- bottom rule ---------------------------------------------------
    for c in range(1, ncol + 1):
        ws.cell(row=last_data, column=c).border = Border(bottom=RULE)

    # --- footnotes (unruled) -------------------------------------------
    for i, note in enumerate(footnotes):
        r = last_data + 1 + i
        cell = ws.cell(row=r, column=1, value=str(note))
        cell.font = Font(name=BASE_FONT, size=FOOT_PT)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
        if ncol > 1:
            ws.merge_cells(f"A{r}:{last_col}{r}")

    for i, w in enumerate(_autowidth(header, rows, spec.get("col_widths")), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{first_data}"


def write_workbook(path: str | Path, tables: list[dict]) -> Path:
    """One workbook, one sheet per table. Use this for the supplementary tables."""
    if not tables:
        raise ValueError("no tables given")
    for spec in tables:
        _check(spec)
    seen = [t["sheet"] for t in tables]
    if len(seen) != len(set(seen)):
        raise ValueError(f"duplicate sheet names: {seen}")
    wb = Workbook()
    wb.remove(wb.active)
    for spec in tables:
        ws = wb.create_sheet(title=spec["sheet"][:31])
        _write_sheet(ws, spec)
    p = wfpaths.paper_output_path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    # 1.9.0: same table -> same bytes (fixed workbook timestamps and zip entry dates), so a
    # rebuild does not look like an edit to the revision credentials.
    import datetime as _dt
    fixed = _dt.datetime(2000, 1, 1)
    wb.properties.created = wb.properties.modified = fixed
    wb.save(p)
    from wfcore import determinism
    determinism.normalize(p)
    print(f"wrote {p}  ({len(tables)} sheet(s))")
    return p


def write_table(path: str | Path, sheet: str, title: str, header: list,
                rows: list, footnotes: list, align: list | None = None,
                col_widths: list | None = None) -> Path:
    """One workbook holding one table. Use this for each main table."""
    return write_workbook(path, [{
        "sheet": sheet, "title": title, "header": header, "rows": rows,
        "footnotes": footnotes, "align": align, "col_widths": col_widths,
    }])


# 1.11.0 (#2c): one missing symbol, one non-estimable wording, no raw statistical statuses.
MISSING = "\u2013"
RAW_STATUS_RE = re.compile(r"\bUndefined\b|\bNot\s+fitted\b|\bnot\s+converged\b|\bNA\b|\bN/A\b|"
                           r"\bNaN\b|(?<![\w.])-?Inf\b|\binfinity\b|\u221e|\bNULL\b")
FOOTNOTE_WORDS_MAX = 120


def not_estimable(lower=None, decimals: int = 2, bound: str = "lower 95% CI bound") -> str:
    """A zero-cell / unbounded estimate: 'Not estimable (lower 95% CI bound 19.60)'."""
    if lower is None:
        return "Not estimable"
    return f"Not estimable ({bound} {fmt(lower, decimals)})"


def fmt(value, decimals: int = 2, pct: bool = False) -> str:
    """Consistent numeric rendering. Keeps trailing zeros, which reviewers expect."""
    if value is None or value == "":
        return MISSING
    try:
        if value != value or abs(float(value)) == float("inf"):  # NaN / Inf never print
            return MISSING
    except (TypeError, ValueError):
        pass
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f"{f:.{decimals}f}%" if pct else f"{f:.{decimals}f}"


def p_value(p) -> str:
    """APA/medical convention: report small p as an inequality, never as 0.000."""
    try:
        f = float(p)
    except (TypeError, ValueError):
        return str(p)
    if f < 0.001:
        # The journal's spacing ("< 0.001") is applied to every cell at build (numfmt).
        return "<0.001"
    return f"{f:.3f}" if f < 0.01 else f"{f:.2f}"


def ci(estimate, low, high, decimals: int = 2) -> str:
    return f"{fmt(estimate, decimals)} ({fmt(low, decimals)} to {fmt(high, decimals)})"

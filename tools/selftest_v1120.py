"""1.12.0 regression tests: lighter engine. False-positive gate fixes (A), deleted/merged gates
(B), CLI fixes (C), slim stage cards (D), `wf handoff` / reviewer report / clinical writer (E)
and cheap early checks (F). Synthetic fixtures only.

Called from tools/selftest.py.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from contextlib import contextmanager
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))


def _w(p: Path, rel: str, text) -> None:
    (p / rel).parent.mkdir(parents=True, exist_ok=True)
    (p / rel).write_text(text if isinstance(text, str) else json.dumps(text), encoding="utf-8")


def _paper(base: Path, at: str) -> Path:
    import selftest_v180 as v180
    return v180._rw_paper(base, at)


def _gate(p: Path, name: str, stage: str, **spec):
    import selftest_v180 as v180
    return v180._gate(p, name, stage, **spec)


@contextmanager
def _env(**kv):
    old = {k: os.environ.get(k) for k in kv}
    os.environ.update({k: str(v) for k, v in kv.items()})
    try:
        yield
    finally:
        for k, v in old.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def _xlsx(p: Path, rel: str, rows) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    (p / rel).parent.mkdir(parents=True, exist_ok=True)
    wb.save(p / rel)


# ------------------------------------------------------------------------------ group A
def run_a(tmp: Path, record, section) -> None:
    section("1.12.0 A: false-positive gates")
    # A1 single_section_written: only the stage being advanced is judged
    p = _paper(tmp / "a1", "S19_human_review")
    for rel in ("07_manuscript/methods.md", "07_manuscript/introduction.md",
                "07_manuscript/discussion.md", "07_manuscript/human_review.md"):
        _w(p, rel, "# x\n\ntext\n")
    old = _gate(p, "single_section_written", "S08_methods")
    q = _paper(tmp / "a1b", "S08_methods")
    _w(q, "07_manuscript/introduction.md", "# Introduction\n\nahead of schedule\n")
    live = _gate(q, "single_section_written", "S08_methods")
    record("A1 single_section_written: a finished stage re-checked later passes; the active "
           "stage still catches a future file", old.ok and not live.ok, old.detail + " | " + live.detail)

    # A2 render/QC race
    from figures import qc
    p = _paper(tmp / "a2", "S11_figures")
    _w(p, "05_figures/out/F1.png", "png")
    with _env(MEDPAPER_PROJECT=p):
        idle = qc.render_in_progress("05_figures/out/F1.png", wait=0)
        _w(p, "05_figures/qc/F1.rendering", "now")
        busy = qc.render_in_progress("05_figures/out/F1.png", wait=0.6)
        old_t = time.time() - 2 * qc.RENDER_STALE_SECONDS
        os.utime(p / "05_figures/qc/F1.rendering", (old_t, old_t))
        stale = qc.render_in_progress("05_figures/out/F1.png", wait=0)
    rtext = (ROOT / "tools/figures/R/medfig.R").read_text(encoding="utf-8")
    record("A2 r_render_provenance: a running render is reported as 'render in progress' "
           "(warn), a stale marker is not; save_medpaper writes and clears the marker",
           not idle and "render in progress" in busy and not stale
           and '".rendering"' in rtext and "unlink(marker)" in rtext
           and 'add("r_render_provenance", False, busy, severity="warn")'
           in (ROOT / "tools/figures/qc.py").read_text(encoding="utf-8"), busy)

    # A3 numbers_have_provenance: lock files, IO errors, footnote constants
    from wfcore import xlsxlite
    p = _paper(tmp / "a3", "S10_tables")
    _w(p, "03_analysis/results/main.json", {"n_reports": 12, "rate": 3.5})
    _xlsx(p, "04_tables/main/T1.xlsx", [["Characteristic", "Value"], ["Reports", "12"],
                                         ["Rate per 1000 reports", "3.5"], ["Denominator", "1000"]])
    _w(p, "04_tables/main/~$T1.xlsx", "lock")
    before = _gate(p, "numbers_have_provenance", "S10_tables", source="tables")
    _w(p, "01_protocol/display_dictionary.json", {"terms": [], "footnote_constants": [
        {"value": "1000", "reason": "rates are expressed per 1000 reports"}]})
    after = _gate(p, "numbers_have_provenance", "S10_tables", source="tables")
    real = xlsxlite.numeric_cell_values

    def locked(path):
        raise PermissionError(13, "locked by another process", str(path))
    xlsxlite.numeric_cell_values = locked
    try:
        io = _gate(p, "numbers_have_provenance", "S10_tables", source="tables")
    finally:
        xlsxlite.numeric_cell_values = real
    record("A3 numbers_have_provenance: '~$' lock files are ignored, a registered footnote "
           "constant passes, a locked workbook is an IO problem with a fix hint",
           not before.ok and "1000" in before.detail and "~$" not in before.detail and after.ok
           and not io.ok and "IO problem" in io.detail and any("Excel" in h for h in io.hints),
           f"{before.detail[:80]} | {after.detail[:60]} | {io.detail[:80]}")

    # A4 reader_review_coverage: always incremental
    import selftest_v180 as v180
    from wfcore import readerreview as rr, reviewrounds
    p, _ = v180._paper(tmp / "a4", None)
    plan = json.loads((p / "01_protocol/artifact_plan.json").read_text(encoding="utf-8"))
    plan["supp_figures"], plan["supp_tables"] = [], []
    _w(p, "01_protocol/artifact_plan.json", plan)
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\nText.\n")
    sections = "".join(f"\n## {h}\n" + "The reviewer read the text, tables and figures carefully. " * 3
                       for h in rr.SECTIONS)
    _w(p, rr.REVIEW_REL, "# Review\n" + sections + "\n" + rr.evidence_block(p) + "\n")
    first = rr.verify(p)
    _w(p, "05_figures/legends.md", "# Figure legends\n\n## Figure 1.\nChanged legend text for the figure.\n")
    stale = rr.verify(p)
    cur = rr.reader_inputs(p)
    rel = "05_figures/legends.md"
    data = reviewrounds.load(p)
    data["manuscript"].append({"id": "M001", "scope": "manuscript", "mode": "delta",
                               "started_at": "2026-01-01T00:00:00Z", "inputs": cur,
                               "reviewed": [rel], "closed_at": "2026-01-01T01:00:00Z",
                               "after": cur, "user_decision": None})
    reviewrounds._save(p, data)
    healed = rr.verify(p)
    record("A4 reader_review_coverage: after a change only the changed file is asked for, and "
           "a closed delta round covering it restores coverage without a full re-read",
           not first and len(stale) == 1 and rel in stale[0] and "full_manuscript" not in stale[0]
           and not healed, f"{stale} | {healed}")

    # A5 submission checklist: warn rows render WARN
    from wfcore import submissionqc
    record("A5 submission_checklist_measured: a warn-severity row is WARN, never FAIL",
           submissionqc.row_status(["close to cap"], warn=True) == "WARN"
           and submissionqc.row_status(["over"], warn=False) == "FAIL"
           and submissionqc.row_status([], warn=True) == "PASS")

    # A6 stacked bars
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from figures import elements
    fig, ax = plt.subplots()
    ax.bar([0, 1, 2], [3, 4, 2])
    ax.bar([0, 1, 2], [1, 2, 5], bottom=[3, 4, 2])
    stacked = elements.baseline_zero(fig)
    plt.close(fig)
    fig, ax = plt.subplots()
    ax.barh([0, 1], [3, 4])
    ax.barh([0, 1], [2, 2], left=[3, 4])
    hstacked = elements.baseline_zero(fig)
    plt.close(fig)
    fig, ax = plt.subplots()
    ax.bar([0, 1, 2], [3, 4, 2], bottom=1)
    ax.set_ylim(1, 6)
    trunc = elements.baseline_zero(fig)
    plt.close(fig)
    record("A6 bar_baseline_zero: stacked bars (vertical and horizontal) pass, a truncated "
           "baseline still fails; the R audit skips continuation segments",
           stacked[0] and hstacked[0] and not trunc[0] and "cont <- vapply" in rtext,
           f"{stacked} {hstacked} {trunc}")
    import shutil
    if shutil.which("Rscript"):
        import selftest as st_main
        P = st_main._new_paper(tmp / "a6r")
        _w(P, "01_protocol/artifact_plan.json", {"main_figures": [
            {"id": "Figure 1", "archetype": "bar_dot", "width": "single",
             "file": "05_figures/out/Stack.png"}], "supp_figures": []})
        st_main._r_script(P / "05_figures/code/stack.R", '''
d <- data.frame(g = rep(c("Early", "Late"), each = 2), part = rep(c("Mild", "Severe"), 2), n = c(3, 1, 4, 2))
p <- ggplot(d, aes(g, n, fill = part)) + geom_col() + labs(x = "Period", y = "Reports, n", fill = "Severity")
save_medpaper(p, "05_figures/out/Stack", width = "single", archetype = "bar_dot")
''')
        r = st_main._wf("run-script", "05_figures/code/stack.R", env=st_main._env(MEDPAPER_PROJECT=P))
        side = P / "05_figures/qc/Stack.artist.json"
        el = json.loads(side.read_text(encoding="utf-8")).get("elements", {}) if side.is_file() else {}
        record("A6/A2 R: a stacked geom_col renders with baseline_zero found and no render "
               "marker left behind", r.returncode == 0 and el.get("baseline_zero", {}).get("found")
               and not (P / "05_figures/qc/Stack.rendering").exists(),
               str(el.get("baseline_zero")) + st_main._out(r)[-200:])

    # A7 bib export -> check passes
    from pubmed import build_library as bl
    from wfcore.checks.refs import bib_format_problems
    bib = bl.to_bibtex([{"citekey": "k2021", "authors": [{"last": "Smith", "first": "Ann"}],
                         "title": "Liver injury after chemotherapy: Results from the FDA "
                                  "Adverse Event Reporting System in Japan",
                         "journal": "Journal X", "journal_abbrev": "J X", "year": "2021",
                         "volume": "3", "pages": "101-109"}])
    hand = ("@article{h,\n  author = {Smith, Ann},\n  title = {Outcomes in Boston: Results},\n"
            "  journal = {J},\n  shortjournal = {J},\n  pages = {1--9}\n}\n")
    record("A7 bib_format_clean uses the exporter's own rule: a fresh export passes (ordinary "
           "word after a colon unbraced), a hand-typed proper noun is still flagged",
           not bib_format_problems(bib) and "Results from" in bib
           and any("Boston" in x for x in bib_format_problems(hand)),
           str(bib_format_problems(bib)) + " | " + str(bib_format_problems(hand)))

    # A8 legends_cover_plan structural
    p = _paper(tmp / "a8", "S07_artifacts")
    _w(p, "01_protocol/artifact_plan.json", {
        "main_figures": [{"id": "Figure 1", "file": "05_figures/out/F1.png"}],
        "main_tables": [{"id": "Table 1", "file": "04_tables/main/T1.xlsx"},
                        {"id": "Table 2", "file": "04_tables/main/T2.xlsx",
                         "footnotes": ["HR, hazard ratio; estimated by Cox regression."]}],
        "supp_figures": [], "supp_tables": []})
    _w(p, "05_figures/legends.md", "# Figure legends\n\n## Figure 1. Flow\n\n"
       "Selection of the analysed reports from the source database.\n")
    _w(p, "04_tables/table_captions.md", "# Table captions\n\n## Table 1. Characteristics\n\n"
       "Values are n (%) of reports; percentages use the column total as denominator.\n")
    ok = _gate(p, "legends_cover_plan", "S07_artifacts")
    _w(p, "04_tables/table_captions.md", "# Table captions\n")
    bad = _gate(p, "legends_cover_plan", "S07_artifacts")
    record("A8 legends_cover_plan: structural per-table footnote entry (captions or plan), no "
           "search for the words 'abbrevi'/'footnote'",
           ok.ok and not bad.ok and "Table 1" in bad.detail and "Table 2" not in bad.detail,
           ok.detail + " | " + bad.detail)


# ------------------------------------------------------------------------------ group B
def run_b(tmp: Path, record, section) -> None:
    import tomllib
    import selftest_v190 as v190
    from wfcore import registry
    section("1.12.0 B: deleted and merged gates")
    raw = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
    word_mins = [g for s in raw["stage"] for g in s.get("gate", [])
                 if g["check"] == "md_wordcount" and ("min_key" in g or "min" in g)]
    record("B1 no section minimum-word warning remains (upper soft budgets only)",
           not word_mins and not any(k.endswith("_words_min") for k in raw["targets"]),
           str(word_mins))

    # B2 deferral needs no decision; the S23 debt gate still counts the unrecorded look
    p = _paper(tmp / "b2", "S23_package")
    debt = _gate(p, "visual_debts_cleared", "S23_package")
    from wfcore.state import State
    st = State(p, ".wf").load()
    for name in ("tables_visually_confirmed", "figures_visually_confirmed"):
        st.record_decision(name, "YES", "looked at the contact sheet and every display at full size")
    cleared = _gate(p, "visual_debts_cleared", "S23_package")
    record("B2 an unrecorded (default-deferred) look is still owed at S23 and clears with YES",
           not debt.ok and "not recorded yet" in debt.detail and "tables_visually_confirmed"
           in debt.detail and cleared.ok, debt.detail[:120] + " | " + cleared.detail[:80])

    # B3 one polish gate, one pending-outputs line, short default gate output
    p = _paper(tmp / "b3", "S22_polish")
    _w(p, "07_manuscript/full_manuscript.md", "# Title\n\n# Introduction\n\nText.\n")
    pol = _gate(p, "polish_complete", "S22_polish")
    tells = _gate(p, "ai_tells_clean", "S22_polish")
    s22 = [g["check"] for g in registry.load().stage("S22_polish").gate]
    record("B3 polish_snapshot_exists/polish_report/polish_preserves_facts are ONE gate with "
           "sub-reasons; report-reading gates do not repeat the missing report",
           not pol.ok and "snapshot:" in pol.detail and "report:" in pol.detail
           and tells.ok and "pending" in tells.detail and "polish_complete" in s22
           and "polish_snapshot_exists" not in s22 and "polish_preserves_facts" not in s22,
           pol.detail[:160])
    q = _paper(tmp / "b3q", "S03_protocol")
    out = v190._run(q, "tools/wf.py", "check")
    record("B3 default `wf check`: failures + one summary line, passes only in the log",
           "[PASS]" not in out.stdout and "outputs_pending" in out.stdout and "RED" in out.stdout
           and (q / ".wf/logs/check_S03_protocol.log").is_file(), out.stdout[-300:])

    # B4 docx_keep_with_next merged into docx_bundle_ready
    s23 = registry.load().stage("S23_package").gate
    record("B4 docx_keep_with_next is part of docx_bundle_ready (one gate line at S23)",
           not any(g["check"] == "docx_keep_with_next" for g in s23)
           and any(g["check"] == "docx_bundle_ready" and g.get("keep_with_next") for g in s23))
    pipe = registry.load()
    names = {s.id: [g["check"] for g in s.gate] for s in pipe.stages}
    dup_free = (all("explicit_confirmation" not in names[s] for s in ("S19_human_review",
                                                                       "S24_package_human_review"))
                and all("review_round_approved" in names[s] for s in ("S19_human_review",
                                                                       "S24_package_human_review"))
                and not {"supplementary_methods_clean", "journal_limits_met",
                         "title_page_reference_count"} & set(names["S23_package"])
                and "submission_checklist_measured" in names["S23_package"]
                and "docx_bundle_ready" in names["S23_package"]
                and not {"journal_limits_met", "title_page_reference_count"} & set(names["S25_submission_audit"]))
    total = sum(len(v) for v in names.values())
    # 1.14.0 adds three new checks (writer bundle completeness, additional-analysis
    # justification, one contact-sheet look before Results); the merged duplicates stay merged.
    new_1140 = sum(1 for v in names.values() for c in v if c in {
        "writer_results_mapped", "additional_analyses_justified", "displays_contact_sheet_seen"})
    record("B4 duplicate gate lines merged (confirmation inside review_round_approved; SI clean "
           "inside docx_bundle_ready; limits/reference count as checklist rows); fewer gate "
           f"entries than 1.11.0 (287 -> {total - new_1140} + {new_1140} new in 1.14.0)",
           dup_free and total - new_1140 < 287, str(total))


# ------------------------------------------------------------------------------ group C
def _mojibake(text: str) -> str:
    for codec in ("gbk", "gb18030", "latin-1"):
        try:
            return text.encode("utf-8").decode(codec)
        except UnicodeDecodeError:
            continue
    return text


def run_c(tmp: Path, record, section) -> None:
    import selftest_v190 as v190
    from wfcore import argtext, reviewrounds as rr
    from wfcore.state import State
    section("1.12.0 C: CLI fixes")
    phrase = "我确认这一版"
    moj = _mojibake(phrase)
    fixed = argtext.fix_argv(["decide", "no_further_review", "--words", moj])
    known = argtext.repair("鏃犻渶缁х画瀹￠槄")
    bad = None
    try:
        argtext.verify("确认\ufffd", "--words")
    except ValueError as exc:
        bad = str(exc)
    p = _paper(tmp / "c1", "S05_analysis")
    why = "用户在会议中口头确认：主要分析采用预先设定的模型，不做额外的事后分组，理由已记录在方案中。"
    _w(p, "temp_why.txt", why + "\n")
    out = v190._run(p, "tools/wf.py", "decide", "model_choice", "YES", "--why-file", p / "temp_why.txt")
    stored = State(p, ".wf").load().decision("model_choice") or {}
    record("C1 mis-decoded Chinese --words/--why is repaired, --why-file carries UTF-8 text, "
           "U+FFFD is refused; the confirmation sentence is still recognised",
           moj != phrase and fixed[-1] == phrase and rr.is_confirmation(fixed[-1])
           and known == "无需继续审阅" and bad and "--words-file" in bad
           and stored.get("rationale") == why and out.returncode == 0,
           f"{fixed[-1]!r} {known!r} {out.stdout[-200:]}")

    # C2 engine pin vs pipeline_version
    m = tmp / "c2"
    m.mkdir(parents=True, exist_ok=True)
    st = State(m, ".wf").create("medpaper", "1.11.0", "S05_analysis")
    st.migrate_pipeline("1.12.0", ["S05_analysis"])
    auto_pin, auto_ver = st.engine_pin, st.data["pipeline_version"]
    detail = st.data["events"][-1]["detail"] if st.data["events"][-1]["kind"] == "pipeline_migration" else \
        next(e["detail"] for e in st.data["events"] if e["kind"] == "pipeline_migration")
    st.migrate_pin("1.12.0")
    q = _paper(tmp / "c2s", "S05_analysis")
    status = v190._run(q, "tools/wf.py", "status", "--brief")
    record("C2 an automatic version event keeps the pin and says so; `wf migrate --yes` moves pin "
           "and pipeline_version together; `wf status` shows both",
           auto_pin == "1.11.0" and auto_ver == "1.12.0" and "stays pinned to 1.11.0" in detail
           and st.engine_pin == st.data["pipeline_version"] == "1.12.0"
           and "paper pinned to" in status.stdout and "state recorded by" in status.stdout,
           detail + " | " + status.stdout[:300])

    # C3 loops on every rewind / re-entry
    r = tmp / "c3"
    r.mkdir(parents=True, exist_ok=True)
    st = State(r, ".wf").create("medpaper", "1.12.0", "S10_tables")
    st.complete("S10_tables", "S11_figures")
    st.complete("S11_figures", "S09_results")
    st.rewind("S10_tables", "fixture rewind")
    st.complete("S10_tables", "S11_figures")
    src = (ROOT / "tools/rework.py").read_text(encoding="utf-8")
    record("C3 stage loops count every rewind and re-entry (incl. the confirmed S19 rewind)",
           st.stage_info("S10_tables")["loops"] == 1 and st.stage_info("S11_figures")["loops"] == 1
           and 's19["loops"] = s19.get("loops", 0) + 1' in src,
           str({k: v.get("loops") for k, v in st.data["stages"].items()}))

    # C4 temp/logs is the log dir, not scratch
    p = _paper(tmp / "c4", "S10_tables")
    _w(p, "temp/logs/repro.log", "log")
    logs_ok = _gate(p, "temp_clean", "S10_tables")
    _w(p, "temp/scratch.csv", "x")
    scratch = _gate(p, "temp_clean", "S10_tables")
    record("C4 temp_clean ignores temp/logs/ (official logs: .wf/logs/), still catches scratch",
           logs_ok.ok and not scratch.ok and "scratch.csv" in scratch.detail, scratch.detail)

    # C5 reference_count show before the title page exists
    p = _paper(tmp / "c5", "S14_introduction")
    show = v190._run(p, "tools/manuscript/reference_count.py", "show")
    record("C5 reference_count.py show reports 'not yet available' instead of crashing",
           show.returncode == 0 and "not yet available" in show.stdout, show.stdout + show.stderr[-200:])

    # C6 status JSON + delta round while the previous awaits the user
    import selftest_v180 as v180
    p, _ = v180._paper(tmp / "c6", None)
    plan = json.loads((p / "01_protocol/artifact_plan.json").read_text(encoding="utf-8"))
    plan["supp_figures"], plan["supp_tables"] = [], []
    _w(p, "01_protocol/artifact_plan.json", plan)
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\nText.\n")
    data = rr.load(p)
    data["manuscript"].append({"id": "M001", "scope": "manuscript", "mode": "full", "stage": "S19",
                               "started_at": "2026-01-01T00:00:00Z", "inputs": {}, "reviewed": [],
                               "closed_at": "2026-01-01T01:00:00Z", "after": {},
                               "user_decision": None, "user_words": "ctrl\x01\x1f" + "x" * 5000})
    rr._save(p, data)
    stat = v190._run(p, "tools/review_round.py", "status", "--scope", "manuscript")
    try:
        parsed = json.loads(stat.stdout)
    except ValueError:
        parsed = None
    refused = ""
    try:
        rr.start(p, "manuscript", "S19_human_review", delta=False)
    except ValueError as exc:
        refused = str(exc)
    nxt = rr.start(p, "manuscript", "S19_human_review", delta=True)
    record("C6 review_round status is valid JSON (control chars, long text); a delta round may "
           "start while the previous awaits the user, linked to it",
           parsed is not None and parsed["latest"][0]["id"] == "M001" and "--delta" in refused
           and nxt.get("linked_to") == "M001" and rr.load(p)["manuscript"][0].get("superseded_by") == nxt["id"],
           stat.stdout[:200] + " | " + refused)

    # C7 Word line numbers
    import zipfile
    from xml.etree import ElementTree as ET
    from docx import Document
    sys.path.insert(0, str(ROOT / "tools/manuscript"))
    import build_docx
    from wfcore import docxaudit, journalspec
    d = tmp / "c7"
    d.mkdir(parents=True, exist_ok=True)
    style = {"font_family": "Arial", "body_font_pt": 11, "title_font_pt": 14,
             "section_heading_font_pt": 12, "subsection_heading_font_pt": 11, "line_spacing": 2.0,
             "paragraph_spacing_after_pt": 0, "margins_in": 1, "paper_size": "A4"}
    roots = {}
    for mode in ("continuous", "per_page", "none"):
        doc = Document()
        doc.add_paragraph("Title")
        doc.add_paragraph("Body text.")
        doc.add_section()
        doc.add_paragraph("Second section.")
        path = d / f"{mode}.docx"
        doc.save(path)
        build_docx.normalize_docx(path, style, "manuscript", line_numbers=mode)
        with zipfile.ZipFile(path) as zf:
            roots[mode] = ET.fromstring(zf.read("word/document.xml"))
    record("C7 target_journal line_numbers -> w:lnNumType in every section; the audit checks it",
           journalspec.line_numbers({"line_numbers": "per-page"}) == "per_page"
           and journalspec.line_numbers({}) == "none"
           and not docxaudit.line_number_problems(roots["continuous"], "continuous")
           and not docxaudit.line_number_problems(roots["per_page"], "per_page")
           and docxaudit.line_number_problems(roots["none"], "continuous")
           and docxaudit.line_number_problems(roots["per_page"], "continuous"))

    # C8 heartbeat cadence
    from wfcore import longrun
    code = ("import sys; sys.path.insert(0, %r); from wfcore import longrun; "
            "raise SystemExit(longrun.run_quiet([sys.executable, '-c', 'import time; time.sleep(3.3)'], "
            "%r, 'slow'))" % (str(ROOT / "tools"), str(tmp / "c8.log")))
    t0 = time.monotonic()
    beat = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                          env={**os.environ, "MEDPAPER_HEARTBEAT": "1"})
    beats = [ln for ln in beat.stdout.splitlines() if "running" in ln]
    with _env(MEDPAPER_HEARTBEAT=120):
        clamped = longrun._heartbeat()
    record("C8 --wait heartbeat <= 30 s (default 25 s, clamped), first line printed at start",
           longrun.HEARTBEAT_SECONDS <= 30 and clamped == 30.0 and len(beats) >= 2
           and "started" in beat.stdout and time.monotonic() - t0 < 30, beat.stdout[-300:])


# ------------------------------------------------------------------------------ group D
def run_d(tmp: Path, record, section) -> None:
    import re
    from wfcore import registry
    section("1.12.0 D: slim stage cards and skill")
    cards = sorted((ROOT / "pipeline/stages").glob("S*.md"))
    sizes = {c.stem: len(c.read_text(encoding="utf-8")) for c in cards}
    over = {k: v for k, v in sizes.items() if v > registry.CARD_MAX_CHARS}
    record(f"D1 every stage card <= {registry.CARD_MAX_CHARS} characters (max {max(sizes.values())})",
           len(cards) == 25 and not over, str(over))
    texts = {c.stem: c.read_text(encoding="utf-8") for c in cards}
    skill = (ROOT / ".kiro/skills/medpaper-pipeline/SKILL.md").read_text(encoding="utf-8")
    refs = sorted({m for t in [*texts.values(), skill]
                   for m in re.findall(r"reference/[\w./-]+\.(?:md|toml)", t)})
    missing = [r for r in refs if not (ROOT / r).is_file()]
    record("D1 every reference/ file a card or the skill points to exists",
           refs and not missing, str(missing))
    writers = [s for s in ("S08_methods", "S09_results", "S14_introduction", "S16_discussion",
                           "S19_human_review")
               if "medpaper-clinical-writer" not in texts[s] or "writer_bundle.py" not in texts[s]]
    stale = [k for k, t in {**texts, "SKILL": skill}.items()
             if "temp/check.log" in t or "temp/repro.log" in t or "visually_confirmed DEFER" in t]
    record("D1 cards delegate prose to the clinical writer, carry no DEFER recording and no "
           "temp/ log paths", not writers and not stale, f"{writers} {stale}")
    record("D2 SKILL.md is slim and has the execution-economy section (handoff, --wait, logs, "
           "one reproduce per round, cheaper models)",
           len(skill) <= 9000 and "Execution economy" in skill and "wf handoff" in skill
           and "--wait" in skill and ".wf/logs/" in skill and "reproduce once per round" in skill
           and "cheaper model" in skill and (ROOT / "reference/skill_details.md").is_file(),
           str(len(skill)))


# ------------------------------------------------------------------------------ group E
def _review_paper(base: Path) -> Path:
    import selftest_v180 as v180
    p, _ = v180._paper(base, None)
    plan = json.loads((p / "01_protocol/artifact_plan.json").read_text(encoding="utf-8"))
    plan["supp_figures"], plan["supp_tables"] = [], []
    _w(p, "01_protocol/artifact_plan.json", plan)
    _w(p, "07_manuscript/full_manuscript.md", "# Title\n\n# Results\n\nText of the results.\n")
    return p


def run_e(tmp: Path, record, section) -> None:
    import selftest_v190 as v190
    from wfcore import adapters, reviewrounds as rr
    section("1.12.0 E: wf handoff, reviewer report on disk, clinical writer")
    p = _paper(tmp / "e1", "S09_results")
    h = v190._run(p, "tools/wf.py", "handoff")
    md = v190._run(p, "tools/wf.py", "handoff", "--md")
    lines = h.stdout.strip().splitlines()
    record("E1 `wf handoff` prints a compact pack: paper, versions, stage + card, gate summary, "
           "questions, last events, next command (<= 60 lines; --md)",
           h.returncode == 0 and len(lines) <= 60 and "stage: S09_results" in h.stdout
           and "paper pinned to" in h.stdout and "gate:" in h.stdout and "next:" in h.stdout
           and "last events:" in h.stdout and md.stdout.startswith("# medpaper handoff"),
           h.stdout[-400:] + h.stderr[-200:])

    # E2 reviewer writes only its report; close reads it by default
    agent = (ROOT / ".kiro/agents/medpaper-reviewer.md").read_text(encoding="utf-8")
    q = _review_paper(tmp / "e2")
    st = v190._run(q, "tools/review_round.py", "start", "--scope", "manuscript")
    rid = (rr.latest(q, "manuscript") or {}).get("id", "")
    _w(q, f".wf/review_bundle/{rid}/report.md", "# Review\n\n## Verdict\nReady; no findings.\n")
    cl = v190._run(q, "tools/review_round.py", "close", "--scope", "manuscript")
    last = rr.latest(q, "manuscript") or {}
    record("E2 reviewer agent may write only .wf/review_bundle/<round>/report.md (paper areas "
           "denied); start names the file; close --report defaults to it",
           '"write"' in agent and ".wf/review_bundle/*/report.md" in agent
           and 'match: ["0*/**"' in agent and "capability: shell" in agent
           and "report.md" in st.stdout and cl.returncode == 0
           and last.get("report") == f".wf/review_reports/{rid}.md",
           st.stdout[-200:] + cl.stdout[-200:] + cl.stderr[-200:])

    # E3 clinical writer agent + writer bundle
    writer = (ROOT / ".kiro/agents/medpaper-clinical-writer.md").read_text(encoding="utf-8")
    sys.path.insert(0, str(ROOT / "tools/manuscript"))
    import writer_bundle as wb
    record("E3 medpaper-clinical-writer is installed (global agents) and cannot read the "
           "analysis layer or run commands",
           "medpaper-clinical-writer.md" in adapters.GLOBAL_AGENTS
           and '"03_analysis/**"' in writer and "capability: shell" in writer
           and "does not apply to you" in writer and "07_manuscript/*.md" in writer)
    p = _paper(tmp / "e3", "S09_results")
    _w(p, "01_protocol/display_dictionary.json", {"terms": [
        {"internal": "comp_a", "aliases": ["Composite A"], "label": "Portal vascular disease",
         "abbreviation": "PVD", "definition": "portal vein thrombosis or nodular hyperplasia"},
        {"internal": "exposed", "label": "Exposed group", "definition": "reports naming the drug"}]})
    _w(p, "03_analysis/results/main.json", {"comp_a": {"exposed": {"hr": 1.87, "hr_low": 1.2,
                                                                   "strict_contrast_x": 4}},
                                            "raw_lines_read": 24812, "environment": {"python": "3"}})
    _w(p, "01_protocol/protocol_final.md", "# Protocol\n\nWe ran `parse_raw.py` from "
       "03_analysis/code/x.py; comp_a was compared using age_years.\n")
    out = wb.build(p, "S09_results")
    texts = {f.name: f.read_text(encoding="utf-8") for f in out.glob("*.md")}
    blob = "\n".join(texts.values())
    lint = wb.lint(p, "S09_results")
    record("E3 writer bundle: clinical labels and result sentences only - no internal keys, "
           "paths or code; its lint passes",
           "Hazard ratio, Portal vascular disease, Exposed group: 1.87" in blob
           and not any(k in blob for k in ("comp_a", "strict_contrast", "raw_lines_read",
                                           "age_years", ".py", "03_analysis", "environment"))
           and "Write: 07_manuscript/results.md" in texts.get("MANIFEST.md", "") and not lint,
           str(lint[:4]) + " | " + texts.get("results_in_words.md", "")[:300])


# ------------------------------------------------------------------------------ group F
def run_f(tmp: Path, record, section) -> None:
    from wfcore import registry, reviewrounds as rr
    section("1.12.0 F: earlier, cheap, deterministic detection")
    p = _paper(tmp / "f1", "S12_reconcile")
    _w(p, "01_protocol/artifact_plan.json", {
        "main_figures": [], "supp_figures": [], "supp_tables": [],
        "main_tables": [{"id": "Table 1", "file": "04_tables/main/T1.xlsx"}]})
    _xlsx(p, "04_tables/main/T1.xlsx", [["Characteristic", "Exposed"], ["Age", "61"],
                                         ["Female sex", "64 (40%)"]])
    _w(p, "04_tables/table_captions.md", "# Table captions\n\n## Table 1. Characteristics\n\n"
       "Values are counts with percentages\n")
    bad = _gate(p, "first_reader_displays", "S12_reconcile",
                include=["abbreviations", "engineering", "flow_units"])
    _xlsx(p, "04_tables/main/T1.xlsx", [["Characteristic", "Exposed"], ["Age, years", "61"],
                                         ["Female sex", "64 (40%)"]])
    _w(p, "04_tables/table_captions.md", "# Table captions\n\n## Table 1. Characteristics\n\n"
       "Values are n (%) of the 160 exposed patients.\n")
    good = _gate(p, "first_reader_displays", "S12_reconcile",
                 include=["abbreviations", "engineering", "flow_units"])
    specs = {s.id: [g for g in s.gate if g["check"] == "first_reader_displays"]
             for s in registry.load().stages}
    record("F1 first_reader_displays: units, denominator, end punctuation per display; warn at "
           "S10/S11, blocking (with abbreviations/language/flow units) at S12",
           not bad.ok and "unit not stated" in bad.detail and "denominator" in bad.detail
           and "punctuation" in bad.detail and good.ok
           and specs["S10_tables"][0].get("severity") == "warn"
           and specs["S11_figures"][0].get("severity") == "warn"
           and "severity" not in specs["S12_reconcile"][0], bad.detail[:200] + " | " + good.detail)

    p = _paper(tmp / "f2", "S07_artifacts")
    plan = {"main_figures": [{"id": "Figure 1", "file": "a.png"}, {"id": "Figure 2", "file": "b.png"}],
            "main_tables": [], "supp_figures": [], "supp_tables": []}
    _w(p, "01_protocol/artifact_plan.json", plan)
    missing = _gate(p, "display_items_justified", "S07_artifacts")
    plan["main_figures"][1]["non_redundancy"] = "shows the time course that Figure 1 cannot"
    _w(p, "01_protocol/artifact_plan.json", plan)
    ok = _gate(p, "display_items_justified", "S07_artifacts")
    wired = [s.id for s in registry.load().stages
             if any(g["check"] == "display_items_justified" for g in s.gate)]
    s12 = next(g for g in registry.load().stage("S12_reconcile").gate
               if g["check"] == "first_reader_displays")
    record("F2 an added display item needs a one-line non-redundancy justification (S07, "
           "re-checked at S12 inside the first-reader line)", not missing.ok
           and "Figure 2" in missing.detail and ok.ok and wired == ["S07_artifacts"]
           and "non_redundancy" in s12.get("include", []), missing.detail + str(wired))

    q = _review_paper(tmp / "f3")
    with _env(MEDPAPER_NO_PREVIEW=1):
        rnd = rr.start(q, "manuscript", "S18_independent_review")
        out = rr.write_bundle(q, rnd)
    man = (out / "MANIFEST.md").read_text(encoding="utf-8")
    built = (out / "draft/manuscript_draft.docx").is_file()
    record("F3 from S17 the review bundle carries a Word build of the draft (+ text copy) or "
           "states why not; previews are skipped/unavailable with a note",
           (built and "manuscript_draft.docx" in man and (out / "files/manuscript_draft.docx.txt").is_file()
            and "preview skipped" in man) or (not built and "DOCX draft build unavailable" in man),
           man[-400:])


def run(tmp: Path, record, section) -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    run_a(tmp, record, section)
    run_b(tmp, record, section)
    run_c(tmp, record, section)
    run_d(tmp, record, section)
    run_e(tmp, record, section)
    run_f(tmp, record, section)

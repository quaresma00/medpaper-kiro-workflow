"""1.13.0 regression tests from two real paper runs: isolation and working path (G1), data
never copied and one reproduction (G2), rewind rule / multi-person checks / protocol deviations
(G3), section polishing and Word without keep markers (G4), extras (G5). Synthetic fixtures
only - no real paper or dataset is read.

Called from tools/selftest.py.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from contextlib import contextmanager, redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
WF = ROOT / "tools/wf.py"


def _env(**extra) -> dict:
    env = {k: v for k, v in os.environ.items() if k not in ("MEDPAPER_PROJECT", "MEDPAPER_ROOT")}
    env["PYTHONIOENCODING"] = "utf-8"
    for k, v in extra.items():
        if v is None:
            env.pop(k, None)
        else:
            env[k] = str(v)
    return env


def _run(cmd, env=None, cwd=None, inp=None, timeout=600) -> subprocess.CompletedProcess:
    return subprocess.run([str(c) for c in cmd], capture_output=True, text=True, encoding="utf-8",
                          errors="replace", env=env or _env(), cwd=str(cwd) if cwd else None,
                          input=inp, timeout=timeout)


def _wf(*args, env=None, cwd=None) -> subprocess.CompletedProcess:
    return _run([sys.executable, WF, *args], env=env, cwd=cwd)


def _out(p) -> str:
    return (p.stdout or "") + (p.stderr or "")


def _paper(path: Path) -> Path:
    from wfcore import registry
    from wfcore.state import State
    pipe = registry.load()
    path.mkdir(parents=True, exist_ok=True)
    State(path, ".wf").create(pipe.meta["name"], pipe.meta["version"], pipe.first().id)
    return path


@contextmanager
def _envset(**kv):
    old = {k: os.environ.get(k) for k in kv}
    for k, v in kv.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = str(v)
    try:
        yield
    finally:
        for k, v in old.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def _w(p: Path, rel: str, text) -> Path:
    t = p / rel
    t.parent.mkdir(parents=True, exist_ok=True)
    t.write_text(text if isinstance(text, str) else json.dumps(text), encoding="utf-8")
    return t


# ------------------------------------------------------------------------------ G1
def run_g1(tmp: Path, record, section) -> None:
    from wfcore import cli, datasources, integrity, paths
    section("1.13.0 G1: working path in place, isolation, engine integrity, engine fork")
    papers_root = Path(os.environ["MEDPAPER_PAPERS_ROOT"])

    ideas = tmp / "ideas"
    ws = ideas / "anticoag-idea"
    ws.mkdir(parents=True)
    (ws / "idea.md").write_text("# idea\n", encoding="utf-8")
    p = _wf("init", cwd=ws)
    record("G1 wf init makes the opened workspace folder the paper, in place (user files kept "
           "and never orphans; paper guard hook installed)",
           p.returncode == 0 and (ws / ".wf/state.json").is_file()
           and (ws / ".kiro/hooks/medpaper-paper-guard.json").is_file()
           and "idea.md" in paths.user_files(ws) and (ws / "idea.md").is_file(), _out(p)[-200:])
    c = _wf("clean", cwd=ws)
    record("G1 the user's own files are not reported as undeclared", "idea.md" not in c.stdout,
           c.stdout[-150:])
    ws2 = tmp / "ideas2" / "second"
    ws2.mkdir(parents=True)
    p = _wf("new", "second-paper", cwd=ws2)
    record("G1 wf new without --root/--central initialises the working directory in place",
           p.returncode == 0 and (ws2 / ".wf/state.json").is_file()
           and not (papers_root / "second-paper").exists(), _out(p)[-160:])
    p = _wf("init", cwd=ideas)
    record("G1 a folder that already holds another paper is refused, asks the user once and "
           "names no other paper", p.returncode == 2 and "already contains" in p.stdout
           and "anticoag-idea" not in p.stdout and "Ask the user once" in p.stdout,
           p.stdout[-220:])
    record("G1 a drive root and the home folder are never a paper",
           "drive root" in (paths.unsuitable_paper_dir(Path(ws.anchor)) or "")
           and "home folder" in (paths.unsuitable_paper_dir(Path.home()) or ""))

    # data sources in place
    data = tmp / "userdata" / "raw"
    data.mkdir(parents=True)
    (data / "a.csv").write_text("x\n1\n", encoding="utf-8")
    inner = ws / "datacopy"
    inner.mkdir()
    p = _wf("data", "add", str(data), "--name", "faers", cwd=ws)
    s = _wf("status", "--brief", cwd=ws)
    h = _wf("handoff", cwd=ws)
    record("G1 wf data add registers a dataset in place; status and handoff print the paper, "
           "the declared data sources and the engine path",
           p.returncode == 0 and "faers" in s.stdout and str(data) in s.stdout
           and "engine  " in s.stdout and "data: faers" in h.stdout and "engine:" in h.stdout
           and "paper:" in h.stdout, s.stdout[:240] + " | " + h.stdout[:200])
    other = _paper(tmp / "neighbours" / "other-paper")
    probs = [datasources.location_problem(ws, ideas), datasources.location_problem(ws, inner),
             datasources.location_problem(ws, other), datasources.location_problem(ws, "rel/x")]
    record("G1 a data folder holding the paper, a data folder whose parent is the paper, "
           "another paper's folder and a relative path are refused",
           all(probs) and "inside the data folder" in probs[0] and "parent" in probs[1]
           and "another paper" in probs[2] and "absolute" in probs[3], str(probs)[:300])
    calls = []
    real_sha = datasources._raw_sha

    def counting(path):
        calls.append(path)
        return real_sha(path)
    datasources._CACHES.clear()
    with patch.object(datasources, "_raw_sha", counting):
        first = datasources.inventory(ws)
        datasources._CACHES.clear()
        again = datasources.inventory(ws)
    record("G1 dataset hashes are computed lazily, once, and cached by size + mtime",
           len(calls) == 1 and first == again and first[0]["path"] == "datasource:faers/a.csv",
           str(first)[:160])

    # engine integrity
    with _envset(MEDPAPER_ENGINE_DEV=None):
        g = tmp / "git_engine"
        (g / "tools").mkdir(parents=True)
        (g / "tools/a.py").write_text("x = 1\n", encoding="utf-8")
        git = ["git", "-c", "user.name=t", "-c", "user.email=t@example.invalid",
               "-c", "commit.gpgsign=false", "-C", str(g)]
        _run(["git", "init", "-q", str(g)])
        _run([*git, "add", "-A"])
        _run([*git, "commit", "-q", "-m", "x"])
        clean = integrity.engine_state(g, use_cache=False)
        (g / "tools/a.py").write_text("x = 2\n", encoding="utf-8")
        modified = integrity.engine_state(g, use_cache=False)
        m = tmp / "zip_engine"
        (m / "tools").mkdir(parents=True)
        (m / "tools/a.py").write_text("x = 1\n", encoding="utf-8")
        unverifiable = integrity.engine_state(m, use_cache=False)[0]
        integrity.write_manifest(m)
        m_clean = integrity.engine_state(m, use_cache=False)[0]
        (m / "tools/b.py").write_text("y = 1\n", encoding="utf-8")
        m_mod = integrity.engine_state(m, use_cache=False)
        wt = tmp / "wt_engine"
        wt.mkdir()
        (wt / ".git").write_text("gitdir: elsewhere\n", encoding="utf-8")
        fk = tmp / "fork_engine"
        fk.mkdir()
        (fk / "ENGINE_FORK.json").write_text("{}", encoding="utf-8")
        kinds = (integrity.engine_state(wt, use_cache=False)[0],
                 integrity.engine_state(fk, use_cache=False)[0])
    with _envset(MEDPAPER_ENGINE_DEV="1"):
        dev = integrity.engine_state(g, use_cache=False)[0]
    record("G1 engine integrity: git clean/modified, release-manifest clean/modified, "
           "unverifiable, maintenance worktree, fork and MEDPAPER_ENGINE_DEV",
           clean == ("clean", []) and modified[0] == "modified" and "tools/a.py" in modified[1]
           and unverifiable == "unverifiable" and m_clean == "clean"
           and m_mod[0] == "modified" and "tools/b.py" in m_mod[1]
           and kinds == ("worktree", "fork") and dev == "dev",
           str([clean, modified, unverifiable, m_clean, m_mod, kinds, dev])[:300])
    buf = StringIO()
    with patch.object(integrity, "engine_state",
                      return_value=("modified", ["tools/wfcore/checks/numbers.py"])), \
            _envset(MEDPAPER_PROJECT=ws):
        try:
            paths.project_dir()
            refused = ""
        except integrity.EngineModified as exc:
            refused = str(exc)
        with redirect_stdout(buf):
            rc = cli.main(["check"])
    record("G1 a modified global engine refuses paper commands and names the legal paths "
           "(paper overrides first, then a project-local fork)",
           "tools/wfcore/checks/numbers.py" in refused and "wf config set" in refused
           and "engine-fork" in refused and rc == 3 and "REFUSED" in buf.getvalue(),
           refused[:200])
    copy = tmp / "engine_copy"
    for tree in ("tools", "pipeline", "reference", ".kiro"):
        shutil.copytree(ROOT / tree, copy / tree,
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    integrity.write_manifest(copy)
    tgt = copy / "tools/wfcore/numfmt.py"
    tgt.write_text(tgt.read_text(encoding="utf-8") + "\n# local edit\n", encoding="utf-8")
    p = _run([sys.executable, copy / "tools/wf.py", "status", "--brief"],
             env=_env(MEDPAPER_ENGINE_DEV=None), cwd=ws)
    q = _run([sys.executable, copy / "tools/reproduce.py", "check"],
             env=_env(MEDPAPER_ENGINE_DEV=None), cwd=ws)
    record("G1 end to end: an engine that differs from its release manifest refuses wf and "
           "reproduce.py", p.returncode == 3 and "REFUSED" in p.stdout
           and "tools/wfcore/numfmt.py" in p.stdout and q.returncode != 0 and "REFUSED" in _out(q),
           _out(p)[-200:] + " | " + _out(q)[-120:])
    shutil.rmtree(copy, ignore_errors=True)

    # project-local engine fork
    fp = tmp / "fork_ws"
    fp.mkdir()
    _wf("init", cwd=fp)
    p = _wf("engine-fork", "create", cwd=fp)
    fork = fp / ".wf/engine_fork"
    try:
        info = json.loads((fork / "ENGINE_FORK.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        info = {}
    steering = (fp / ".kiro/steering/medpaper-pipeline.md").read_text(encoding="utf-8")
    guard = (fp / ".kiro/hooks/medpaper-paper-guard.json").read_text(encoding="utf-8")
    from wfcore import registry
    ver = str(registry.load().meta.get("version"))
    record("G1 wf engine-fork create copies tools/pipeline/reference into <paper>/.wf/"
           "engine_fork (no .git, no .venv), records base version + hashes and re-points the "
           "paper's adapters", p.returncode == 0 and (fork / "tools/wf.py").is_file()
           and not (fork / ".git").exists() and not (fork / ".venv").exists()
           and info.get("base_version") == ver and len(info.get("files", {})) > 50
           and fork.as_posix() in steering and fork.as_posix() in guard, _out(p)[-200:])
    FW = fork / "tools/wf.py"
    s = _run([sys.executable, FW, "status", "--brief"], cwd=fp)
    record("G1 the fork drives its own paper (only the fork folder counts as engine)",
           s.returncode == 0 and "ENGINE FORK" in s.stdout, _out(s)[-200:])
    code = ("import sys; sys.path.insert(0, sys.argv[1]); from wfcore import paths; "
            "print(paths.in_engine(sys.argv[2]), paths.in_engine(sys.argv[3]), "
            "paths.in_engine(sys.argv[4]))")
    gp = _run([sys.executable, "-c", code, fork / "tools", ROOT / "tools/x.py", fp,
               fork / "tools/x.py"], cwd=fp)
    record("G1 from a fork the global engine stays protected and the paper is not 'over the "
           "engine'", gp.stdout.split() == ["True", "False", "True"], _out(gp)[-160:])
    o = _run([sys.executable, FW, "status", "--brief"], env=_env(MEDPAPER_PROJECT=other), cwd=tmp)
    record("G1 a fork drives no other paper", o.returncode != 0 and "belongs to" in _out(o),
           _out(o)[-160:])
    edit = fork / "tools/wfcore/numfmt.py"
    edit.write_text(edit.read_text(encoding="utf-8") + "\n# paper-local fix\n", encoding="utf-8")
    d = _run([sys.executable, FW, "engine-fork", "diff"], cwd=fp)
    diff_file = fp / ".wf/engine_fork.diff"
    hd = _run([sys.executable, FW, "handoff"], cwd=fp)
    import selftest_v180 as v180
    gate = v180._gate(fp, "engine_fork_recorded", "S25_submission_audit")
    record("G1 wf engine-fork diff records the change for the maintainer; handoff and the S25 "
           "check show it", d.returncode == 0 and "M tools/wfcore/numfmt.py" in d.stdout
           and diff_file.is_file() and "+# paper-local fix" in diff_file.read_text(encoding="utf-8")
           and "1 file(s) differ" in hd.stdout and gate.ok and "1 changed file" in gate.detail,
           _out(d)[-160:] + " | " + gate.detail)
    r = _wf("engine-fork", "remove", cwd=fp)
    steering = (fp / ".kiro/steering/medpaper-pipeline.md").read_text(encoding="utf-8")
    record("G1 wf engine-fork remove deletes the fork, keeps the diff and points the adapters "
           "back at the global engine", r.returncode == 0 and not fork.exists()
           and diff_file.is_file() and ROOT.as_posix() in steering
           and fork.as_posix() not in steering, _out(r)[-160:])
    run_guard(tmp, record, ws)


def run_guard(tmp: Path, record, hook_paper: Path) -> None:
    from wfcore import datasources
    sys.path.insert(0, str(ROOT / "tools/hooks"))
    import paper_guard
    hook = json.loads((hook_paper / ".kiro/hooks/medpaper-paper-guard.json").read_text(
        encoding="utf-8"))["hooks"][0]
    record("G1 the paper guard is a v1 PreToolUse command hook matching the file, search, "
           "shell and subagent tools", hook["trigger"] == "PreToolUse"
           and hook["action"]["type"] == "command" and "paper_guard.py" in hook["action"]["command"]
           and "--paper" in hook["action"]["command"]
           and all(t in hook["matcher"] for t in ("read_file", "list_directory", "grep_search",
                                                   "file_search", "fs_write", "str_replace",
                                                   "delete_file", "smart_relocate",
                                                   "execute_pwsh", "[Ss]ub_?[Aa]gent")))
    g = tmp / "guard"
    cur = _paper(g / "ws" / "paper_a")
    other = _paper(g / "ws" / "paper_b")
    central = g / "central"
    _paper(central / "c1")
    data = g / "data"
    _w(data, "raw.csv", "x\n1\n")
    datasources.add(cur, data, "d")
    archive = g / "medpaper_archive_2026"
    _w(archive, "x.md", "x")
    systemp = g / "systemp"
    (systemp / "medpaper_repro_abc").mkdir(parents=True)
    E = ROOT.as_posix()

    def sh(cmd, cwd=None):
        return {"tool_name": "execute_pwsh", "tool_input": {"command": cmd, "cwd": str(cwd or cur)}}

    cases = [
        ("read own file", {"tool_name": "read_file", "tool_input": {"path": "01_protocol/idea.json"}}, False),
        ("read an engine card", {"tool_name": "read_file", "tool_input": {"path": f"{E}/pipeline/stages/S01_intake.md"}}, False),
        ("fs_write into the engine", {"tool_name": "fs_write", "tool_input": {"path": f"{E}/tools/wfcore/checks/numbers.py", "text": "x"}}, True),
        ("str_replace in the engine", {"tool_name": "str_replace", "tool_input": {"path": f"{E}/tools/wfcore/bibrules.py", "oldStr": "a", "newStr": "b"}}, True),
        ("read another paper", {"tool_name": "read_file", "tool_input": {"path": str(other / "02_data/acquisition/01_acquire.py")}}, True),
        ("list a folder holding other papers", {"tool_name": "list_directory", "tool_input": {"path": str(g / "ws")}}, True),
        ("relative escape to a sibling paper", {"tool_name": "read_file", "tool_input": {"path": "../paper_b/notes.md"}}, True),
        ("grep inside another paper", {"tool_name": "grep_search", "tool_input": {"query": "x", "includePattern": str(other) + "/**/*.py"}}, True),
        ("glob over sibling papers", {"tool_name": "file_search", "tool_input": {"query": "x", "includePattern": str(g / "ws") + "/*/02_data"}}, True),
        ("the central papers root", {"tool_name": "list_directory", "tool_input": {"path": str(central)}}, True),
        ("an archive folder", {"tool_name": "read_file", "tool_input": {"path": str(archive / "x.md")}}, True),
        ("read a declared data source", {"tool_name": "read_file", "tool_input": {"path": str(data / "raw.csv")}}, False),
        ("write a declared data source", {"tool_name": "fs_write", "tool_input": {"path": str(data / "raw.csv"), "text": "x"}}, True),
        ("write to an invented location", {"tool_name": "fs_write", "tool_input": {"path": str(g / "studies/new/a.md"), "text": "x"}}, True),
        ("delete a paper file", {"tool_name": "delete_file", "tool_input": {"targetFile": "temp/x.txt"}}, False),
        ("relocate out of the paper", {"tool_name": "smart_relocate", "tool_input": {"sourcePath": "temp/x.txt", "destinationPath": str(g / "elsewhere.txt")}}, True),
        ("shell: read another paper", sh(f'Get-Content "{other}\\notes.md"'), True),
        ("shell: engine command, log in the paper", sh(f'python "{E}/tools/wf.py" status *> .wf/logs/status.txt'), False),
        ("shell: write into the engine", sh(f'Set-Content -Path "{E}/tools/x.py" -Value 1'), True),
        ("shell: git checkout in the engine", sh(f'git -C "{E}" checkout -- tools/wfcore/bibrules.py'), True),
        ("shell: python one-liner writes the engine", sh(f"python -c \"from pathlib import Path; p = Path(r'{E}/tools/x.py'); p.write_text('1')\""), True),
        ("shell: stray log in the system temp", sh(f'python x.py *> "{systemp}\\repro12.txt"'), True),
        ("shell: clean an engine-managed sandbox", sh(f'Remove-Item -Recurse -Force "{systemp}\\medpaper_repro_abc"'), False),
        ("shell: relative escape", sh("Get-ChildItem ..\\paper_b\\02_data"), True),
        ("shell: cwd in another paper", sh("Get-ChildItem", cwd=other), True),
        ("shell: copy an engine template into the paper", sh(f'Copy-Item "{E}/tools/analysis/results.py" temp/results_copy.py'), False),
        ("subagent prompt naming another paper", {"tool_name": "invoke_subagent", "tool_input": {"prompt": f"Compare with {other}\\07_manuscript"}}, True),
        ("subagent prompt about this paper", {"tool_name": "invoke_subagent", "tool_input": {"prompt": "Review .wf/review_bundle/M001/report.md"}}, False),
    ]
    wrong = []
    with _envset(MEDPAPER_PAPERS_ROOT=central, MEDPAPER_GUARD_TEMP=systemp):
        guard = paper_guard.Guard(cur)
        for label, payload, blocked in cases:
            try:
                got = guard.decide(payload)
            except Exception as exc:  # noqa: BLE001
                got = f"ERROR {exc!r}"
                blocked = not blocked
            if bool(got) != blocked:
                wrong.append(f"{label}: {'allowed' if not got else got[:90]}")
    record(f"G1 paper guard: {len(cases)} tool shapes allow the paper/engine reads/data reads/"
           "sandboxes and block other papers, the central root, archives, stray temp, data "
           "writes, invented locations and every engine write", not wrong, "; ".join(wrong)[:600])
    env = _env(MEDPAPER_PAPERS_ROOT=central, MEDPAPER_GUARD_TEMP=systemp)
    gpath = ROOT / "tools/hooks/paper_guard.py"
    deny = _run([sys.executable, gpath, "--paper", cur], env=env, inp=json.dumps(cases[2][1]))
    junk = _run([sys.executable, gpath, "--paper", cur], env=env, inp="not json {")
    ok_ = _run([sys.executable, gpath, "--paper", cur], env=env, inp=json.dumps(cases[0][1]))
    record("G1 paper guard as a process: exit 2 with the reason on stderr; unparseable payload "
           "and allowed calls exit 0", deny.returncode == 2
           and "BLOCKED by the medpaper paper guard" in deny.stderr and "engine-fork" in deny.stderr
           and junk.returncode == 0 and ok_.returncode == 0, deny.stderr[-200:])


# ------------------------------------------------------------------------------ G2
def run_g2(tmp: Path, record, section) -> None:
    import tempfile
    import time
    import zipfile
    from wfcore import dataproof, datasources, logs, reproproof as rp
    section("1.13.0 G2: data used in place (never copied), one reproduction, logs")
    P = _paper(tmp / "repro_paper")
    D = tmp / "userdata"
    D.mkdir(parents=True)
    (D / "big.csv").write_text("x,y\n1,2\n2,4\n", encoding="utf-8")
    with zipfile.ZipFile(D / "arch.zip", "w") as zf:
        zf.writestr("inner/rows.txt", "a$b\n1$2\n3$4\n5$6\n")
    datasources.add(P, D, "src")
    _w(P, "02_data/raw/small.csv", "k\n1\n")
    C = "03_analysis/code/"
    _w(P, C + "lib/util.py", "def total(rows):\n    return sum(len(r) for r in rows)\n")
    _w(P, C + "01_load.py",
       "from wfcore.paths import project_dir\nfrom analysis.datasources import source_path, open_member\n"
       "P = project_dir()\n(P / '02_data/derived').mkdir(parents=True, exist_ok=True)\n"
       "text = source_path('src', 'big.csv').read_text()\n"
       "with open_member('src', 'arch.zip', 'inner/rows.txt', encoding='utf-8') as fh:\n"
       "    n = sum(1 for _ in fh)\n"
       "(P / '02_data/derived/clean.csv').write_text(text + f'#{n}\\n')\n")
    _w(P, C + "02_model.py",
       "import sys, os\nsys.path.insert(0, os.path.dirname(__file__))\nfrom lib.util import total\n"
       "from wfcore.paths import project_dir\nfrom analysis.results import write_results\n"
       "rows = (project_dir() / '02_data/derived/clean.csv').read_text().splitlines()\n"
       "write_results('primary', {'n': len(rows), 'total': total(rows)})\n")
    _w(P, C + "03_other.py", "from wfcore.paths import project_dir\nfrom analysis.results import "
                             "write_results\nP = project_dir()\nwrite_results('other', {'k': 1})\n")
    entries = [{"path": C + "01_load.py", "inputs": ["datasource:src"],
                "outputs": ["02_data/derived/clean.csv"]},
               {"path": C + "02_model.py", "inputs": ["02_data/derived/clean.csv"],
                "outputs": ["03_analysis/results/primary.json"]},
               {"path": C + "03_other.py", "inputs": [], "outputs": ["03_analysis/results/other.json"]}]
    _w(P, "03_analysis/analysis_manifest.json", {"schema": 1, "scripts": entries})
    env = _env(MEDPAPER_PROJECT=P)
    rs = [_wf("run-script", C + s, env=env, cwd=tmp) for s in ("01_load.py", "02_model.py", "03_other.py")]
    RP = [sys.executable, ROOT / "tools/reproduce.py"]
    report = P / "03_analysis/reproducibility_report.json"

    def rep():
        return json.loads(report.read_text(encoding="utf-8")) if report.is_file() else {}

    def started(r):
        return [ln.split()[2] for ln in r.stdout.splitlines() if " start " in ln]
    stale = P / rp.SANDBOX_DIR / "run_stale"
    _w(stale, "x.bin", "x" * 100)
    legacy = Path(tempfile.gettempdir()) / "medpaper_repro_selftest_legacy"
    _w(legacy, "y.bin", "y")
    old = time.time() - 7 * 3600
    os.utime(legacy, (old, old))
    r1 = _run([*RP, "run", "--project", P], cwd=tmp)
    r = rep()
    sb = r.get("sandbox", {})
    inputs = [i["path"] for i in r.get("inputs", [])]
    record("G2 reproduce: declared data read in place, in-paper data hard-linked, 0 bytes of "
           "data copied, sandbox inside the paper and deleted after the run",
           all(x.returncode == 0 for x in rs) and r1.returncode == 0 and r.get("all_match")
           and sb.get("data_bytes_copied") == 0 and sb.get("data_files_linked", 0) >= 1
           and sb.get("data_sources_in_place") == 2 and "datasource:src/big.csv" in inputs
           and "02_data/raw/small.csv" in inputs
           and not any((P / rp.SANDBOX_DIR).iterdir()), _out(r1)[-300:])
    record("G2 leftovers of interrupted runs are removed at the next start (paper sandbox and "
           "old medpaper_repro_* in the system temp)", not stale.exists() and not legacy.exists()
           and "leftover sandbox" in r1.stdout, r1.stdout[:200])
    eng = r.get("engine_inputs", {})
    record("G2 the proof records only the engine helpers the analysis uses (no figure helpers)",
           set(eng) == {"tools/analysis/results.py", "tools/analysis/datasources.py"}, str(sorted(eng)))
    record("G2 an engine release that leaves the used helpers unchanged keeps the proof "
           "(older, wider helper records are tolerated)",
           rp.helpers_match({**eng, "tools/R/medresults.R": "old", "tools/analysis/spotcheck.py": "x"}, eng)
           and not rp.helpers_match({**eng, "tools/analysis/results.py": "changed"}, eng))
    _w(P, "02_data/codebook.md", "# codebook\n")
    _w(P, "02_data/acquisition/fetch.py", "print('acquire')\n")
    _w(P, "03_analysis/notes.md", "notes\n")
    g = _gate(P, "analysis_reproduces", "S17_assemble")
    old_rep = rep()
    old_rep["inputs"].append({"path": "02_data/provenance.md", "sha256": "0" * 64})
    report.write_text(json.dumps(old_rep), encoding="utf-8")
    g_old = _gate(P, "analysis_reproduces", "S23_package")
    record("G2 documentation, acquisition code and notes are not analysis inputs; an older proof "
           "that listed them stays valid; the gate says so by hash",
           g.ok and g_old.ok and "nothing to re-run" in g.detail, g.detail[:160] + " | " + g_old.detail[:80])
    r2 = _run([*RP, "run", "--project", P], cwd=tmp)
    record("G2 nothing changed: REUSED (no second full reproduction)",
           r2.returncode == 0 and "REUSED" in r2.stdout and not started(r2), _out(r2)[-200:])
    _w(P, C + "04_new.py", "from wfcore.paths import project_dir\nfrom analysis.results import "
                           "write_results\nP = project_dir()\nwrite_results('newres', {'k': 2})\n")
    entries.append({"path": C + "04_new.py", "inputs": [], "outputs": ["03_analysis/results/newres.json"]})
    _w(P, "03_analysis/analysis_manifest.json", {"schema": 1, "scripts": entries})
    _wf("run-script", C + "04_new.py", env=env, cwd=tmp)
    hint = _gate(P, "analysis_reproduces", "S17_assemble")
    r3 = _run([*RP, "run", "--project", P], cwd=tmp)
    record("G2 a new analysis script re-runs alone (the others are carried); the gate names the "
           "cheapest next command", r3.returncode == 0 and started(r3) == [C + "04_new.py"]
           and any("Cheapest next step" in h and "1 of 4" in h for h in hint.hints),
           _out(r3)[-200:] + " | " + "; ".join(hint.hints)[:200])
    _w(P, C + "lib/util.py", "def total(rows):\n    return sum(len(r) for r in rows)\n# note\n")
    r4 = _run([*RP, "run", "--project", P], cwd=tmp)
    record("G2 a changed helper re-runs only the scripts that use it",
           r4.returncode == 0 and started(r4) == [C + "02_model.py"], _out(r4)[-200:])
    (D / "big.csv").write_text("x,y\n1,2\n2,4\n3,9\n", encoding="utf-8")
    for s in ("01_load.py", "02_model.py"):
        _wf("run-script", C + s, env=env, cwd=tmp)
    r5 = _run([*RP, "run", "--project", P], cwd=tmp)
    record("G2 a changed data source re-runs only the scripts that declare it and their "
           "consumers", r5.returncode == 0 and started(r5) == [C + "01_load.py", C + "02_model.py"],
           _out(r5)[-200:])
    lock = P / rp.LOCK_REL
    lock.write_text(json.dumps({"pid": os.getpid(), "started_at": "now"}), encoding="utf-8")
    r6 = _run([*RP, "run", "--project", P], cwd=tmp)
    lock.unlink(missing_ok=True)
    record("G2 a second concurrent reproduction of the same paper is refused",
           r6.returncode == 1 and "another reproduce.py run is active" in r6.stdout, _out(r6)[-160:])
    src = D / "big.csv"
    with patch("os.link", side_effect=OSError("no hard links here")):
        try:
            rp.provide(src, tmp / "copy_dst" / "big.csv", rp.Budget(5))
            refused = ""
        except rp.DataCopyRefused as exc:
            refused = str(exc)
    record("G2 copying more input data than data_copy_limit_mb is refused with the legal path",
           "wf data add" in refused and "limit" in refused and not (tmp / "copy_dst/big.csv").exists(),
           refused[:160])
    keep, side = rp.split_paper_changes(["03_analysis/notes.md", "02_data/codebook.md",
                                         "01_protocol/x.md", "x.txt", "03_analysis/code/a.py"])
    record("G2 edits to notes and documentation during a run never invalidate it",
           keep == ["x.txt", "03_analysis/code/a.py"] and len(side) == 3, str(keep))
    raw_paths = {e["path"] for e in dataproof.raw_inventory(P)}
    n = dataproof.evidence_count(P, {"path": "datasource:src/arch.zip", "kind": "zip_member_rows",
                                     "member": "inner/rows.txt", "delimiter": "$", "header": True},
                                 raw_paths)
    nl = dataproof.evidence_count(P, {"path": "datasource:src/arch.zip", "kind": "zip_member_lines",
                                      "member": "inner/rows.txt", "header": False}, raw_paths)
    _w(P, "02_data/acquisition_manifest.json", {"schema_version": 1})
    v = dataproof.validate(P)
    record("G2 acquisition proof inventories declared data in place and counts archive members "
           "by streaming (no extraction)", n == 3 and nl == 4
           and "datasource:src/arch.zip" in raw_paths
           and not any("no acquired payload" in p for p in v.problems), f"{n} {nl} {v.problems[:2]}")
    shutil.copy2(D / "big.csv", P / "02_data/raw/big.csv")
    dup = _gate(P, "data_not_copied", "S04_data")
    (P / "02_data/raw/big.csv").unlink()
    clean = _gate(P, "data_not_copied", "S04_data")
    record("G2 data_not_copied fails on a copy of a declared dataset inside the paper",
           not dup.ok and dup.severity == "fail" and "big.csv" in dup.detail and clean.ok,
           dup.detail[:120] + " | " + clean.detail[:80])
    sys.path.insert(0, str(ROOT / "tools/hooks"))
    import paper_guard
    with _envset(MEDPAPER_GUARD_TEMP=tmp / "systemp"):
        guard = paper_guard.Guard(P)
        cp = guard.decide({"tool_name": "execute_pwsh", "tool_input": {
            "command": f'Copy-Item "{D}\\big.csv" 02_data/raw/', "cwd": str(P)}})
        ex = guard.decide({"tool_name": "execute_pwsh", "tool_input": {
            "command": f'Expand-Archive "{D}\\arch.zip" -DestinationPath 02_data/raw', "cwd": str(P)}})
        rd = guard.decide({"tool_name": "execute_pwsh", "tool_input": {
            "command": f'Get-Content "{D}\\big.csv" -TotalCount 3', "cwd": str(P)}})
    record("G2 the paper guard blocks copying or extracting a declared dataset into the paper "
           "but allows reading it", bool(cp) and bool(ex) and "never copy" in (cp or "") and rd is None,
           f"{cp} | {ex} | {rd}"[:240])
    _w(P, ".wf/config.toml", "[targets]\nlog_cap_mb = 0.001\n")
    for i in range(5):
        _w(P, f".wf/logs/old{i}.log", "z" * 400)
        t = time.time() - 100 + i
        os.utime(P / f".wf/logs/old{i}.log", (t, t))
    target = logs.log_path(P, "reproduce", tmp / "elsewhere" / "repro12.txt")
    total = sum(f.stat().st_size for f in (P / ".wf/logs").rglob("*") if f.is_file())
    record("G2 logs: a requested path elsewhere lands in .wf/logs/ and the folder is pruned "
           "oldest-first to log_cap_mb", target == P / ".wf/logs/repro12.txt"
           and total <= 1100 and not (tmp / "elsewhere").exists(), f"{target} {total}")


# ------------------------------------------------------------------------------ G3
def run_g3(tmp: Path, record, section) -> None:
    from wfcore import declarations, presentation as pr, scifreeze
    from wfcore.checks.text import deviation_hits
    section("1.13.0 G3: rewind rule, multi-person checks only for manual coding, no deviations")
    P = _paper(tmp / "rw")
    _w(P, "03_analysis/results/primary.json", {"ror": 3.2, "lo": 2.1, "hi": 4.8, "n": 159,
                                               "adj_ror": 2.9, "pct_ind": 20.4})
    rel = "07_manuscript/full_manuscript.md"
    base = ("# Results\n\nAmong 159 reports the ROR was 3.2 (95% CI 2.1 to 4.8). The association "
            "persisted in sensitivity analyses.\n\nDuplicates were rare. We found 12 duplicate "
            "reports.\n")
    _w(P, rel, base)
    before = pr.snapshot_facts(P)

    def verdict(text, changed=(rel,)):
        _w(P, rel, text)
        return pr.classify(P, before, list(changed))
    v1 = verdict(base.replace("The association persisted", "After adjustment for cancer type the "
                              "ROR was 2.9, and the association persisted").replace(
        "were rare.", "were rare; indications were recorded in 20.4% of reports."))
    record("G3 a reviewer-driven edit that adds numbers traceable to the frozen results is "
           "presentation; each added number is listed with its source",
           v1["class"] == "presentation" and {a["number"] for a in v1["added_numbers"]} == {"2.9", "20.4"}
           and all("primary.json" in a["source"] for a in v1["added_numbers"])
           and any("2.9" in ln for ln in pr.change_summary(v1)), str(v1)[:300])
    v2 = verdict(base.replace("Among 159 reports", "Among 160 reports"))
    record("G3 a changed reported value is science", v2["class"] == "scientific"
           and any("value changed: 159 -> 160" in r for r in v2["reasons"]), str(v2["reasons"]))
    v3 = verdict(base.replace("were rare.", "were rare (7.7% of 159)."))
    record("G3 a new number that cannot be traced to the results is science",
           v3["class"] == "scientific" and any("7.7" in r and "cannot be traced" in r
                                               for r in v3["reasons"]), str(v3["reasons"]))
    _w(P, rel, base)
    v4 = pr.classify(P, before, [rel, "03_analysis/results/primary.json"])
    record("G3 a changed results JSON is science", v4["class"] == "scientific"
           and any("results changed" in r for r in v4["reasons"]), str(v4["reasons"]))
    v5 = verdict(base.replace(" We found 12 duplicate reports.", ""))
    record("G3 deleting a sentence (and its numbers) is presentation, listed for the user",
           v5["class"] == "presentation" and {d["number"] for d in v5["deleted_numbers"]} == {"12"},
           str(v5)[:200])
    legacy = {k: {kk: vv for kk, vv in f.items() if kk != "seq"} for k, f in before.items()}
    _w(P, rel, base.replace("were rare.", "were rare in 20.4% of reports."))
    v6 = pr.classify(P, legacy, [rel])
    record("G3 a baseline taken before 1.13 (no reading order) still classifies traceable "
           "additions as presentation", v6["class"] == "presentation", str(v6["reasons"]))
    v7 = verdict(base.replace("analyses.", "analyses, as in a cohort of 412 reports [@smith2020]."))
    record("G3 a new number in a cited sentence is traced to the citation (presentation)",
           v7["class"] == "presentation"
           and [a["source"] for a in v7["added_numbers"]] == ["cited reference"], str(v7)[:200])
    # journal copy: removing a repeat of a value that is still stated is not dropped content
    _w(P, "07_manuscript/full_manuscript.md", "# Results\n\nThere were 12 reports. Of the 12 "
                                              "reports, 5 were fatal.\n")
    _w(P, "08_submission/manuscript/full_manuscript.md", "# Results\n\nThere were 12 reports, 5 "
                                                         "of them fatal.\n")
    ok_dup, probs_dup, _ = scifreeze.variant_delta(P)
    _w(P, "08_submission/manuscript/full_manuscript.md", "# Results\n\nThere were 12 reports.\n")
    ok_drop, probs_drop, _ = scifreeze.variant_delta(P)
    record("G3 journal copy: a merged duplicate value is not a drop; a value that disappears is",
           ok_dup and not ok_drop and "5" in " ".join(probs_drop), f"{probs_dup} | {probs_drop}")
    st = json.loads((P / ".wf/state.json").read_text(encoding="utf-8"))
    st.setdefault("decisions", {})["journal_trim_authorised"] = {"value": "YES"}
    (P / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    record("G3 rebind honours a user-authorised journal trim (no hand-edited variant.json)",
           scifreeze._trim_authorised(P))
    # final copy uses the same rule
    from docx import Document
    from wfcore import finalcopy
    a, b = tmp / "orig.docx", tmp / "copy.docx"
    for path, text in ((a, "Among 159 reports the ROR was 3.2. We found 12 duplicate reports."),
                       (b, "Among 159 reports the ROR was 3.2; after adjustment it was 2.9.")):
        d = Document()
        d.add_paragraph(text)
        d.save(path)
    fc_ok = finalcopy.compare(P, a, b)
    d = Document()
    d.add_paragraph("Among 160 reports the ROR was 3.2. We found 12 duplicate reports.")
    d.save(b)
    fc_bad = finalcopy.compare(P, a, b)
    record("G3 final copy: traceable additions and deletions pass; a changed value is refused",
           not fc_ok and any("value changed" in x for x in fc_bad), f"{fc_ok} | {fc_bad}")

    # multi-person checks only for manual coding
    pub = {"study_design": "case_control", "data_source": {"type": "public_database", "name": "FAERS"}}
    meta = {"study_design": "meta_analysis", "data_source": {"type": "published_studies", "name": "x"}}
    chart = {"study_design": "cohort_retrospective",
             "data_source": {"type": "own_centre_retrospective", "name": "H", "manual_abstraction": True}}
    _w(P, "01_protocol/idea.json", pub)
    rules_pub = " ".join(declarations.review_rules(P))
    _w(P, "01_protocol/idea.json", chart)
    rules_chart = " ".join(declarations.review_rules(P))
    toml = (ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8")
    cards = {n: (ROOT / f"pipeline/stages/{n}.md").read_text(encoding="utf-8")
             for n in ("S01_intake", "S03_protocol", "S21_authors")}
    record("G3 automated extraction (public database) asks no coder/double-entry/agreement "
           "question and the reviewer is told not to ask; manual abstraction still does",
           declarations.needs_human_coding(pub) is False and declarations.needs_human_coding(meta) is True
           and declarations.needs_human_coding(chart) is True
           and "Do not request double data entry" in rules_pub
           and "double data entry" not in rules_chart and "human_coders" not in toml
           and all("manual_abstraction" in c for c in cards.values()), rules_pub[:160])

    # protocol deviations never reach the paper
    phrases = ["Departures from it are reported in the Discussion.",
               "This was a protocol deviation.",
               "The analysis deviated from the original protocol.",
               "As originally planned, we used Firth regression.",
               "The protocol was archived but not registered.",
               "本研究存在方案偏离。", "与研究方案不一致的分析"]
    caught = [p for p in phrases if deviation_hits(p)]
    clean = deviation_hits("The analysis followed a logistic model with prespecified covariates.")
    meta_p = _paper(tmp / "meta")
    _w(meta_p, "01_protocol/idea.json", meta)
    _w(meta_p, "01_protocol/protocol_v1.md", "# Ethics\n\nPROSPERO CRD42024000001\n")
    allowed = deviation_hits("Deviations from the registered protocol are listed.", meta_p)
    record("G3 protocol-deviation narration is blocked in manuscript text (EN and ZH); a "
           "registered meta-analysis (PROSPERO) may report it (PRISMA 2020 item 24c)",
           len(caught) == len(phrases) and not clean and not allowed, str(caught))
    _w(P, "07_manuscript/methods.md", "# Methods\n\nThe protocol was written before any "
                                      "contrast was computed; departures from it are reported "
                                      "in the Discussion.\n")
    g = _gate(P, "no_ai_boilerplate", "S08_methods", path="07_manuscript/methods.md")
    record("G3 no_ai_boilerplate blocks plan-history text and simply asks to describe the "
           "analysis performed", not g.ok and any("describe the analysis performed" in h
                                                  for h in g.hints), g.detail[:160])
    s06 = next(s for s in __import__("tomllib").loads(toml)["stage"] if s["id"] == "S06_protocol_final")
    heads = [x for gg in s06["gate"] if gg.get("check") == "md_sections" for x in gg["headings"]]
    from manuscript import writer_bundle
    _w(P, "01_protocol/protocol_final.md", "# Objective\n\nX.\n\n## Deviations from protocol v1\n\n"
                                           "We changed the outcome.\n\n## Design\n\nCase-control.\n")
    src = (ROOT / "tools/manuscript/writer_bundle.py").read_text(encoding="utf-8")
    record("G3 S06 no longer requires a deviations section; the writer bundle drops one; "
           "outputs no longer include protocol_diff.md",
           "Deviations from protocol v1" not in heads
           and "01_protocol/protocol_diff.md" not in s06["outputs"]
           and "Deviations\\b" in src and hasattr(writer_bundle, "build"), str(heads))


# ------------------------------------------------------------------------------ G4
def run_g4(tmp: Path, record, section) -> None:
    import tomllib
    import zipfile
    from xml.etree import ElementTree as ET
    section("1.13.0 G4: each section polished at its stage; Word without keep markers")
    P = _paper(tmp / "pol")
    _w(P, "03_analysis/results/primary.json", {"n": 159, "ror": 3.2})
    _w(P, "07_manuscript/results.md", "# Results\n\nAmong 159 reports the reporting odds ratio "
                                      "was 3.2.\n")
    ok = _gate(P, "section_language_clean", "S09_results", path="07_manuscript/results.md")
    _w(P, "07_manuscript/results.md", "# Results\n\nThe pipeline stage parsed 159 rows; it is "
                                      "worth noting that the reporting odds ratio was 3.2. "
                                      "Departures from the protocol are listed below.\n")
    bad = _gate(P, "section_language_clean", "S09_results", path="07_manuscript/results.md")
    record("G4 section_language_clean passes clinical prose and blocks engineering, AI and "
           "backstage/plan-history language in the section",
           ok.ok and not bad.ok and bad.severity == "fail" and "blocking" in bad.detail,
           ok.detail[:80] + " | " + bad.detail[:200])
    pipe = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
    wired = {s["id"]: [g.get("path") for g in s["gate"] if g.get("check") == "section_language_clean"
                       and g.get("since") == "1.13.0"] for s in pipe["stage"]}
    record("G4 every writing stage polishes its own section (S08 Methods, S09 Results, S14 "
           "Introduction, S16 Discussion, S17 Abstract); S22 only harmonises",
           "07_manuscript/methods.md" in wired["S08_methods"]
           and wired["S09_results"] == ["07_manuscript/results.md"]
           and wired["S14_introduction"] == ["07_manuscript/introduction.md"]
           and wired["S16_discussion"] == ["07_manuscript/discussion.md"]
           and wired["S17_assemble"] == ["07_manuscript/abstract.md"]
           and "cross-section consistency" in (ROOT / "pipeline/stages/S22_polish.md").read_text(
               encoding="utf-8"), str(wired)[:300])
    # Word: build a synthetic manuscript with an embedded figure and a Word table
    if not shutil.which("pandoc"):
        record("G4 Word build needs pandoc", False, "pandoc missing")
        return
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from manuscript import build_docx as bd
    from tables.threeline import write_table
    from wfcore import docxaudit
    W = tmp / "word"
    Q = _paper(W)
    (Q / "05_figures/out").mkdir(parents=True)
    fig = plt.figure(figsize=(6, 7))
    plt.plot([0, 1], [0, 1])
    fig.savefig(Q / "05_figures/out/Figure1.png", dpi=60)
    plt.close(fig)
    with _envset(MEDPAPER_PROJECT=Q):
        write_table(Q / "04_tables/main/Table1.xlsx", "Table 1", "Table 1. Characteristics.",
                    ["Characteristic", "Value"], [["Age, years", "62"], ["Women, n (%)", "40 (25)"]],
                    ["Data are n (%)."])
    _w(Q, "01_protocol/artifact_plan.json", {
        "main_figures": [{"id": "Figure 1", "file": "05_figures/out/Figure1.png"}],
        "supp_figures": [], "main_tables": [{"id": "Table 1", "file": "04_tables/main/Table1.xlsx"}],
        "supp_tables": []})
    _w(Q, "04_tables/table_captions.md", "## Table 1. Characteristics of the reports\n\nData are n (%).\n")
    _w(Q, "08_submission/target_journal.json", {"journal": "J", "embed_figures": True,
                                                "word_tables": True, "limits": {"source": "x"}})
    legend = " ".join(["The figure shows the cumulative reports over time."] * 25)
    _w(Q, "07_manuscript/full_manuscript.md",
       "# Title of the study\n\n# Results\n\nText (Table 1; Figure 1).\n\n# Figure legends\n\n"
       f"## Figure 1. Reports over time\n\n{legend}\n")
    out = W / "manuscript.docx"
    style = bd.load_style_defaults()
    with _envset(MEDPAPER_PROJECT=Q):
        work = W / "render"
        work.mkdir()
        rendered = bd.prepare_source("manuscript", [Q / "07_manuscript/full_manuscript.md"], Q, work, style)
        bd.run_pandoc(rendered, out, None, None)
        bd.normalize_docx(out, style, "manuscript")
    with zipfile.ZipFile(out) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
        styles = ET.fromstring(zf.read("word/styles.xml"))
        raw = zf.read("word/document.xml") + zf.read("word/styles.xml")
    probs = docxaudit.keep_problems(root, styles)
    tbl = docxaudit.table_layout_problems(root)
    ext = root.find(f".//{{{docxaudit.WP}}}extent")
    W_ = docxaudit.W
    # paragraphs and tables only (Pandoc puts heading bookmarks between them)
    kids = [el for el in root.find(f"{{{W_}}}body") if el.tag in (f"{{{W_}}}p", f"{{{W_}}}tbl")]

    def _brk(el):
        return el.find(f".//{{{W_}}}br[@{{{W_}}}type='page']") is not None

    def _sect(el):
        return el.find(f"{{{W_}}}pPr/{{{W_}}}sectPr") is not None
    # a page break right after a next-page section break is a blank page (Word COM check)
    blank = any(_sect(a) and _brk(b) for a, b in zip(kids, kids[1:]))
    fig_i = next((i for i, el in enumerate(kids) if el.find(f".//{{{docxaudit.WP}}}extent")
                  is not None), 0)
    j = fig_i - 1
    while j >= 0 and not _sect(kids[j]) and not _brk(kids[j]) and "".join(kids[j].itertext()).strip():
        j -= 1
    new_page = j >= 0 and (_sect(kids[j]) or _brk(kids[j]))
    record("G4 the built DOCX has zero keepNext/keepLines/pageBreakBefore (document and styles), "
           "cantSplit rows with a repeated header, the table on its own section, and the figure "
           "on a new page without a blank page, sized to fit with its long legend",
           not probs and not tbl and b"keepNext" not in raw and b"keepLines" not in raw
           and b"pageBreakBefore" not in raw and b"cantSplit" in raw and b"tblHeader" in raw
           and new_page and not blank and ext is not None
           and int(ext.get("cy")) / 914400 < 9.0, f"{probs} | {tbl} | blank={blank} new={new_page}")
    import copy
    marked = copy.deepcopy(root)
    for p in marked.iter(f"{{{docxaudit.W}}}p"):
        ppr = p.find(f"{{{docxaudit.W}}}pPr")
        if ppr is None:
            ppr = ET.SubElement(p, f"{{{docxaudit.W}}}pPr")
        ET.SubElement(ppr, f"{{{docxaudit.W}}}keepNext")
        break
    record("G4 the audit (docx_bundle_ready keep_with_next) reports a keep marker as a black square",
           any("black squares" in x for x in docxaudit.keep_problems(marked, styles)))
    # A long legend never turns the figure into a thumbnail: it keeps >= the minimum page share
    # (or its natural size) and the legend continues on the next page.
    import struct
    import tempfile
    import zlib
    bdx = bd

    def _png(path: Path, w: int, h: int) -> None:
        raw = b"".join(b"\x00" + b"\xff" * w for _ in range(h))
        def chunk(t: bytes, d: bytes) -> bytes:
            return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))
        path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 0, 0, 0, 0))
                         + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    with tempfile.TemporaryDirectory(prefix="medpaper_selftest_") as tdir:
        sq, wide = Path(tdir) / "sq.png", Path(tdir) / "wide.png"
        _png(sq, 400, 400)
        _png(wide, 400, 200)
        style = {"body_font_pt": 12, "line_spacing": 2}
        long_legend = "Figure 1. " + " ".join(["word"] * 300)
        _pw, _ph, _pm = bdx._page_geometry(style)
        usable_h = (_ph - 2 * _pm) / 1440
        w_long, h_long = bdx.figure_size(sq, long_legend, style)
        w_short, h_short = bdx.figure_size(wide, "Figure 2. A short legend.", style)
        floor = usable_h * bdx.FIGURE_MIN_PAGE_FRACTION
    record("G4 a 300-word legend keeps the figure readable (>= minimum page share), never a "
           "thumbnail; a short legend leaves the figure at its natural full width",
           h_long >= floor - 0.01 and w_long > 3.0 and w_short >= 5.5,
           f"long {w_long}x{h_long} (floor {floor:.2f}); short {w_short}x{h_short}")


def _gate(p: Path, name: str, stage: str, **spec):
    import selftest_v180 as v180
    return v180._gate(p, name, stage, **spec)


# ------------------------------------------------------------------------------ G5
def run_g5(tmp: Path, record, section) -> None:
    sys.path.insert(0, str(ROOT / "tools/hooks"))
    import paper_guard
    from wfcore.state import State
    section("1.13.0 G5: engine state not hand-edited, coined labels, numeral lint, S24 rule, migration")
    P = _paper(tmp / "paper")
    cases = [
        ("fs_write .wf/state.json", {"tool_name": "fs_write", "tool_input": {"path": ".wf/state.json", "text": "{}"}}, True),
        ("str_replace .wf/review_rounds.json", {"tool_name": "str_replace", "tool_input": {"path": str(P / ".wf/review_rounds.json"), "oldStr": "a", "newStr": "b"}}, True),
        ("shell Set-Content .wf/progress.md", {"tool_name": "execute_pwsh", "tool_input": {"command": "Set-Content .wf/progress.md x", "cwd": str(P)}}, True),
        ("python one-liner writes a revision round", {"tool_name": "execute_pwsh", "tool_input": {"command": "python -c \"open('.wf/revisions/R003.json','w').write('{}')\"", "cwd": str(P)}}, True),
        ("read .wf/state.json", {"tool_name": "read_file", "tool_input": {"path": ".wf/state.json"}}, False),
        ("reviewer writes its report.md", {"tool_name": "fs_write", "tool_input": {"path": ".wf/review_bundle/M001/report.md", "text": "x"}}, False),
        ("paper allowlist edit", {"tool_name": "fs_write", "tool_input": {"path": ".wf/language_allowlist.json", "text": "{}"}}, False),
        ("engine command redirects to .wf/logs", {"tool_name": "execute_pwsh", "tool_input": {"command": "python tools/wf.py check *> .wf/logs/c.txt", "cwd": str(P)}}, False),
    ]
    guard = paper_guard.Guard(P)
    wrong = [f"{lab}: {got or 'allowed'}"[:120] for lab, pay, blk in cases
             if bool(got := guard.decide(pay)) != blk]
    record("G5 paper guard blocks hand edits of engine-owned .wf state (state, progress, rounds, "
           "revisions) and keeps reads, report.md, allowlist and logs writable", not wrong,
           "; ".join(wrong))

    _w(P, "01_protocol/display_dictionary.json", {"terms": [
        {"internal": "rule2", "aliases": ["second screen"], "label": "Confirmed bleeding",
         "definition": "bleeding confirmed in the case narrative"}]})
    _w(P, "07_manuscript/results.md", "# Results\n\nUnder the \u201coriginal rule\u201d 12 cases "
       "remained; the \"second screen\" and \"Confirmed bleeding\" groups are named.\n")
    res = _gate(P, "engineering_language_clean", "S09_results", paths=["07_manuscript/results.md"])
    record("G5 engineering_language_clean warns on a quoted coined label absent from the display "
           "dictionary (dictionary labels and aliases pass)",
           not res.ok and res.severity == "warn" and "original rule" in res.detail
           and "second screen" not in res.detail and "Confirmed bleeding" not in res.detail,
           res.detail[:200])

    import importlib.util
    spec = importlib.util.spec_from_file_location("pol_g5", ROOT / "tools/text/polish.py")
    pol = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(pol)
    ok_txt = "The protocol was approved (No. 2026J014). See Fig. 2 for details; vs. 3 controls.\n"
    bad_txt = "The protocol was approved. 2026 patients were screened.\n"
    kinds = lambda t: {x["kind"] for x in pol.lint_style("a.md", t)}  # noqa: E731
    record("G5 numeral lint: 'No. 2026J', 'Fig. 2', 'vs. 3' are not sentence-initial numerals; a "
           "real one still is", "sentence_starts_with_numeral" not in kinds(ok_txt)
           and "sentence_starts_with_numeral" in kinds(bad_txt))

    card = (ROOT / "pipeline/stages/S24_package_human_review.md").read_text(encoding="utf-8")
    src = (ROOT / "tools/package_content.py").read_text(encoding="utf-8")
    record("G5 S24 card and package_content.py agree: classify first, recapture only in a round "
           "that reopened S23", "rework.py check" in card and "capture --replace" in card
           and "revision.permits(project, \"S23_package\")" in src and len(card) <= 4000,
           f"{len(card)} chars")

    q = tmp / "old"
    q.mkdir(parents=True, exist_ok=True)
    st = State(q, ".wf").create("medpaper", "1.12.0", "S10_tables")
    st.migrate_pipeline("1.13.0", ["S10_tables"])
    st.migrate_pipeline("1.13.0", ["S10_tables"])
    notes = [e for e in st.data["events"] if e["kind"] == "requirements_changed"
             and str(e.get("detail", "")).startswith("1.13.0")]
    record("G5 a 1.12 paper opened by 1.13 records exactly one 1.13.0 migration event (paper "
           "files untouched)", len(notes) == 1 and st.data.get("v1130_migration_noted") is True,
           str(len(notes)))


def run(tmp: Path, record, section) -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    run_g1(tmp / "g1", record, section)
    run_g2(tmp / "g2", record, section)
    run_g3(tmp / "g3", record, section)
    run_g4(tmp / "g4", record, section)
    run_g5(tmp / "g5", record, section)

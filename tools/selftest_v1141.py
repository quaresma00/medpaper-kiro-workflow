"""1.14.1 regression tests: the user's three complaints (P1 workspace = the opened folder, P2 one
guard hook on writes and folder searches, P3 no black squares in Word) and the bugs from a full
S01-S25 run on 1.14.0 (B1-B13). Synthetic fixtures only - no real paper or dataset is read.

Called from tools/selftest.py.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from selftest_v1140 import _env, _envset, _out, _paper, _run, _w  # noqa: E402


def run(tmp: Path, record, section, only: str = "") -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    for name, fn in GROUPS:
        if not only or name in only:
            fn(tmp / name, record, section)


def _wf(cwd: Path, *args, **extra) -> subprocess.CompletedProcess:
    return _run(None, "tools/wf.py", *args, env=_env(None, **extra), cwd=cwd)


def _docx(path: Path, paras: list[tuple[str, str]], extra: dict[str, str] | None = None) -> Path:
    """A minimal DOCX: (style id, text) paragraphs; `extra` adds zip parts (docProps...)."""
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = "".join(f'<w:p><w:pPr><w:pStyle w:val="{s}"/></w:pPr><w:r><w:t>{t}</w:t></w:r></w:p>'
                   for s, t in paras)
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.'
                    'openxmlformats.org/package/2006/content-types"/>')
        zf.writestr("word/document.xml", f'<w:document xmlns:w="{w}"><w:body>{body}</w:body>'
                    "</w:document>")
        for name, text in (extra or {}).items():
            zf.writestr(name, text)
    return path


# ------------------------------------------------------------------------------ P1
def run_p1(tmp: Path, record, section) -> None:
    from wfcore import paths, registry
    from wfcore.state import State
    section("1.14.1 P1: the folder the user opened is the paper")
    fresh = tmp / "fresh"
    _w(fresh, "my idea.md", "Idea: drug A and portal vein thrombosis.\n")
    p = _wf(fresh, "init")
    users = json.loads((fresh / ".wf/user_files.json").read_text(encoding="utf-8")) \
        if (fresh / ".wf/user_files.json").is_file() else {}
    subdirs = sorted(x.name for x in fresh.iterdir() if x.is_dir())
    record("P1 wf init makes the opened folder the paper in place: user notes stay where they "
           "are, no paper subfolder, no 'reopen another folder' advice",
           p.returncode == 0 and (fresh / ".wf/state.json").is_file()
           and (fresh / "my idea.md").is_file() and "my idea.md" in users.get("files", [])
           and not any(n.startswith(("oxaliplatin", "paper")) for n in subdirs)
           and "open this folder" not in _out(p), _out(p)[-300:] + str(subdirs))

    older = tmp / "older"
    _w(older, "notes.md", "my notes\n")
    old_paper = older / "first-attempt"
    State(old_paper, ".wf").create("medpaper", "1.13.0", "S05_analysis")
    _w(old_paper, "07_manuscript/draft.md", "old draft\n")
    before = (old_paper / ".wf/state.json").read_bytes()
    p = _wf(older, "init")
    users = json.loads((older / ".wf/user_files.json").read_text(encoding="utf-8"))
    from wfcore.checks import Ctx, get, load_all
    load_all()
    pipe = registry.load()
    orphans = get("no_orphans")(Ctx(pipeline=pipe, state=State(older, ".wf").load(), project=older,
                                    stage=pipe.stage("S01_intake"), spec={"check": "no_orphans"}))
    record("P1 an older paper attempt in a subfolder does not make the folder unsuitable: init "
           "succeeds, mentions it once, leaves it untouched and outside the paper's own files",
           p.returncode == 0 and "older paper attempt in first-attempt/" in _out(p)
           and (old_paper / ".wf/state.json").read_bytes() == before
           and not any(f.startswith("first-attempt/") for f in users.get("files", []))
           and paths.foreign_prefixes(older)[1:] == ("first-attempt/",)
           and orphans.ok and paths.unsuitable_paper_dir(older) is None,
           _out(p)[-300:] + orphans.detail)
    sys.path.insert(0, str(ROOT / "tools/hooks"))
    import paper_guard
    g = paper_guard.Guard(older)
    read_old = g.decide({"tool_name": "list_directory", "tool_input": {"path": str(old_paper)}})
    write_old = g.decide({"tool_name": "fs_write", "tool_input": {
        "path": str(old_paper / "07_manuscript/draft.md"), "text": "x"}})
    wide = g.decide({"tool_name": "grep_search", "tool_input": {"query": "draft"}})
    scoped = g.decide({"tool_name": "grep_search", "tool_input": {
        "query": "draft", "includePattern": "07_manuscript/**"}})
    excluded = g.decide({"tool_name": "file_search", "tool_input": {
        "query": "x", "excludePattern": "first-attempt/**"}})
    own = g.decide({"tool_name": "fs_write", "tool_input": {
        "path": str(older / "07_manuscript/x.md"), "text": "x"}})
    record("P1 the older attempt stays isolated: listing/writing it and a whole-workspace search "
           "are blocked; a search scoped to this paper (include or exclude) and the paper's own "
           "writes are allowed",
           bool(read_old) and bool(write_old) and bool(wide) and "first-attempt" in str(wide)
           and not scoped and not excluded and not own, str([read_old, write_old, wide])[:400])

    again = _wf(fresh, "init")
    force = _wf(fresh, "init", "--force")
    st_before = (fresh / ".wf/state.json").read_bytes()
    _w(fresh, "03_analysis/notes.md", "old notes\n")
    yes = _wf(fresh, "init", "--force", "--yes")
    arch = sorted((fresh / paths.ARCHIVE_DIR).glob("*")) if (fresh / paths.ARCHIVE_DIR).is_dir() else []
    record("P1 a folder that is already a paper is resumed (nothing reset); a fresh restart "
           "needs an explicit --yes and archives the old run in place (.wf_archive/<UTC>/, not "
           "a paper), the user's notes stay",
           again.returncode == 0 and "already a paper" in _out(again) and "resum" in _out(again)
           and force.returncode == 2 and "Ask the user once" in _out(force)
           and yes.returncode == 0 and len(arch) == 1 and (arch[0] / "wf_state/state.json").is_file()
           and (arch[0] / "03_analysis/notes.md").is_file()
           and not (fresh / "03_analysis/notes.md").is_file()
           and (fresh / "my idea.md").is_file() and (fresh / ".wf/state.json").is_file()
           and (fresh / ".wf/state.json").read_bytes() != st_before
           and not paths.nested_papers(fresh),
           _out(again)[-200:] + _out(force)[-200:] + _out(yes)[-200:])

    central_root = tmp / "central_root"
    c = _wf(tmp, "new", "my-paper", "--central", MEDPAPER_PAPERS_ROOT=central_root)
    card = (ROOT / "pipeline/stages/S01_intake.md").read_text(encoding="utf-8")
    skill = (ROOT / ".kiro/skills/medpaper-pipeline/SKILL.md").read_text(encoding="utf-8")
    from wfcore import adapters
    steer = dict(adapters.paper_adapters(ROOT, fresh))[".kiro/steering/medpaper-pipeline.md"]
    record("P1 an explicit central root still creates <root>/<slug> (only on the user's "
           "request); S01 card, SKILL, steering and the refusal text say: never a subfolder, "
           "never 'reopen another folder'",
           c.returncode == 0 and (central_root / "my-paper/.wf/state.json").is_file()
           and "asked for this location explicitly" in _out(c)
           and "Never create a" in card and "reopen another folder" in card
           and len(card) <= registry.CARD_MAX_CHARS and "never ask the" in skill
           and "reopen" in steer and "Never create a subfolder" in paths.ASK_ONCE,
           _out(c)[-300:])


# ------------------------------------------------------------------------------ P2
def run_p2(tmp: Path, record, section) -> None:
    from wfcore import adapters
    section("1.14.1 P2: one guard hook, on writes and folder searches only")
    ws = _paper(tmp / "ws")
    files = dict(adapters.paper_adapters(ROOT, ws))
    hooks = {k: json.loads(v) for k, v in files.items() if k.startswith(".kiro/hooks/")}
    pre = [(k, h["hooks"][0]) for k, h in hooks.items() if h["hooks"][0]["trigger"] == "PreToolUse"
           and h["hooks"][0].get("matcher") != "disclose_context"]
    tokens = set(pre[0][1]["matcher"].split("|")) if len(pre) == 1 else set()
    record("P2 a paper gets ONE merged PreToolUse guard (no separate frozen-guard hook); its "
           "matcher is the write/delete/rename tools and the discovery tools only - never "
           "read_file, read_code, execute_pwsh or control_pwsh_process",
           len(pre) == 1 and set(adapters.GUARD_TOOLS) <= tokens
           and not tokens & {"read_file", "read_code", "execute_pwsh", "control_pwsh_process",
                             "readFile", "executePwsh"}
           and not any(".*" in t for t in tokens)
           and ".kiro/hooks/medpaper-frozen-guard.json" not in files
           and ".kiro/hooks/medpaper-frozen-guard.json" not in dict(adapters.engine_adapters()),
           str(sorted(tokens)))

    home = tmp / "home"
    reports = tmp / "reports"
    reports.mkdir(parents=True)
    other = _paper(tmp / "other")
    cfg = _wf(tmp, "config", "--user", "report_dirs", "add", str(reports), MEDPAPER_HOME=home)
    bad = [_wf(tmp, "config", "--user", "report_dirs", "add", str(d), MEDPAPER_HOME=home)
           for d in (ROOT, Path(ROOT.anchor), other)]
    text = (home / "config.toml").read_text(encoding="utf-8") if (home / "config.toml").is_file() else ""
    record("P2 `wf config --user report_dirs add <dir>` writes the user config; the engine, a "
           "drive root and a paper are refused",
           cfg.returncode == 0 and "report_dirs" in text and str(reports).replace("\\", "\\\\") in text
           and all(b.returncode == 2 for b in bad) and str(ROOT) not in text,
           _out(cfg)[-200:] + " | " + " | ".join(_out(b)[-80:] for b in bad))
    sys.path.insert(0, str(ROOT / "tools/hooks"))
    import paper_guard
    nested_report = _paper(reports / "a-paper")
    # the selftest itself runs inside the system temp; point the guard's temp rule elsewhere
    with _envset(MEDPAPER_HOME=home, MEDPAPER_GUARD_TEMP=tmp / "no_such_temp"):
        g = paper_guard.Guard(ws)

        def d(tool, path, **kw):
            return g.decide({"tool_name": tool, "tool_input": {"path": str(path), **kw}})
        cases = {
            "own write": (d("fs_write", ws / "07_manuscript/a.md", text="x"), False),
            "own str_replace": (d("str_replace", ws / "07_manuscript/a.md", oldStr="a", newStr="b"), False),
            "engine write": (d("fs_write", ROOT / "tools/x.py", text="x"), True),
            "engine delete": (d("delete_file", ROOT / "README.md"), True),
            "other paper write": (d("fs_write", other / "07_manuscript/a.md", text="x"), True),
            "other paper list": (d("list_directory", other), True),
            "other paper search": (g.decide({"tool_name": "file_search", "tool_input": {
                "query": "x", "includePattern": str(other) + "/**"}}), True),
            "report .md": (d("fs_write", reports / "retro.md", text="x"), False),
            "report .txt": (d("fs_write", reports / "notes.txt", text="x"), False),
            "report .py": (d("fs_write", reports / "x.py", text="x"), True),
            "paper inside report dir": (d("fs_write", nested_report / "x.md", text="x"), True),
            "outside write": (d("fs_write", tmp / "elsewhere/x.md", text="x"), True),
        }
    record("P2 synthetic hook payloads: own writes and .md/.txt into a report folder pass; "
           "engine writes/deletes, another paper (write, list, search), other file types in the "
           "report folder, a paper inside it and other outside writes are blocked",
           all(bool(v) == want for v, want in cases.values()),
           str({k: bool(v) for k, (v, _) in cases.items()}))

    # merged frozen-source guard: a write to a round-certified file is blocked in the same hook
    from wfcore import revision
    _w(ws, "05_figures/code/fig1.R", "# fig1\n")
    rec = revision.receipt(ws, "05_figures/code/fig1.R")
    _w(ws, ".wf/revisions/R001.json", {"round_id": "R001", "status": "complete", "schema_version": 1,
                                       "items": [{"id": "I1", "status": "done", "validated_by": ["x"],
                                                  "changed_files": [rec]}]})
    payload = json.dumps({"tool_name": "fs_write", "tool_input": {
        "path": str(ws / "05_figures/code/fig1.R"), "text": "# edited\n"}})
    hook = subprocess.run([sys.executable, "-B", str(ROOT / "tools/hooks/paper_guard.py"),
                           "--paper", str(ws)], input=payload, capture_output=True, text=True,
                          env=_env(None, MEDPAPER_HOME=home), timeout=60)
    read = subprocess.run([sys.executable, "-B", str(ROOT / "tools/hooks/paper_guard.py"),
                           "--paper", str(ws)], input=json.dumps({"tool_name": "fs_write",
                                                                  "tool_input": {"path": str(ws / "07_manuscript/b.md"), "text": "x"}}),
                          capture_output=True, text=True, env=_env(None, MEDPAPER_HOME=home), timeout=60)
    record("P2 the merged hook still refuses a write to a round-certified source outside an open "
           "round (the 1.8.0 frozen guard) and lets an ordinary paper write through",
           hook.returncode == 2 and "certified" in hook.stderr and read.returncode == 0,
           hook.stderr[-300:] + read.stderr[-200:])

    # adapter auto-refresh
    q = tmp / "refresh"
    q.mkdir()
    init = _wf(q, "init")
    guard_file = q / ".kiro/hooks/medpaper-paper-guard.json"
    old = json.loads(guard_file.read_text(encoding="utf-8"))
    old["hooks"][0]["matcher"] = "read_file|execute_pwsh|fs_write"
    guard_file.write_text(json.dumps(old), encoding="utf-8")
    _w(q, ".kiro/hooks/medpaper-frozen-guard.json", {"version": "v1", "hooks": [
        {"name": "medpaper: stop edits", "trigger": "PreToolUse", "action": {
            "type": "command", "command": f'python "{ROOT.as_posix()}/tools/hooks/frozen_guard.py"'}}]})
    s1 = _wf(q, "status", "--brief")
    s2 = _wf(q, "status", "--brief")
    events = [e for e in json.loads((q / ".wf/state.json").read_text(encoding="utf-8"))["events"]
              if e.get("kind") == "adapters_refreshed"]
    foreign = tmp / "foreign_engine"
    foreign.mkdir()
    _wf(foreign, "init")
    for f in (foreign / ".kiro").rglob("*"):
        if f.is_file():
            f.write_text(f.read_text(encoding="utf-8").replace(ROOT.as_posix(), "C:/elsewhere/mp")
                         .replace("read_file|", ""), encoding="utf-8")
    fg = foreign / ".kiro/hooks/medpaper-paper-guard.json"
    alien = fg.read_text(encoding="utf-8")
    _wf(foreign, "status", "--brief")
    record("P2 every wf command refreshes stale paper adapters (one line, one event; the old "
           "frozen-guard hook file is removed); current adapters are left alone; adapters that "
           "call another engine checkout are never re-pointed",
           init.returncode == 0 and "adapters refreshed" in s1.stderr
           and "adapters refreshed" not in s2.stderr and len(events) == 1
           and not (q / ".kiro/hooks/medpaper-frozen-guard.json").exists()
           and guard_file.read_text(encoding="utf-8") == dict(adapters.paper_adapters(ROOT, q.resolve()))[
               ".kiro/hooks/medpaper-paper-guard.json"]
           and fg.read_text(encoding="utf-8") == alien, s1.stderr[-300:])


# ------------------------------------------------------------------------------ P3
def run_p3(tmp: Path, record, section) -> None:
    from xml.etree import ElementTree as ET
    from docx import Document
    from manuscript import build_docx as bd
    from wfcore import docxaudit
    from wfcore.checks.manuscript import _footnote_keep_allowed
    section("1.14.1 P3: no keep markers (no black squares) unless the journal opts in")
    tmp.mkdir(parents=True, exist_ok=True)

    def make(name: str) -> Path:
        doc = Document()
        doc.add_paragraph("Table 1. Baseline characteristics")
        t = doc.add_table(rows=3, cols=2)
        for i, row in enumerate(t.rows):
            row.cells[0].text, row.cells[1].text = f"r{i}", "1"
        doc.add_paragraph("Data are n (%).")
        out = tmp / name
        doc.save(out)
        return out

    def markers(path: Path) -> tuple[int, list[str], list[str]]:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
            styles = ET.fromstring(zf.read("word/styles.xml"))
        root = ET.fromstring(xml)
        return (xml.count(b"<w:keepNext") + xml.count(b"<w:keepLines")
                + xml.count(b"<w:pageBreakBefore"),
                docxaudit.keep_problems(root, styles), docxaudit.keep_problems(root, styles, True))
    default, opt = make("default.docx"), make("opt.docx")
    bd._sanitize_package_xml(default, "Arial")
    bd._sanitize_package_xml(opt, "Arial", True)
    n0, prob0, _ = markers(default)
    n1, prob1, prob1_on = markers(opt)
    with zipfile.ZipFile(default) as zf:
        styles_xml = zf.read("word/styles.xml")
    record("P3 the default build has zero keepNext/keepLines/pageBreakBefore in document and "
           "styles (the footnote is a normal paragraph after the table); keep_table_footnotes "
           "puts keepNext on the last row only, and only the switch makes the audit accept it",
           n0 == 0 and b"keepNext" not in styles_xml and not any("black squares" in x for x in prob0)
           and n1 > 0 and any("black squares" in x for x in prob1)
           and not any("black squares" in x for x in prob1_on), f"{n0} {n1} {prob0} {prob1}")
    p = _paper(tmp / "spec")
    _w(p, "08_submission/target_journal.json", {"journal": "J", "keep_table_footnotes": True})
    spec_on = docxaudit.audit_spec(p)
    _w(p, "08_submission/target_journal.json", {"journal": "J"})
    spec_off = docxaudit.audit_spec(p)
    pinned = _footnote_keep_allowed(SimpleNamespace(state=SimpleNamespace(engine_pin="1.14.0")), spec_off)
    current = _footnote_keep_allowed(SimpleNamespace(state=SimpleNamespace(engine_pin="1.14.1")), spec_off)
    switched = _footnote_keep_allowed(SimpleNamespace(state=SimpleNamespace(engine_pin="1.14.1")), spec_on)
    src = (ROOT / "tools/manuscript/build_docx.py").read_text(encoding="utf-8")
    record("P3 target_journal.json keep_table_footnotes (default false) drives build and audit; "
           "a paper pinned to 1.14.0 keeps the D3 exception until wf migrate; Word 2013 layout "
           "mode stays",
           spec_on["keep_table_footnotes"] is True and spec_off["keep_table_footnotes"] is False
           and pinned and not current and switched and "modern_layout(blob)" in src,
           str([pinned, current, switched]))


# ------------------------------------------------------------------------------ B
def run_b(tmp: Path, record, section) -> None:
    from wfcore import (determinism, docxaudit, journalspec, meaning, nonfacts, packagefreeze,
                        presentation, readerreview, revision)
    from wfcore.checks import numbers, text as textchecks
    section("1.14.1 B1-B13: bugs from a full S01-S25 run on 1.14.0")
    tmp.mkdir(parents=True, exist_ok=True)
    # B1 --------------------------------------------------------------------------------
    report = ("P01: meaning unchanged\nP02: meaning changed - the secondary outcome was dropped\n"
              "- **P03** - **meaning unchanged**\n| P04 | meaning unchanged | wording only |\n"
              "P05: meaning unchanged.\nP06 — meaning changed: hedging removed\n")
    got = meaning.parse_report(report)
    record("B1 the meaning report is parsed line by line: a verdict without a reason no longer "
           "swallows the next line; mixed list/table/dash formats all count",
           sorted(got) == ["P01", "P02", "P03", "P04", "P05", "P06"]
           and got["P01"] == ("unchanged", "") and got["P02"][0] == "changed"
           and "secondary outcome" in got["P02"][1] and got["P04"] == ("unchanged", "wording only")
           and got["P06"][0] == "changed", str(got))
    # B2 --------------------------------------------------------------------------------
    import selftest_v180 as v180
    p = v180._rw_paper(tmp / "b2", "S24_package_human_review")
    _w(p, "08_submission/bundle/manuscript.docx", "v1")
    h1 = revision.receipt(p, "08_submission/bundle/manuscript.docx")
    _w(p, "08_submission/bundle/manuscript.docx", "v2")
    h2 = revision.receipt(p, "08_submission/bundle/manuscript.docx")
    ok_val = {"ok": True, "input_signature": "s", "results": [{"check": "x", "ok": True}]}

    def rnd(rid, items, **extra):
        _w(p, f".wf/revisions/{rid}.json", {"round_id": rid, "status": "complete",
                                             "schema_version": 2, "items": items, **extra})
    item = lambda rel_rec: {"id": "I1", "status": "done", "validated_by": ["x"],  # noqa: E731
                            "changed_files": [rel_rec]}
    _w(p, "07_manuscript/x.md", "x\n")
    xr = revision.receipt(p, "07_manuscript/x.md")
    rnd("R001", [item(h1)], validation=ok_val)
    rnd("R002", [item(xr)], validation={**ok_val, "artifact_hashes": {
        "08_submission/bundle/manuscript.docx": h2["sha256"]}})
    rnd("R003", [item(xr)], validation=ok_val)
    gate = v180._gate(p, "revision_rounds_closed", "S24_package_human_review")
    drift = revision.certified_drift(p)
    record("B2 per file the latest receipt across ALL closed rounds counts: R002 re-certified the "
           "Word file after R001, so the newest round R003 not listing it is no drift (gate and "
           "run-script drift check)", gate.ok and "08_submission/bundle/manuscript.docx" not in drift,
           gate.detail[:300] + str(drift))
    # an open round: the builder's output is recorded into it automatically
    _w(p, ".wf/revisions/R004.json", {"round_id": "R004", "status": "active", "schema_version": 2,
                                       "items": [item(xr)], "allowed_files": []})
    st = json.loads((p / ".wf/state.json").read_text(encoding="utf-8"))
    st["active_revision_round"] = "R004"
    (p / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    _w(p, "08_submission/bundle/manuscript.docx", "v3")
    _w(p, "08_submission/build_receipts/manuscript.json", "{}")
    done = revision.record_derived(p, [p / "08_submission/bundle/manuscript.docx",
                                       p / "08_submission/build_receipts/manuscript.json"])
    data = json.loads((p / ".wf/revisions/R004.json").read_text(encoding="utf-8"))
    data.update(status="complete", validation=ok_val)
    (p / ".wf/revisions/R004.json").write_text(json.dumps(data), encoding="utf-8")
    st["active_revision_round"] = None
    (p / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    gate2 = v180._gate(p, "revision_rounds_closed", "S24_package_human_review")
    _w(p, "08_submission/bundle/manuscript.docx", "hand edit")
    _w(p, "07_manuscript/x.md", "edited\n")
    gate3 = v180._gate(p, "revision_rounds_closed", "S24_package_human_review")
    srcs = {rel: (ROOT / rel).read_text(encoding="utf-8") for rel in (
        "tools/manuscript/build_docx.py", "tools/text/polish.py", "tools/wfcore/submissionqc.py",
        "tools/package_content.py")}
    record("B2 derived outputs written during an open round are recorded into it (no "
           "bookkeeping round); a later hand edit is drift and every drifted file is listed at "
           "once; build_docx, polish, submission_qc and package_content capture record theirs",
           len(done) == 2 and "derived_outputs" in data and gate2.ok and not gate3.ok
           and "manuscript.docx" in gate3.detail and "x.md" in gate3.detail
           and all("record_derived" in s for s in srcs.values()), gate2.detail[:200] + gate3.detail[:300])
    # B3 --------------------------------------------------------------------------------
    from manuscript import build_docx as bd
    q = _paper(tmp / "b3")
    _w(q, "08_submission/target_journal.json", {"journal": "J"})
    _w(q, "00_input/author_info.json", {"authors": [{"name": "A B"}], "corresponding": {
        "name": "A B", "email": "a@b.org", "address": "Dept of Hepatology, City Hospital"}})
    page_addr = bd._info_page(q)
    _w(q, "00_input/author_info.json", {"authors": [{"name": "A B"}], "corresponding": {
        "name": "A B", "email": "a@b.org", "affiliation": "Dept of Oncology"}})
    page_aff = bd._info_page(q)
    card = (ROOT / "pipeline/stages/S21_authors.md").read_text(encoding="utf-8")
    record("B3 the supplementary information page shows the corresponding author's affiliation "
           "from 'affiliation' or (older files) 'address'; S21 writes 'affiliation'",
           "Dept of Hepatology" in page_addr and "Dept of Oncology" in page_aff
           and "affiliation}" in card, page_addr)
    # B4 --------------------------------------------------------------------------------
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, TextStringObject
    pdf = tmp / "esm.pdf"
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w._root_object[NameObject("/Lang")] = TextStringObject("zh")
    with pdf.open("wb") as fh:
        w.write(fh)
    before = determinism.pdf_lang(pdf)
    ok = determinism.set_pdf_lang(pdf, docxaudit.pdf_lang({}))
    determinism.normalize(pdf)
    record("B4 an engine-exported PDF gets /Lang from target_journal.json language (default "
           "en-US) after export; the bundle audit checks it (since 1.14.1)",
           before == "zh" and ok and determinism.pdf_lang(pdf) == "en-US"
           and docxaudit.pdf_lang({"language": "en-GB"}) == "en-GB"
           and docxaudit.pdf_lang({"language": "English"}) == "en-US"
           and "set_pdf_lang(pdf_path" in (ROOT / "tools/manuscript/build_docx.py").read_text(encoding="utf-8")
           and "PDF language /Lang" in (ROOT / "tools/wfcore/checks/manuscript.py").read_text(encoding="utf-8"),
           f"{before} {determinism.pdf_lang(pdf)}")
    # B5 --------------------------------------------------------------------------------
    dx = _docx(tmp / "props.docx", [("Normal", "Text")], {
        "docProps/custom.xml": '<?xml version="1.0"?><Properties xmlns="http://schemas.openxml'
                               'formats.org/officeDocument/2006/custom-properties" xmlns:vt="x">'
                               '<property name="csl"><vt:lpwstr>G:\\temp\\medpaper_render_1\\'
                               'patched_journal.csl</vt:lpwstr></property><property name="bibli'
                               'ography"><vt:lpwstr>06_refs\\refs.bib</vt:lpwstr></property>'
                               '</Properties>',
        "docProps/core.xml": '<?xml version="1.0"?><cp:coreProperties xmlns:cp="c" xmlns:dc="d">'
                             '<dc:title>Drug A</dc:title><dc:description>F:\\work\\x.md'
                             '</dc:description></cp:coreProperties>'})
    hits0 = docxaudit.docprops_problems(dx)
    c0 = determinism.content_sha256(dx)
    determinism.scrub_properties(dx)
    hits1 = docxaudit.docprops_problems(dx)
    with zipfile.ZipFile(dx) as zf:
        core = zf.read("docProps/core.xml").decode()
    record("B5 Word builds drop custom properties (csl/bibliography paths) and local paths in "
           "core/app; the upload check now scans docProps/* and customXml/*; the visible-content "
           "hash is unchanged",
           hits0 and not hits1 and "Drug A" in core and "F:\\" not in core
           and determinism.content_sha256(dx) == c0, str(hits0))
    # B6 --------------------------------------------------------------------------------
    txt = ("Analyses used Python 3.13.12, R 4.5.2, Stata 18.0 and pandas 3.0.5 (version 2.1). "
           "Data: https://www.pmda.go.jp/safety/0003/x.html and doi:10.1000/abc123; approval "
           "no. 2023-0012; NCT01234567. The hazard ratio was 1.34 and median survival 3.5 months.")
    nums = numbers.NUM_RE.findall(numbers._scrub(txt))
    seq = [t for t in presentation.token_seq(txt) if presentation._isnum(t)]
    record("B6 software/package versions, URLs, DOIs and approval/registry identifiers are not "
           "extracted as numbers (no provenance demand, no coincidental match); real values stay",
           nums == ["1.34", "3.5"] and seq == ["1.34", "3.5"]
           and "0003" not in nonfacts.strip(txt), f"{nums} {seq}")
    # B7 --------------------------------------------------------------------------------
    reg = ("The protocol was not registered. The review was not registered in PROSPERO. The "
           "study was not prospectively registered.")
    dev = "We deviated from the protocol and originally planned a different outcome."
    r = _paper(tmp / "b7")
    _w(r, "08_submission/manuscript/polish_allowlist.tsv",
       "independent review\tthe journal's name for its external review board\n")
    allow = textchecks.polish_allowed(r)
    record("B7 registration-status statements required by reporting guidelines are not deviation "
           "narratives; real deviation narratives still are; polish_allowlist.tsv in "
           "08_submission/manuscript/ applies, also to backstage rules",
           textchecks.deviation_hits(reg) == [] and textchecks.deviation_hits(dev)
           and "independent review" in allow
           and textchecks.unallowed(textchecks.backstage_hits("an independent review board"), allow) == []
           and textchecks.unallowed(["earlier draft"], allow) == ["earlier draft"],
           str(textchecks.deviation_hits(reg)))
    # B8 --------------------------------------------------------------------------------
    from docx import Document
    s = _paper(tmp / "b8")
    for name in ("a", "b"):
        d_ = Document()
        d_.add_paragraph(f"text {name}")
        (s / "08_submission/bundle").mkdir(parents=True, exist_ok=True)
        d_.save(s / f"08_submission/bundle/{name}.docx")
    readerreview.write_text_copies(s, ["08_submission/bundle/a.docx", "08_submission/bundle/b.docx"])
    readerreview.write_text_copies(s, ["08_submission/bundle/a.docx"], round_id="M002")
    top = sorted(x.name for x in (s / readerreview.TEXT_DIR_REL).glob("*.txt"))
    per = sorted(x.name for x in (s / readerreview.TEXT_DIR_REL / "M002").glob("*.txt"))
    record("B8 a delta round no longer wipes .wf/review_text/; each round keeps its own copies",
           len(top) == 2 and per == ["08_submission__bundle__a.docx.txt"], f"{top} {per}")
    # B9 --------------------------------------------------------------------------------
    _w(s, "08_submission/bundle/manifest.json", {"items": [
        {"file": "08_submission/bundle/a.docx", "role": "manuscript"},
        {"file": "08_submission/bundle/b.docx", "role": "evidence", "upload": False}]})
    ups = packagefreeze.upload_files(s)
    record("B9 package freeze roles follow the manifest upload flag (upload:false is evidence)",
           "08_submission/bundle/a.docx" in ups and "08_submission/bundle/b.docx" not in ups, str(ups))
    # B10 -------------------------------------------------------------------------------
    from openpyxl import Workbook, load_workbook
    sys.path.insert(0, str(ROOT / "tools/figures"))
    import qc
    wb = Workbook()
    wb.active.title = "Table S1"
    wb.active["A1"] = "S1 content"
    ws2 = wb.create_sheet("Table S2")
    ws2["A1"] = "S2 content"
    book = tmp / "supp.xlsx"
    wb.save(book)
    entries = [{"id": "Table S1", "file": "04_tables/supplementary/supp.xlsx"},
               {"id": "Table S2", "file": "04_tables/supplementary/supp.xlsx"}]
    unnamed = [{"id": "eTable A", "file": "x.xlsx"}, {"id": "eTable B", "file": "x.xlsx"}]
    one = qc.table_sheet_copy(book, entries[1], entries, tmp)
    only = load_workbook(one)
    record("B10 the displays contact sheet renders each table's own sheet of a shared workbook "
           "(by sheet key, matching name or position), not Table S1 every time",
           qc.table_sheet(book, entries[0], entries) == "Table S1"
           and qc.table_sheet(book, entries[1], entries) == "Table S2"
           and qc.table_sheet(book, unnamed[1], unnamed) == "Table S2"
           and only.sheetnames == ["Table S2"] and only.active["A1"].value == "S2 content",
           str(only.sheetnames))
    # B11 -------------------------------------------------------------------------------
    paras = [("SectionHeading", "Abstract"), ("Normal", "one two three"),
             ("SectionHeading", "Introduction"), ("Normal", "alpha beta"),
             ("SectionHeading", "Discussion"), ("Normal", "gamma delta epsilon"),
             ("SectionHeading", "Statements and Declarations"), ("Normal", "funding none at all"),
             ("SectionHeading", "References"), ("Normal", "1. Ref")]
    doc1 = _docx(tmp / "counts1.docx", paras)
    c1 = journalspec.docx_counts(doc1, ["Introduction", "Methods", "Results", "Discussion"], {})
    paras2 = [x if x[1] != "Statements and Declarations" else ("SectionHeading", "Ethics and Funding")
              for x in paras]
    doc2 = _docx(tmp / "counts2.docx", paras2)
    c2 = journalspec.docx_counts(doc2, ["Introduction", "Discussion"],
                                 {"declarations_heading": "Ethics and Funding"})
    record("B11 the Word count recognises the declarations section by role ('Statements and "
           "Declarations', the journal's declarations_heading), so it is not counted as main text",
           c1["main"] == 5 and c1["abstract"] == 3 and c1["references"] == 1 and c2["main"] == 5,
           f"{c1} {c2}")
    # B12 -------------------------------------------------------------------------------
    from manuscript import checklist_docx
    from wfcore.checks.manuscript import _audit_docx
    t = _paper(tmp / "b12")
    _w(t, "08_submission/reporting_checklist.md", "Guideline: STROBE\n\n| Item | Status | "
       "Location | Evidence | Reason |\n|---|---|---|---|---|\n| STROBE 1a | Reported | Title | "
       "title | |\n| STROBE 2 | Reported | Introduction | text | |\n")
    with _envset(MEDPAPER_PROJECT=t):
        out = checklist_docx.build(t / "08_submission/reporting_checklist.md",
                                   t / "08_submission/bundle/STROBE_checklist.docx")
    style = bd.load_style_defaults()
    as_checklist = _audit_docx(out, style, "checklist")
    as_supp = _audit_docx(out, style, "supplementary")
    record("B12 the reporting checklist (a table form) passes the upload audit as role "
           "'checklist'; listing it as 'supplementary' names the right role",
           not as_checklist and any("role 'checklist'" in x for x in as_supp), str(as_checklist)[:300])
    # B13 -------------------------------------------------------------------------------
    sys.path.insert(0, str(ROOT / "tools/manuscript"))
    import preview
    from PIL import Image
    imgs = []
    for i in range(30):
        im = tmp / f"page-{i + 1:03d}.png"
        Image.new("RGB", (20, 28), "white").save(im)
        imgs.append(im)
    outdir = tmp / "sheets"
    outdir.mkdir()
    made = preview.sheets(imgs, outdir, "manuscript-stamp", 4)
    record("B13 preview renders every page by default, several contact sheets with unique names "
           "when needed",
           preview._pages(None, 30) == list(range(30)) and len(made) == 3
           and len({m.name for m in made}) == 3 and all(m.is_file() for m in made),
           str([m.name for m in made]))


# ------------------------------------------------------------------------------ R
def run_r(tmp: Path, record, section) -> None:
    from wfcore import registry
    from wfcore.state import State
    section("1.14.1 release: version, one migration event, pinned paper policy, card length")
    q = tmp / "old"
    q.mkdir(parents=True, exist_ok=True)
    st = State(q, ".wf").create("medpaper", "1.14.0", "S25_submission_audit")
    st.migrate_pipeline("1.14.1", ["S25_submission_audit"])
    st.migrate_pipeline("1.14.1", ["S25_submission_audit"])
    notes = [e for e in st.data["events"] if e["kind"] == "requirements_changed"
             and str(e.get("detail", "")).startswith("1.14.1")]
    pipe = registry.load()
    sizes = {s.id: len((ROOT / "pipeline" / s.card).read_text(encoding="utf-8")) for s in pipe.stages}
    from wfcore.submissionqc import _since_1141
    old_ctx = SimpleNamespace(state=SimpleNamespace(engine_pin=st.engine_pin))
    new_ctx = SimpleNamespace(state=SimpleNamespace(engine_pin="1.14.1"))
    record("R pipeline 1.14.1; a 1.14.0 paper (e.g. waiting at S25) gets exactly one 1.14.1 "
           "event, stays pinned (stricter checks only after wf migrate --yes); cards <= limit",
           pipe.meta.get("version") >= "1.14.1" and len(notes) == 1 and st.engine_pin == "1.14.0"
           and not _since_1141(old_ctx) and _since_1141(new_ctx)
           and max(sizes.values()) <= registry.CARD_MAX_CHARS, f"{len(notes)} {max(sizes.values())}")


GROUPS = [("P1", run_p1), ("P2", run_p2), ("P3", run_p3), ("B", run_b), ("R", run_r)]

#!/usr/bin/env python3
"""Reporting checklist upload DOCX (1.9.0, #23) - only when the journal requires one.

    .\\.venv\\Scripts\\python.exe tools/manuscript/checklist_docx.py \\
        --output 08_submission/bundle/STROBE_checklist.docx [--official official_items.json] [--pages]

Reads 08_submission/reporting_checklist.md (the gated source). The upload has fixed column
widths (Item narrow and never wrapped, Topic / Location wide), no internal paths (the
`Manuscript:` line and any project path are dropped from the header), black Times New Roman,
landscape A4. `--official` maps item id -> the official item wording (supplied by the author
from the official checklist; the engine never invents it). `--pages` (JSON item id -> page,
read off the rendered manuscript, e.g. with tools/manuscript/preview.py) adds a Page column.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
from wfcore import determinism, paths as wfpaths  # noqa: E402

PATH_RE = re.compile(r"(?:\b\d\d_[a-z]+/[\w./-]+|[A-Za-z]:[\\/][^\s|]+|\.wf/[\w./-]+)")


def parse(text: str) -> list[dict]:
    """[{guideline, rows: [{item, status, location, evidence, reason}]}]."""
    blocks: list[dict] = []
    cur = None
    for line in text.splitlines():
        m = re.match(r"^\s*Guideline:\s*(.+)$", line)
        if m:
            cur = {"guideline": m.group(1).strip(), "rows": []}
            blocks.append(cur)
            continue
        if cur is None or not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not cells or cells[0].casefold() == "item" or set(cells[0]) <= set("-: "):
            continue
        cells += [""] * (5 - len(cells))
        cur["rows"].append(dict(zip(("item", "status", "location", "evidence", "reason"), cells)))
    return blocks


def _titles() -> dict[str, dict[str, str]]:
    import tomllib
    reg = tomllib.loads((ROOT / "reference/reporting_guidelines.toml").read_text(encoding="utf-8"))
    return {gid: {str(i["id"]): i.get("title", "") for i in g.get("items", [])}
            for gid, g in reg.get("guideline", {}).items()}


def build(src: Path, out: Path, official: dict[str, str] | None = None,
          pages: dict[str, str] | None = None) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    blocks = parse(src.read_text(encoding="utf-8"))
    if not blocks:
        raise ValueError(f"{src.name}: no 'Guideline:' block with a table")
    titles = _titles()
    font = "Times New Roman"
    try:  # 1.14.1 (B12): the journal's Word font, so the upload audit passes as role checklist
        style = json.loads((wfpaths.project_dir(check_engine=False)
                            / "08_submission/docx_style.json").read_text(encoding="utf-8"))
        font = str(style.get("font_family") or font)
    except (OSError, ValueError, SystemExit):
        pass

    def fmt(run, size: float = 10):
        """Explicit black journal font on every run (what docx_bundle_ready audits)."""
        run.font.name, run.font.size = font, Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        return run
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    cols = [("Item", 1.6), ("Topic", 6.5 if official else 5.0)]
    if official:
        cols.append(("Recommendation", 8.0))
    cols += [("Reported on / location", 7.0 if not official else 5.5), ("Status", 2.4)]
    if pages is not None:
        cols.append(("Page", 1.6))
    for block in blocks:
        head = doc.add_paragraph()
        run = fmt(head.add_run(f"{block['guideline']} checklist"), 12)
        run.bold = True
        gid = block["guideline"].split("+")[0].strip().upper()
        table = doc.add_table(rows=1, cols=len(cols))
        table.autofit = False
        tblpr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        for i, (name, width) in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.width = Cm(width)
            fmt(cell.paragraphs[0].add_run(name)).bold = True
        for row in block["rows"]:
            item = re.sub(r"^(?:STROBE|RECORD|PRISMA|CONSORT)\s+", "", row["item"], flags=re.I)
            values = [row["item"], titles.get(gid, {}).get(item, "")]
            if official:
                values.append(official.get(row["item"], official.get(item, "")))
            loc = row["location"] if row["status"].casefold() == "reported" else row["reason"]
            values += [PATH_RE.sub("", loc).strip(), row["status"]]
            if pages is not None:
                values.append(pages.get(row["item"], ""))
            cells = table.add_row().cells
            for i, (value, (_, width)) in enumerate(zip(values, cols)):
                cells[i].width = Cm(width)
                fmt(cells[i].paragraphs[0].add_run(value))
                if i == 0:  # the item id stays on one line
                    pr = cells[i]._tc.get_or_add_tcPr()
                    nowrap = OxmlElement("w:noWrap")
                    pr.append(nowrap)
        grid = table._tbl.tblGrid
        for gc, (_, width) in zip(grid.findall(qn("w:gridCol")), cols):
            gc.set(qn("w:w"), str(int(Cm(width).twips)))
        doc.add_paragraph()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(out)
    _strip_template_controls(out)
    determinism.scrub_properties(out)  # 1.14.1 (B5)
    determinism.normalize(out, blind=True)
    return out


def _strip_template_controls(path: Path) -> None:
    """1.14.1 (B12): the python-docx template's styles carry keepNext/keepLines/outlineLvl
    (headings) and a paragraph border (Title); the upload audit refuses all of them (black
    squares, outline levels), so they go, as in build_docx."""
    import io
    import zipfile
    with zipfile.ZipFile(path) as src:
        entries = [(i, src.read(i.filename)) for i in src.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, blob in entries:
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                for tag in (b"keepNext", b"keepLines", b"outlineLvl", b"pageBreakBefore", b"pBdr"):
                    blob = re.sub(rb"<w:" + tag + rb"\b[^>]*/>", b"", blob)
                    blob = re.sub(rb"<w:" + tag + rb"\b[^>]*>.*?</w:" + tag + rb">", b"", blob,
                                  flags=re.S)
            dst.writestr(info, blob)
    path.write_bytes(buf.getvalue())


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="build the reporting-checklist upload DOCX")
    ap.add_argument("--input", type=Path, default=None)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--official", type=Path, help="JSON {item id: official wording} (author-supplied)")
    ap.add_argument("--pages", type=Path, help="JSON {item id: page} to add a Page column")
    args = ap.parse_args(argv)
    proj = wfpaths.project_dir()
    src = args.input or proj / "08_submission/reporting_checklist.md"
    src = src if src.is_absolute() else proj / src
    out = wfpaths.paper_output_path(args.output if args.output.is_absolute()
                                    else "project/" + args.output.as_posix())
    official = json.loads(args.official.read_text(encoding="utf-8")) if args.official else None
    pages = json.loads(args.pages.read_text(encoding="utf-8")) if args.pages else None
    try:
        build(src, out, official, pages)
    except (OSError, ValueError) as exc:
        print(f"checklist_docx: {exc}", file=sys.stderr)
        return 2
    print(f"built reporting checklist -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

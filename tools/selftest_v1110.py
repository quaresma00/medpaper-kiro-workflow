"""1.11.0 regression tests: the clinical presentation layer (display dictionary, engineering
language, self-contained displays, flow-diagram units, Word keep-with-next, closed review
findings), early journal parameters, result presentation, concise footnotes, revision preview
and rewind confirmation, reproduction decoupling, quiet long commands, freeze layers,
traceable final-copy numbers, proportion denominators, role-based gates, card encoding,
review bundles and the in-paper progress file. Synthetic fixtures only.

Called from tools/selftest.py.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))


def _w(p: Path, rel: str, text) -> None:
    (p / rel).parent.mkdir(parents=True, exist_ok=True)
    (p / rel).write_text(text if isinstance(text, str) else json.dumps(text), encoding="utf-8")


def _paper(base: Path, at: str) -> Path:
    import selftest_v180 as v180
    return v180._rw_paper(base, at)


def _gate(p: Path, name: str, stage: str, **spec):
    import selftest_v180 as v180
    return v180._gate(p, name, stage, **spec)


def _table(p: Path, rel: str, header, rows, footnotes, title="Table 1. Characteristics of reports") -> None:
    from openpyxl import Workbook
    from tables import threeline
    wb = Workbook()
    ws = wb.active
    threeline._write_sheet(ws, {"sheet": "Table 1", "title": title, "header": header, "rows": rows,
                                "footnotes": footnotes})
    (p / rel).parent.mkdir(parents=True, exist_ok=True)
    wb.save(p / rel)


def _display_paper(base: Path, at: str = "S11_figures") -> Path:
    p = _paper(base, at)
    _w(p, "01_protocol/artifact_plan.json", {
        "main_figures": [{"id": "Figure 1", "file": "05_figures/out/F1.png", "archetype": "flow_diagram"}],
        "main_tables": [{"id": "Table 1", "file": "04_tables/main/T1.xlsx"}],
        "supp_figures": [], "supp_tables": []})
    _w(p, "01_protocol/display_dictionary.json", {"terms": [
        {"internal": "comp_a", "aliases": ["Composite A"], "label": "Portal vascular disease composite",
         "abbreviation": "PVD", "definition": "nodular regenerative hyperplasia or portal vein thrombosis"}]})
    _w(p, "05_figures/legends.md", "# Figure legends\n\n## Figure 1. Selection of reports\n\n"
       "Flow of adverse-event reports from the database to the analysis.\n")
    _w(p, "04_tables/table_captions.md", "# Table captions\n\n## Table 1. Characteristics\n\n"
       "Characteristics of the analysed reports by drug group.\n")
    return p


def _sidecar(p: Path, texts, min_font=8.0, stem="F1") -> None:
    _w(p, f"05_figures/qc/{stem}.artist.json", {"figure": stem, "figure_text": list(texts),
                                                  "min_font_pt": min_font})


# ------------------------------------------------------------------------------ group 1
def run_dictionary(tmp: Path, record, section) -> None:
    from wfcore import clinical
    section("1.11.0 #1a display dictionary + clinical labels")
    p = _display_paper(tmp / "dict", "S07_artifacts")
    ok = _gate(p, "display_dictionary_ready", "S07_artifacts")
    _w(p, "01_protocol/display_dictionary.json", {"terms": [
        {"internal": "comp_a", "label": "strict_contrast_a", "definition": "x"},
        {"internal": "grp_b", "label": "Group B", "definition": "a group of reports"}]})
    bad = _gate(p, "display_dictionary_ready", "S07_artifacts")
    (p / "01_protocol/display_dictionary.json").unlink()
    missing = _gate(p, "display_dictionary_ready", "S07_artifacts")
    record("S07 requires a display dictionary: internal name -> clinical label, abbreviation, definition",
           ok.ok and not bad.ok and "code identifier" in bad.detail and "letter code" in bad.detail
           and not missing.ok, bad.detail)
    data = {"terms": [{"internal": "comp_a", "aliases": ["Composite A"], "label": "Portal vascular disease composite",
                       "definition": "defined outcome"},
                      {"internal": "reference", "label": "Irinotecan reference group", "definition": "comparator"}]}
    hits = clinical.label_problems("Composite A rose; comp_a; Group B differed; Model 3 was used.", data)
    clean = clinical.label_problems("The portal vascular disease composite and the Irinotecan reference "
                                    "group; Group A streptococcal infection; type 2 diabetes.", data)
    record("internal names and letter-code labels are flagged; clinical labels and exceptions are not",
           len(hits) == 4 and not clean, f"{hits} | {clean}")
    _display_paper(tmp / "dict2", "S11_figures")
    q = tmp / "dict2"
    _sidecar(q, ["Composite A", "Reports analysed (n = 120)"])
    fig = _gate(q, "clinical_labels_used", "S11_figures", displays="figures")
    _sidecar(q, ["Portal vascular disease composite", "Reports analysed (n = 120)"])
    fig_ok = _gate(q, "clinical_labels_used", "S11_figures", displays="figures")
    record("in-figure text from the render sidecar must use the clinical label",
           not fig.ok and "Composite A" in fig.detail and fig_ok.ok, fig.detail)


def run_engineering(tmp: Path, record, section) -> None:
    from wfcore import clinical
    section("1.11.0 #1b engineering-language lint (body, tables, figures, legends)")
    p = _display_paper(tmp / "eng", "S17_assemble")
    full = ("# Title\n\n# Abstract\n\nWe studied reports.\n\n# Methods\n\nNames underwent Unicode "
            "normalization and exact match against a vocabulary frozen from all strings observed.\n\n"
            "# Results\n\nWe analysed 1,204 reports.\n\n# Discussion\n\nDeduplication of reports and "
            "frozen section histology were standard; analyses used Python 3.11 and R 4.5.\n\n# References\n")
    _w(p, "07_manuscript/full_manuscript.md", full)
    warn = _gate(p, "engineering_language_clean", "S17_assemble", paths=["07_manuscript/full_manuscript.md"])
    _w(p, "07_manuscript/full_manuscript.md", full.replace(
        "We analysed 1,204 reports.", "In total 24812425 demographic lines read; all analyses were "
        "re-executed from the raw files and reproduced identically."))
    block = _gate(p, "engineering_language_clean", "S17_assemble", paths=["07_manuscript/full_manuscript.md"])
    _w(p, "07_manuscript/supplementary_methods.md", "Names were parsed with a regex; first 15 alphanumeric characters.\n")
    supp = _gate(p, "engineering_language_clean", "S17_assemble", paths=["07_manuscript/supplementary_methods.md"])
    record("Methods implementation terms warn; Results/Discussion terms block; Supplementary Methods is free",
           not warn.ok and warn.severity == "warn" and "normalization" in warn.detail.casefold()
           and "Deduplication" not in warn.detail and "frozen section" not in warn.detail
           and not block.ok and block.blocking and "demographic lines" in block.detail
           and "re-executed" in block.detail and supp.ok,
           warn.detail[:200] + " | " + block.detail[:200])
    _w(p, ".wf/language_allowlist.json", {"terms": [
        {"term": "normalization", "reason": "the journal's data standard names this step", "scope": "methods"},
        {"term": "string", "reason": "too short"}]})
    allow = clinical.allowlist(p)
    record("a paper-local allowlist needs a reason (>= 15 chars) and a scope",
           allow == {"normalization": "methods"}, str(allow))
    _sidecar(p, ["Exact duplicates collapsed", "Reports analysed (n = 120)"])
    fig = _gate(p, "engineering_language_clean", "S11_figures", displays="figures")
    record("in-figure text is linted (sidecar figure_text)", not fig.ok and "collapsed" in fig.detail, fig.detail)
    terms = {e["term"] for e in clinical.rules()["engineering"]}
    record("the maintained term list covers the reported engineering words",
           {"normalization", "string", "field", "record", "snapshot", "collapsed", "pipeline", "script",
            "heuristic", "cluster rule", "re-executed", "reproduced identically", "lines read", "frozen",
            "parsed", "regex", "alphanumeric", "token", "deduplicated by key"} <= terms, str(sorted(terms)))
    card = (ROOT / "pipeline/stages/S08_methods.md").read_text(encoding="utf-8")
    record("S08 card states the layering rule (principles in Methods, implementation in supplement)",
           "Supplementary Methods" in card and "principle" in card.casefold() and "1.11.0" in card)


def run_abbreviations(tmp: Path, record, section) -> None:
    from wfcore.state import State
    section("1.11.0 #1c each table/figure defines its abbreviations; central list kept")
    p = _display_paper(tmp / "abbr", "S17_assemble")
    _table(p, "04_tables/main/T1.xlsx", ["Characteristic", "ROR (95% CI)"], [["Age", "1.20 (1.01-1.43)"]],
           ["Data are n (%) unless otherwise indicated."])
    _sidecar(p, ["Reports screened (n = 1,204)", "PVD reports (n = 30)"])
    missing = _gate(p, "display_abbreviations_defined", "S17_assemble")
    _table(p, "04_tables/main/T1.xlsx", ["Characteristic", "ROR (95% CI)"], [["Age", "1.20 (1.01-1.43)"]],
           ["Data are n (%) unless otherwise indicated.",
            "Abbreviations: CI, confidence interval; ROR, reporting odds ratio."])
    _w(p, "05_figures/legends.md", "# Figure legends\n\n## Figure 1. Selection of reports\n\n"
       "Flow of reports. Abbreviations: PVD, portal vascular disease composite.\n")
    local = _gate(p, "display_abbreviations_defined", "S17_assemble")
    nocentral = _gate(p, "display_abbreviations_defined", "S17_assemble", central_list=True)
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\n# Abstract\n\nx\n\n## Abbreviations\n\nCI, confidence "
       "interval; PVD, portal vascular disease composite; ROR, reporting odds ratio.\n")
    central = _gate(p, "display_abbreviations_defined", "S17_assemble", central_list=True)
    record("a display that uses an abbreviation defines it locally; the manuscript also keeps a central list",
           not missing.ok and "ROR" in missing.detail and "PVD" in missing.detail and local.ok
           and not nocentral.ok and central.ok, missing.detail + " | " + nocentral.detail)
    _w(p, "08_submission/target_journal.json", {"abbreviation_placement": "central_required"})
    journal = _gate(p, "display_abbreviations_defined", "S21_authors", central_list=True)
    st = State(p, ".wf").load()
    st.data["engine_pin"] = "1.10.0"
    st.save()
    _w(p, "08_submission/target_journal.json", {"abbreviation_placement": "central"})
    legacy = _gate(p, "display_abbreviations_defined", "S21_authors", central_list=True)
    record("only an explicit journal rule centralises; a pinned paper's legacy placement keeps the old rule",
           not journal.ok and "central list" in journal.detail and "legacy placement" in legacy.detail,
           journal.detail + " | " + legacy.detail)
    import tomllib
    raw = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
    used = {g["check"] for s in raw["stage"] for g in s["gate"]}
    record("abbreviations_centralized is no longer wired; display_abbreviations_defined is",
           "abbreviations_centralized" not in used and "display_abbreviations_defined" in used)


def run_flow(tmp: Path, record, section) -> None:
    import tomllib
    from figures import qc
    from wfcore import clinical
    section("1.11.0 #1d flow diagrams are schematics: boxes, counts, arrows and analysis units")
    reg = tomllib.loads((ROOT / "reference/archetypes.toml").read_text(encoding="utf-8"))
    found = {"box_nodes": {"found": True}, "figure_statistics": {"found": True},
             "arrows": {"found": True}, "axis_labels": {"found": False}}
    checks = qc._archetype_checks("flow_diagram", found, reg)
    names = {c["name"]: c for c in checks}
    km = {c["name"] for c in qc._archetype_checks("km_survival", {"axis_labels": {"found": True}}, reg)}
    record("a flow diagram needs boxes, counts and arrows - not axis titles; plots still need axes",
           "element:axis_labels" not in names and names["element:arrows"]["ok"]
           and reg["archetype"]["flow_diagram"]["kind"] == "schematic" and "element:axis_labels" in km)
    probs = clinical.flow_unit_problems(["Database quarterly files (n = 90)", "24812425 demographic lines read",
                                         "Reports screened (n = 1,204)", "Excluded (n = 30)",
                                         "Patients analysed\nn = 412"])
    record("each count box names its analysis unit; processing-log units are rejected",
           len(probs) == 3 and "files" in probs[0] and "lines" in probs[1] and "Excluded" in probs[2], str(probs))
    p = _display_paper(tmp / "flow", "S11_figures")
    _sidecar(p, ["24812425 demographic lines read", "Reports analysed (n = 120)"])
    bad = _gate(p, "flow_diagram_units", "S11_figures")
    _sidecar(p, ["Reports received (n = 2,400)", "Reports analysed (n = 120)"])
    good = _gate(p, "flow_diagram_units", "S11_figures")
    record("flow_diagram_units gate reads the in-figure text", not bad.ok and good.ok, bad.detail)


def run_keep(tmp: Path, record, section) -> None:
    import zipfile
    from xml.etree import ElementTree as ET
    from docx import Document
    from manuscript import build_docx as bd
    from wfcore import docxaudit as da
    section("1.11.0 #1e Word keeps caption -> table -> footnote and legend -> figure together")
    out = tmp / "keep.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.add_heading("Results", level=1)
    d.add_paragraph("Table 1. Characteristics of reports")
    t = d.add_table(rows=3, cols=2)
    for r in range(3):
        for c in range(2):
            t.cell(r, c).text = f"r{r}c{c}"
    d.add_paragraph("Data are n (%).")
    d.add_paragraph("Abbreviations: CI, confidence interval.")
    d.add_heading("Discussion", level=1)
    d.save(out)
    with zipfile.ZipFile(out) as zf:
        before = da.keep_problems(ET.fromstring(zf.read("word/document.xml")))
    d = Document(out)
    for p in d.paragraphs:
        bd._remove_heading_metadata(p, strip_numbering=False)
    bd._keep_together(d)
    d.save(out)
    bd._sanitize_package_xml(out, "Arial")
    with zipfile.ZipFile(out) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
        styles = zf.read("word/styles.xml")
    after = da.keep_problems(root)
    record("unbound captions/footnotes are detected; the builder binds them and keeps them after sanitising",
           before and not after and b"keepNext" not in styles
           and b"keepNext" in ET.tostring(root), f"{before} | {after}")
    record("headings carry no keepNext/keepLines collapse markers", not da.heading_keep_problems(root))
    record("FORBIDDEN controls are outline level and page-break-before only; keep is heading/style-only",
           bd.FORBIDDEN_PARAGRAPH_CONTROLS == ("outlineLvl", "pageBreakBefore")
           and bd.HEADING_ONLY_CONTROLS == ("keepNext", "keepLines"))


def run_review(tmp: Path, record, section) -> None:
    import selftest_v181 as v181
    from wfcore import reviewrounds as rr
    section("1.11.0 #1f every numbered finding closed; neutral brief; #3f review bundle")
    p, env = v181._review_paper(tmp / "rev")
    env = {**env, "MEDPAPER_NO_PREVIEW": "1"}

    def cli(*args):
        return subprocess.run([sys.executable, "-B", str(ROOT / "tools/review_round.py"), *args,
                               "--project", str(p)], capture_output=True, text=True,
                              encoding="utf-8", errors="replace", env=env)
    start = cli("start")
    rnd = rr.latest(p, "manuscript")
    manifest = p / rr.BUNDLE_DIR / rnd["id"] / "MANIFEST.md"
    mtext = manifest.read_text(encoding="utf-8") if manifest.is_file() else ""
    record("start writes one self-contained bundle and prints a neutral brief with both perspectives",
           start.returncode == 0 and "MANIFEST.md" in start.stdout and "first-time clinical reader" in mtext
           and "engineering language" in mtext and "known deviation" not in mtext.casefold()
           and list((manifest.parent / "files").glob("*")), (start.stdout + start.stderr)[-300:])
    report = tmp / "report.md"
    report.write_text("# Verdict\n\n- F1 [numbers] HR rounding in the abstract\n\n# Reader comprehension\n\n"
                      "F2 [other] 'Composite A' slows reading\n\n# Medical presentation\n\n"
                      "**F3** [other] Figure 2 legend lacks the abbreviation list\n", encoding="utf-8")
    none = cli("close", "--findings", "major=1", "--finding", "F1", "numbers", "changed", "rounded")
    partial = cli("close", "--findings", "major=1", "--finding", "F1", "numbers", "changed", "rounded",
                  "--report", str(report))
    noreason = cli("close", "--findings", "major=1,minor=2", "--finding", "F1", "numbers", "changed", "rounded",
                   "--finding", "F2", "other", "changed", "clinical label", "--finding", "F3", "other",
                   "kept", "legend", "--report", str(report))
    ok = cli("close", "--findings", "major=1,minor=2", "--finding", "F1", "numbers", "changed", "rounded",
             "--finding", "F2", "other", "changed", "clinical label", "--finding", "F3", "other", "kept",
             "legend", "--kept", "F3", "the journal puts abbreviations in one list", "--report", str(report))
    record("close needs the saved report; Reader/Medical findings must be closed; kept needs a reason",
           none.returncode != 0 and "--report" in none.stderr and partial.returncode != 0
           and "F2" in partial.stderr and "F3" in partial.stderr and noreason.returncode != 0
           and "without a reason" in noreason.stderr and ok.returncode == 0,
           " | ".join(x.stderr[-160:] for x in (none, partial, noreason, ok)))
    g = _gate(p, "review_findings_closed", "S19_human_review", scope="manuscript")
    record("review_findings_closed gate passes with all F-numbers closed", g.ok and "3 numbered" in g.detail, g.detail)
    agent = (ROOT / ".kiro/agents/medpaper-reviewer.md").read_text(encoding="utf-8")
    card = (ROOT / "pipeline/stages/S18_independent_review.md").read_text(encoding="utf-8")
    record("reviewer agent and S18 card carry the perspectives from round 1 and the neutral-brief rule",
           "first-time clinical reader" in agent and "every section" in agent.casefold()
           and "known deviations" in card and "MANIFEST.md" in card)


# ------------------------------------------------------------------------------ group 2
def run_params(tmp: Path, record, section) -> None:
    from wfcore import numfmt, presentparams as pp
    section("1.11.0 #2a provisional journal parameters from S07; #2b print-size fonts")
    p = _display_paper(tmp / "params", "S07_artifacts")
    miss = _gate(p, "provisional_journal_params", "S07_artifacts")
    prof = {"status": "provisional", "journal_or_type": "clinical journal type", "font_min_pt": 8,
            "figure_width_mm": {"single": 84, "1.5": 129, "double": 174}, "panel_label_case": "lower",
            "line_art_dpi": 600, "number_format": {"thousands": ","}, "spelling": "american",
            "heading_names": {}, "source": "journal-type default"}
    _w(p, pp.PROVISIONAL, prof)
    ok = _gate(p, "provisional_journal_params", "S07_artifacts")
    eff = pp.params(p)
    record("S07 records provisional presentation parameters; builders read them before S20",
           not miss.ok and ok.ok and eff["figure_width_mm"]["double"] == 174
           and eff["panel_label_case"] == "lower" and numfmt.style(p)["thousands"] == ",", ok.detail)
    _w(p, pp.TARGET, {"presentation": {**prof, "figure_width_mm": {"single": 90, "double": 180}}})
    conf = _gate(p, "presentation_params_confirmed", "S20_journal")
    record("S20 confirms or replaces them; a replacement is a presentation re-render",
           conf.ok and "figure_width_mm" in conf.detail and "presentation" in conf.detail
           and pp.params(p)["figure_width_mm"]["double"] == 180, conf.detail)
    _sidecar(p, ["Reports (n = 1)"], min_font=6.5)
    small = _gate(p, "figure_print_font", "S11_figures")
    _sidecar(p, ["Reports (n = 1)"], min_font=8.0)
    fine = _gate(p, "figure_print_font", "S11_figures")
    old = os.environ.get("MEDPAPER_PROJECT")
    os.environ["MEDPAPER_PROJECT"] = str(p)
    try:
        import matplotlib
        from figures import style
        style.apply_style()
        base = matplotlib.rcParams["font.size"]
        letter = style.panel_letter(1)
    finally:
        if old is None:
            os.environ.pop("MEDPAPER_PROJECT", None)
        else:
            os.environ["MEDPAPER_PROJECT"] = old
    medfig = (ROOT / "tools/figures/R/medfig.R").read_text(encoding="utf-8")
    record("6.5 pt text fails the 8 pt print floor; Python/R defaults are >= 8 pt; panel case from params",
           not small.ok and fine.ok and base >= 8 and letter == "b"
           and "medpaper_font_floor" in medfig and "mf_tag_levels" in medfig
           and "base_size = 7" not in medfig, small.detail)


def run_values(tmp: Path, record, section) -> None:
    from tables import threeline
    section("1.11.0 #2c result presentation standard; #2d concise footnotes, diagnostics to supplement")
    p = _display_paper(tmp / "vals", "S10_tables")
    _table(p, "04_tables/main/T1.xlsx", ["Outcome", "ROR (95% CI)", "Collapse step"],
           [["Outcome 1", "Undefined (19.6 to infinity)", "3"], ["Outcome 2", "Not fitted", "NA"],
            ["Outcome 3", "-", "\u2013"]],
           ["Data are ROR (95% CI). Woolf intervals were estimated", "for nonzero cells. Firth penalized "
            "regression was used. The model was adjusted for age. Iterations converged."])
    raw = _gate(p, "display_values_presented", "S10_tables", displays="tables")
    foot = _gate(p, "table_footnotes_concise", "S10_tables")
    try:
        threeline._check({"sheet": "T", "title": "Table 1. X", "header": ["a"], "rows": [["Not fitted"]],
                          "footnotes": ["Data are n."]})
        raised = False
    except ValueError as exc:
        raised = "raw status" in str(exc)
    record("raw statuses, mixed missing symbols are blocked; not_estimable() states the bound",
           not raw.ok and "Undefined" in raw.detail and "missing" in raw.detail and raised
           and threeline.not_estimable(19.6) == "Not estimable (lower 95% CI bound 19.60)"
           and threeline.fmt(float("nan")) == threeline.MISSING, raw.detail)
    record("split footnote sentences, method-heavy footnotes and diagnostic columns in a main table fail",
           not foot.ok and "split" in foot.detail and "method-explanation" in foot.detail
           and "Collapse step" in foot.detail, foot.detail)
    long_note = "Data are n (%). " + "Values are counts of reports with the stated characteristic. " * 6
    _table(p, "04_tables/main/T1.xlsx", ["Characteristic", "Reports"], [["Age", "12"]], [long_note])
    tl = _gate(p, "tables_threeline", "S10_tables")
    record("a single long footnote row is no longer failed by the 300-character cell rule",
           tl.ok or "chars of prose" not in tl.detail, tl.detail)


def run(tmp: Path, record, section) -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    old = os.environ.get("MEDPAPER_PROJECT")
    try:
        run_dictionary(tmp, record, section)
        run_engineering(tmp, record, section)
        run_abbreviations(tmp, record, section)
        run_flow(tmp, record, section)
        run_keep(tmp, record, section)
        run_review(tmp, record, section)
        run_params(tmp, record, section)
        run_values(tmp, record, section)
    finally:
        if old is None:
            os.environ.pop("MEDPAPER_PROJECT", None)
        else:
            os.environ["MEDPAPER_PROJECT"] = old

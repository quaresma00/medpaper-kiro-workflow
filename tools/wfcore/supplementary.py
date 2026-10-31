"""1.15.0 (E5): engine builders for supplementary files that had no builder.

- figures_pdf: the supplementary figures as ONE PDF - an information page (header
  "Online Resource n", title, authors - left out when the journal is double-blind -,
  corresponding author, journal), then one page per figure with its legend under it.
  Deterministic: same inputs, same bytes.
- xlsx_info: a supplementary workbook with an "Info" sheet first (the same header block and a
  list of the sheets); the table sheets are copied unchanged.

Both write a build receipt (08_submission/build_receipts/<name>.json) so `wf rebuild`
rebuilds them when a figure, legend, workbook or the author information changes. The header
label comes from target_journal.json `supplementary_label` (default "Online Resource"), the
same label build_docx --esm-number uses for Word supplementary files.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import tempfile
import textwrap
from datetime import datetime, timezone
from pathlib import Path

FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest().upper()


def info_lines(project: Path, number: int | None) -> list[str]:
    from . import journalspec
    project = Path(project)
    target = journalspec.load(project)
    head = f"{journalspec.esm_label(target)} {number}" if number else "Supplementary information"
    out = [head]
    title_path = project / "07_manuscript/title.md"
    if title_path.is_file():
        title = re.sub(r"^#\s+", "", title_path.read_text(encoding="utf-8").strip())
        if title:
            out.append(f"Article title: {title}")
    if target.get("journal"):
        out.append(f"Journal: {target['journal']}")
    if journalspec.blinded(target):
        return out
    try:
        info = json.loads((project / "00_input/author_info.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        info = {}
    names = [a.get("name", "") if isinstance(a, dict) else str(a)
             for a in (info.get("authors") or []) if a]
    if any(names):
        out.append("Authors: " + ", ".join(n for n in names if n))
    corr = info.get("corresponding")
    if isinstance(corr, dict):
        corr = {**corr, "affiliation": corr.get("affiliation") or corr.get("address")}
        corr = ", ".join(str(corr[k]) for k in ("name", "email", "affiliation") if corr.get(k))
    if corr:
        out.append(f"Corresponding author: {corr}")
    return out


def _plan(project: Path) -> dict:
    try:
        return json.loads((Path(project) / "01_protocol/artifact_plan.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _canon(fid: str) -> str:
    m = re.search(r"(S?\d+)\s*$", str(fid).strip(), re.I)
    return m.group(1).upper() if m else str(fid).strip().upper()


def legends(project: Path) -> dict[str, str]:
    """'S1' -> legend text (heading and body) from 05_figures/legends.md."""
    path = Path(project) / "05_figures/legends.md"
    if not path.is_file():
        return {}
    text = path.read_text(encoding="utf-8")
    out = {}
    heads = list(re.finditer(r"(?mi)^#{2,6}\s+(?:Figure|Fig\.?)\s+(S?\d+)\b\.?[ \t]*(.*)$", text))
    for i, m in enumerate(heads):
        end = heads[i + 1].start() if i + 1 < len(heads) else len(text)
        nxt = re.search(r"(?m)^#\s", text[m.end():end])
        body = text[m.end():m.end() + nxt.start()] if nxt else text[m.end():end]
        title = m.group(2).strip()
        out[m.group(1).upper()] = " ".join(f"Figure {m.group(1).upper()}. {title} {body}".split())
    return out


def _image(project: Path, entry: dict) -> Path | None:
    raw = str(entry.get("file", "")).strip()
    if not raw:
        return None
    base = Path(project) / raw
    for cand in (base.with_suffix(".png"), base, base.with_suffix(".tif"), base.with_suffix(".tiff")):
        if cand.is_file():
            return cand
    return None


def figure_entries(project: Path, ids: list[str] | None) -> list[dict]:
    plan = _plan(project)
    pool = list(plan.get("supp_figures", []) or [])
    if ids:
        want = [_canon(i) for i in ids]
        pool += [e for e in plan.get("main_figures", []) or [] if _canon(e.get("id", "")) in want]
        by = {_canon(e.get("id", "")): e for e in pool}
        missing = [i for i in want if i not in by]
        if missing:
            raise ValueError("not in artifact_plan.json: " + ", ".join(missing))
        return [by[i] for i in want]
    return pool


def _receipt(project: Path, output: Path, kind: str, sources: list[Path], args: dict) -> Path:
    from .packagecontent import build_receipt_path
    project = Path(project)
    rec_path = build_receipt_path(project, output)
    aux = [project / r for r in ("08_submission/target_journal.json", "00_input/author_info.json",
                                 "07_manuscript/title.md") if (project / r).is_file()]
    rows = lambda ps: [{"path": p.relative_to(project).as_posix(), "sha256": _sha(p),  # noqa: E731
                        "bytes": p.stat().st_size} for p in ps]
    payload = {"schema_version": 1, "kind": kind, "built_at": datetime.now(timezone.utc).isoformat(
        timespec="seconds"), "output": output.relative_to(project).as_posix(),
        "output_sha256": _sha(output), "sources": rows(sources), "auxiliaries": rows(aux),
        "args": args}
    rec_path.parent.mkdir(parents=True, exist_ok=True)
    rec_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return rec_path


def figures_pdf(project: Path, output: Path, number: int | None = None,
                ids: list[str] | None = None) -> tuple[Path, Path]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from PIL import Image
    from . import determinism, docxaudit, journalspec
    from .revision import guard_outputs, record_derived
    project, output = Path(project), Path(output)
    entries = figure_entries(project, ids)
    if not entries:
        raise ValueError("no supplementary figures in artifact_plan.json (supp_figures)")
    texts = legends(project)
    images = []
    for e in entries:
        img = _image(project, e)
        if img is None:
            raise ValueError(f"{e.get('id')}: figure file missing ({e.get('file')})")
        key = _canon(e.get("id", ""))
        if key not in texts:
            raise ValueError(f"{e.get('id')}: no legend in 05_figures/legends.md")
        images.append((e, img, texts[key]))
    from .packagecontent import build_receipt_path
    guard_outputs(project, [output, build_receipt_path(project, output)])
    output.parent.mkdir(parents=True, exist_ok=True)
    head = info_lines(project, number)
    matplotlib.rcParams.update({"pdf.fonttype": 42, "font.family": "DejaVu Sans",
                                "svg.hashsalt": "medpaper"})
    a4 = (8.27, 11.69)
    with tempfile.TemporaryDirectory(prefix="medpaper_esm_") as work:
        tmp_pdf = Path(work) / "esm.pdf"
        with PdfPages(tmp_pdf, metadata={"Creator": None, "Producer": None,
                                         "CreationDate": None, "ModDate": None}) as pdf:
            fig = plt.figure(figsize=a4)
            y = 0.92
            for i, line in enumerate(head):
                for part in textwrap.wrap(line, 80) or [""]:
                    fig.text(0.08, y, part, fontsize=14 if i == 0 else 10,
                             weight="bold" if i == 0 else "normal", va="top")
                    y -= 0.035 if i == 0 else 0.025
            pdf.savefig(fig)
            plt.close(fig)
            for e, img, legend in images:
                fig = plt.figure(figsize=a4)
                fig.text(0.08, 0.965, head[0], fontsize=9, weight="bold", va="top")
                with Image.open(img) as im:
                    arr = im.convert("RGB")
                    w, h = arr.size
                    lines = textwrap.wrap(legend, 105)
                    room = 0.86 - 0.02 * len(lines) - 0.03
                    ax_h = max(0.2, room)
                    ax_w = min(0.84, ax_h * a4[1] / a4[0] * w / h)
                    ax_h = ax_w * a4[0] / a4[1] * h / w
                    ax = fig.add_axes([0.08, 0.94 - ax_h, ax_w, ax_h])
                    ax.imshow(arr, interpolation="none")
                    ax.axis("off")
                yy = 0.94 - ax_h - 0.03
                for line in lines:
                    fig.text(0.08, yy, line, fontsize=8.5, va="top")
                    yy -= 0.02
                pdf.savefig(fig)
                plt.close(fig)
        shutil.copyfile(tmp_pdf, output)
    determinism.set_pdf_lang(output, docxaudit.pdf_lang(journalspec.load(project)))
    determinism.normalize(output)
    sources = [project / "05_figures/legends.md", *[img for _e, img, _l in images]]
    rec = _receipt(project, output, "esm_figures", sources,
                   {"number": number, "figures": [str(e.get("id")) for e, _i, _l in images]})
    record_derived(project, [output, rec])
    return output, rec


def xlsx_info(project: Path, src: Path, output: Path, number: int | None = None) -> tuple[Path, Path]:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from .revision import guard_outputs, record_derived
    from .packagecontent import build_receipt_path
    project, src, output = Path(project), Path(src), Path(output)
    if not src.is_file():
        raise ValueError(f"workbook missing: {src}")
    guard_outputs(project, [output, build_receipt_path(project, output)])
    wb = load_workbook(src)
    if "Info" in wb.sheetnames:
        del wb["Info"]
    ws = wb.create_sheet("Info", 0)
    lines = info_lines(project, number)
    for r, line in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=line)
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    row = len(lines) + 2
    ws.cell(row=row, column=1, value="Contents").font = Font(bold=True)
    for name in wb.sheetnames[1:]:
        row += 1
        ws.cell(row=row, column=1, value=name)
    ws.column_dimensions["A"].width = 100
    wb.active = 0
    wb.properties.created = FIXED_TIME.replace(tzinfo=None)
    wb.properties.modified = FIXED_TIME.replace(tzinfo=None)
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    try:
        from . import determinism
        determinism.normalize(output)
    except Exception:  # noqa: BLE001 - normalisation is best effort for workbooks
        pass
    rec = _receipt(project, output, "esm_xlsx", [src],
                   {"number": number, "input": src.relative_to(project).as_posix()})
    record_derived(project, [output, rec])
    return output, rec


def rebuild_from(project: Path, receipt: Path) -> Path:
    rec = json.loads(Path(receipt).read_text(encoding="utf-8"))
    args = rec.get("args") or {}
    out = Path(project) / rec["output"]
    if rec.get("kind") == "esm_figures":
        return figures_pdf(project, out, args.get("number"), args.get("figures"))[0]
    if rec.get("kind") == "esm_xlsx":
        return xlsx_info(project, Path(project) / args["input"], out, args.get("number"))[0]
    raise ValueError(f"not a supplementary-builder receipt: {receipt}")

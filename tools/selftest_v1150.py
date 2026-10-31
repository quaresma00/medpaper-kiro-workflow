"""1.15.0 regression tests from a full S01-S25 run on 1.14.0 whose S23-S25 packaging, review and
audit loop took more than half of the working time: incremental meaning checks (E1), review
reports written by the reviewer and checked (E2), one-command package rebuild (E3), S25
convergence (E4), supplementary builders (E5), separate display files and the abbreviation list
(E6), a fully generated submission_qc.md and dated guideline snapshots (E7), final-copy coverage
of PDFs and figure files (E8), soft page-derived word caps (E9), the AI/LLM-use fact (E10),
checklist numbering (E11) and reference-list digits and overrides (E12). Synthetic fixtures
only - no real paper or dataset is read.

Called from tools/selftest.py.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from selftest_v1140 import _env, _envset, _out, _paper, _run, _w  # noqa: E402


def run(tmp: Path, record, section, only: str = "") -> None:
    tmp.mkdir(parents=True, exist_ok=True)
    for name, fn in GROUPS:
        if not only or name in only.split():
            fn(tmp / name, record, section)


def _set_round(p: Path, rid: str | None) -> None:
    st = json.loads((p / ".wf/state.json").read_text(encoding="utf-8"))
    st["active_revision_round"] = rid
    (p / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")


# ------------------------------------------------------------------------------ E1
FROZEN = ("# T\n\n# Results\n\nStatin use was associated with lower mortality (hazard ratio 0.80). "
          "Causal inference is not possible from these data. The cohort included older adults "
          "treated in two hospitals.\n\nThe secondary outcome did not differ between groups.\n")


def run_e1(tmp: Path, record, section) -> None:
    from wfcore import meaning, scifreeze
    import selftest_v1140 as v1140
    section("1.15.0 E1: meaning checks judge sentences, increments and one bundle per round")
    terms: dict = {}
    old = ("Statin use was associated with lower mortality. Causal inference is not possible "
           "from these data. The cohort included older adults.")
    trimmed = "Statin use was associated with lower mortality. The cohort included older adults."
    pairs = meaning.pairs_for(old, trimmed, "x.md", terms)
    edited = "Statin use was associated with lower mortality. The cohort comprised older adults."
    pairs2 = meaning.pairs_for(old, edited, "x.md", terms)
    record("E1b a paragraph is judged per aligned sentence: only the sentence whose meaning cues "
           "changed is a pair; a later one-word edit elsewhere in the paragraph yields the SAME "
           "normalised pair hash (no re-judging)",
           len(pairs) == 1 and "Causal inference" in pairs[0]["old"] and pairs[0]["new"] == ""
           and meaning.pair_hash(pairs[0]) == meaning.pair_hash(pairs2[0]) and len(pairs2) == 1
           and meaning.pair_hash({"old": "No  effect, was seen [@a].", "new": ""})
           == meaning.pair_hash({"old": "no effect was seen", "new": "", "file": "other.md"}),
           str(pairs2)[:300])

    p = v1140._frozen(tmp / "variant", FROZEN)
    var = _w(p, "08_submission/manuscript/full_manuscript.md", FROZEN.replace(
        " Causal inference is not possible from these data.", ""))
    first = scifreeze.variant_meaning(p)
    reg = meaning.known_pairs(p)
    rec = meaning.record_author(p, reg[0], "unchanged", "The caveat is kept in the Discussion.")
    second = scifreeze.variant_meaning(p)
    var.write_text(var.read_text(encoding="utf-8").replace("two hospitals", "two university "
                                                           "hospitals"), encoding="utf-8")
    third = scifreeze.variant_meaning(p)
    var.write_text(var.read_text(encoding="utf-8").replace("did not differ", "differed"),
                   encoding="utf-8")
    fourth = scifreeze.variant_meaning(p)
    reg2 = meaning.known_pairs(p)
    record("E1a the journal copy is compared with the frozen text once; after every pair is "
           "resolved, later checks compare the last confirmed copy with the current one: a "
           "one-word edit re-opens nothing, a new meaning change is the only new pair (the "
           "accepted trim is not asked again)",
           first and not second and not third and fourth and len(reg) == 1 and len(reg2) == 2
           and "differed" in reg2[-1]["new"] and "Causal" not in fourth[0]
           and (p / meaning.CONFIRMED_REL).is_file() and rec["by"] == "author",
           str(fourth)[:300])

    q = _paper(tmp / "round")
    _set_round(q, "R007")
    a = meaning.pairs_for("No association was seen.", "An association was seen.", "a.md", terms)
    b = meaning.pairs_for("Effects may be causal.", "Effects are causal.", "b.md", terms)
    meaning.gate(q, "R007", a)
    meaning.gate(q, "variant", a + b)
    meaning.gate(q, "final", b)
    bundles = sorted(x.name for x in (q / meaning.BUNDLE_DIR).glob("meaning_*"))
    inside = json.loads((q / meaning.BUNDLE_DIR / "meaning_R007/pairs.json").read_text(encoding="utf-8"))
    record("E1c one meaning bundle per rework round for all contexts (round edits, submission "
           "copy, final copy), deduplicated, with stable ids",
           bundles == ["meaning_R007"] and len(inside) == 2
           and [x["id"] for x in inside] == ["P01", "P02"], f"{bundles} {len(inside)}")

    _w(q, ".wf/review_bundle/meaning_R007/report.md",
       "P01: meaning changed - association reversed\nP02: meaning unchanged - same claim\n")
    unresolved, changed = meaning.resolve(q, "R007", list(a))
    h = changed[0]["hash"] if changed else ""
    line = meaning.changed_line(changed[0]) if changed else ""
    env = _env(q)
    dec = _run(q, "tools/rework.py", "meaning-decide", "--pair", h[:8], "--verdict", "unchanged",
               "--words", "作者确认：原文措辞有误，修改后意思一致", env=env)
    _u2, changed2 = meaning.resolve(q, "R007", list(a))
    over = meaning.overrides(q)
    record("E1d a reviewer 'changed' verdict names its id and hash and can be overridden by the "
           "author by hash (exact words recorded, shown as an override); without an override it "
           "stays 'changed' (SCIENCE path with explicit rewind confirmation)",
           not unresolved and len(changed) == 1 and h[:8] in line and "meaning-decide" in line
           and dec.returncode == 0 and "overrides reviewer 'changed'" in _out(dec)
           and not changed2 and over and "作者确认" in over[0] and "reviewer 'changed'" in over[0],
           _out(dec)[-300:] + str(over)[:200])
    src = (ROOT / "tools/rework.py").read_text(encoding="utf-8")
    record("E1d rework check routes an unoverridden 'changed' pair to SCIENCE and prints author "
           "overrides in the round summary and at close",
           "meaning.changed_line(p)" in src and "author override:" in src
           and "author meaning override" in src, "")


# ------------------------------------------------------------------------------ E2
def run_e2(tmp: Path, record, section) -> None:
    from wfcore import adapters, meaning, reviewguard
    section("1.15.0 E2: the reviewer writes its report; nothing else may change during a review")
    agent = (ROOT / ".kiro/agents/medpaper-reviewer.md").read_text(encoding="utf-8")
    head = agent.split("\n---", 2)[0]
    record("E2 medpaper-reviewer may write its report.md; the deny list no longer holds patterns "
           "that also match .wf/review_bundle/<round>/report.md ('*.md', '.wf/*'); scope check "
           "passes; fallback markers documented",
           'tools: ["read", "write"]' in head and '"*.md"' not in head and '".wf/*"' not in head
           and adapters.agent_scope_problems("medpaper-reviewer", agent) == []
           and "<<<REPORT>>>" in agent, str(adapters.agent_scope_problems("medpaper-reviewer", agent)))

    p = _paper(tmp / "guard", "S24_package_human_review")
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\ntext\n")
    _w(p, "08_submission/cover_letter.md", "Dear editor\n")
    _w(p, ".wf/review_rounds.json", {"manuscript": [], "package": [{
        "id": "P001", "scope": "package", "stage": "S24_package_human_review", "mode": "full",
        "started_at": "2026-01-01T00:00:00Z", "inputs": {}, "reviewed": [], "closed_at": None,
        "findings": {}, "changes": [], "not_changed": [], "after": {}, "user_decision": None,
        "protocol": "1.11.0"}]})
    reviewguard.take(p, "P001")
    time.sleep(1.2)
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\nchanged by the reviewer\n")
    msg = _w(tmp, "reply.txt", "Here is my report.\n<<<REPORT>>>\n## Verdict\nF1 [other] Typo in "
             "Methods.\nF2 [numbers] Count differs.\n<<<END>>>\nThanks.")
    time.sleep(1.2)
    ing = _run(p, "tools/review_round.py", "ingest", "--scope", "package", "--from", msg)
    report = p / ".wf/review_bundle/P001/report.md"
    bad = reviewguard.problems(p, "P001", report)
    later = _w(p, "08_submission/cover_letter.md", "Dear editor, fixed after the report\n")
    t = time.time() + 30
    os.utime(later, (t, t))
    close = _run(p, "tools/review_round.py", "close", "--scope", "package", "--findings",
                 "critical=0,major=1,minor=1", "--finding", "F1", "other", "changed", "typo",
                 "--finding", "F2", "numbers", "author", "count")
    record("E2 review_round ingest --from copies the text between <<<REPORT>>> markers verbatim; "
           "close refuses a report written while a paper file changed (the reviewer edited the "
           "manuscript); edits made after the report are not 'during the review'",
           ing.returncode == 0 and report.read_text(encoding="utf-8").startswith("## Verdict\nF1 [other]")
           and "Thanks" not in report.read_text(encoding="utf-8")
           and bad == ["07_manuscript/full_manuscript.md"] and close.returncode == 2
           and "changed while the review was running" in _out(close)
           and reviewguard.problems(p, "P001", report, ["07_manuscript/full_manuscript.md"]) == [],
           _out(ing)[-200:] + str(bad) + _out(close)[-300:])

    state = {"executions": [{"steps": [{"messages": [{"role": "assistant", "content":
             "x <<<REPORT>>>\nP01: meaning unchanged - same\n<<<END>>> y"}]}]}]}
    ws = _w(tmp, "workflow-state.json", state)
    q = _paper(tmp / "mean")
    meaning.gate(q, "variant", meaning.pairs_for("No effect.", "An effect.", "a.md", {}))
    ing2 = _run(q, "tools/rework.py", "meaning-ingest", "--context", "variant", "--from", ws)
    un, ch = meaning.resolve(q, "variant", meaning.pairs_for("No effect.", "An effect.", "a.md", {}))
    record("E2 rework.py meaning-ingest --from accepts a host workflow-state JSON and stores the "
           "verdicts verbatim (no paraphrase); the check then has no open pair",
           ing2.returncode == 0 and not un and not ch
           and reviewguard.extract_report(json.dumps(state)) == "P01: meaning unchanged - same\n",
           _out(ing2)[-300:])


# ------------------------------------------------------------------------------ helpers
def _png(path: Path, color=(30, 90, 160), size=(240, 160), dpi=None, fmt=None) -> Path:
    from PIL import Image
    path.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", size, "white")
    for x in range(20, size[0] - 20):
        for y in range(40, 60):
            im.putpixel((x, y), color)
    kw = {"dpi": (dpi, dpi)} if dpi else {}
    im.save(path, format=fmt, **kw)
    return path


def _supp_paper(base: Path, target: dict | None = None) -> Path:
    p = _paper(base, "S23_package")
    _w(p, "01_protocol/artifact_plan.json", {
        "main_figures": [{"id": "Figure 1", "file": "05_figures/out/Figure1.png"}],
        "supp_figures": [{"id": "Figure S1", "file": "05_figures/out/FigureS1.png"},
                         {"id": "Figure S2", "file": "05_figures/out/FigureS2.png"}],
        "main_tables": [{"id": "Table 1", "file": "04_tables/main/Table1.xlsx"}],
        "supp_tables": []})
    _png(p / "05_figures/out/Figure1.png")
    _png(p / "05_figures/out/Figure1.tif", dpi=600, fmt="TIFF")
    _png(p / "05_figures/out/FigureS1.png", (200, 40, 40))
    _png(p / "05_figures/out/FigureS2.png", (40, 160, 40))
    _w(p, "05_figures/legends.md", "# Figure legends\n\n## Figure 1. Main.\nMain legend.\n\n"
       "## Figure S1. Sensitivity analysis.\nOdds ratios by year; ROR, reporting odds ratio.\n\n"
       "## Figure S2. Negative control.\nNo signal for the control drug.\n")
    _w(p, "07_manuscript/title.md", "# Drug A and portal vein injury\n")
    _w(p, "00_input/author_info.json", {"authors": [{"name": "Li Wei"}, {"name": "Ann Roe"}],
                                         "corresponding": {"name": "Li Wei", "email": "li@x.org",
                                                           "address": "Dept of Hepatology"}})
    _w(p, "08_submission/target_journal.json", target or {"journal": "Eur J Test"})
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Table 1"
    wb.active["A1"] = "Table 1. Baseline"
    wb.active["A2"] = "Age"
    wb.active["B2"] = 61
    (p / "04_tables/main").mkdir(parents=True, exist_ok=True)
    wb.save(p / "04_tables/main/Table1.xlsx")
    return p


# ------------------------------------------------------------------------------ E3
def run_e3(tmp: Path, record, section) -> None:
    from wfcore import rebuild, supplementary
    section("1.15.0 E3: one command rebuilds the affected package outputs")
    p = _supp_paper(tmp / "rb")
    with _envset(MEDPAPER_PROJECT=p):
        out, rec = supplementary.figures_pdf(p, p / "08_submission/bundle/ESM_3.pdf", 3)
    before = out.read_bytes()
    todo0, _n = rebuild.plan(p)
    leg = p / "05_figures/legends.md"
    leg.write_text(leg.read_text(encoding="utf-8").replace("No signal", "No disproportionality "
                                                           "signal"), encoding="utf-8")
    todo1, _n = rebuild.plan(p)
    dry = _run(p, "tools/wf.py", "rebuild", "--dry-run")
    done = _run(p, "tools/wf.py", "rebuild", "--wait")
    receipt = json.loads(rec.read_text(encoding="utf-8"))
    record("E3 wf rebuild finds the outputs whose build inputs changed (here a legend -> the "
           "supplementary figures PDF), rebuilds only those in dependency order, waits with "
           "heartbeats and prints failures + one summary line (details in .wf/logs/rebuild.log)",
           not todo0 and [r["output"] for r in todo1] == ["08_submission/bundle/ESM_3.pdf"]
           and "would rebuild 08_submission/bundle/ESM_3.pdf" in _out(dry)
           and done.returncode == 0 and "[rebuild] 1 output(s) rebuilt" in _out(done)
           and "[wf rebuild] rc=0" in _out(done) and out.read_bytes() != before
           and (p / ".wf/logs/rebuild.log").is_file() and receipt["kind"] == "esm_figures"
           and not rebuild.plan(p)[0], _out(dry)[-300:] + _out(done)[-500:])
    fake = {"kind": "manuscript", "output": "08_submission/bundle/manuscript.docx",
            "sources": [{"path": "08_submission/manuscript/full_manuscript.md"}],
            "auxiliaries": [{"path": "06_refs/refs.bib"}, {"path": "C:/x/journal.csl"},
                            {"path": "08_submission/docx_style.json"},
                            {"path": "08_submission/target_journal.json"}],
            "pdf_output": None, "_receipt": p / "r.json"}
    cmd = rebuild.build_command(p, fake)
    cards = {s: (ROOT / f"pipeline/stages/{s}.md").read_text(encoding="utf-8")
             for s in ("S23_package", "S24_package_human_review", "S25_submission_audit")}
    record("E3 an older receipt (no stored arguments) still rebuilds with the same bibliography, "
           "CSL and style; the S23/S24 cards use wf rebuild",
           "--bibliography" in cmd and str(p / "06_refs/refs.bib") in cmd and "--csl" in cmd
           and str(p / "08_submission/docx_style.json") in cmd
           and "wf.py rebuild" in cards["S23_package"] and "wf.py rebuild" in cards["S24_package_human_review"],
           " ".join(cmd)[-300:])


# ------------------------------------------------------------------------------ E4
def run_e4(tmp: Path, record, section) -> None:
    from wfcore import audit, reviewrounds as rr
    section("1.15.0 E4: S25 converges - full first audit, then re-audits of open findings")
    p = _paper(tmp / "audit", "S25_submission_audit")
    _w(p, "08_submission/package_review_freeze.json", {"freeze_id": "a" * 64})
    _w(p, "08_submission/bundle/manuscript.docx", "x")
    b1, m1 = audit.start(p, ["08_submission/bundle/manuscript.docx"])
    man1 = (b1 / "MANIFEST.md").read_text(encoding="utf-8")
    _w(p, f".wf/review_bundle/{b1.name}/report.md",
       "## Verdict\nREVISE\nFreeze ID: " + "a" * 64 + "\n## Issues requiring correction\n"
       "- F1 MINOR manuscript.docx p3: typo 'teh'\n- F2 BLOCKING title page: ORCID missing\n"
       "- F3 PREFERENCE manuscript.docx: shorter title\n")
    saved = audit.save(p)
    offer = audit.revise_offer(p)
    _w(p, "08_submission/package_review_freeze.json", {"freeze_id": "b" * 64})
    b2, m2 = audit.start(p, ["08_submission/bundle/manuscript.docx"])
    man2 = (b2 / "MANIFEST.md").read_text(encoding="utf-8")
    prev = (b2 / "previous_findings.md").read_text(encoding="utf-8")
    record("E4 the first S25 audit is full; after REVISE the next one is a re-audit of the "
           "previous findings plus new BLOCKING issues; engine-owned formatting is an 'Engine "
           "issue', never REVISE by itself; at the first REVISE both routes (S24 round or author "
           "final copy) are offered with the open items",
           m1 == "full" and m2 == "recheck" and "FULL" in man1 and "RE-AUDIT" in man2
           and "Engine issues" in man1 and "PDF /Lang" in man1 and "F1 MINOR" in prev
           and "F2 BLOCKING" in prev and len(audit.history(p)) == 1 and saved.is_file()
           and any("author final copy" in x for x in offer)
           and sum(1 for x in offer if x.startswith("  - F")) == 2,
           man2[:400] + str(offer))
    rid = "P001"
    _w(p, ".wf/review_rounds.json", {"manuscript": [], "package": [{
        "id": rid, "scope": "package", "stage": "S24_package_human_review", "mode": "full",
        "started_at": "2026-01-01T00:00:00Z", "inputs": {}, "reviewed": [], "closed_at": None,
        "findings": {}, "changes": [], "not_changed": [], "after": {}, "user_decision": None,
        "protocol": "1.11.0"}]})
    rep = _w(tmp, "rep.md", "F1 [other] MINOR typo in Methods\nF2 [other] PREFERENCE shorter "
             "title\nF3 [other] PREFERENCE fewer subheadings\n")
    real_inputs = rr.inputs
    rr.inputs = lambda project, scope: {}  # the package evidence set is not under test here
    try:
        rnd = rr.close(p, "package", {"minor": 1}, [{"file": "x.md", "summary": "typo fixed"}], [],
                       [{"id": "F1", "category": "other", "status": "changed", "summary": "typo"}],
                       report=rep, preference_reason="stylistic preferences; the journal has no rule")
    finally:
        rr.inputs = real_inputs
    record("E4 preference findings close as ONE group with one reason (not counted, never a "
           "protected category); review_round close --preference-group exposes it",
           rnd["preference_group"]["ids"] == ["F2", "F3"] and not rr._closure(rnd)
           and "--preference-group" in (ROOT / "tools/review_round.py").read_text(encoding="utf-8"),
           str(rnd.get("preference_group")))


# ------------------------------------------------------------------------------ E5
def run_e5(tmp: Path, record, section) -> None:
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from wfcore import supplementary
    from manuscript import build_docx as bd
    section("1.15.0 E5: supplementary figures PDF, workbook info sheet, one ESM header")
    p = _supp_paper(tmp / "esm", {"journal": "Eur J Test",
                                   "supplementary_label": "Electronic Supplementary Material"})
    with _envset(MEDPAPER_PROJECT=p):
        a, _r = supplementary.figures_pdf(p, p / "08_submission/bundle/ESM_3.pdf", 3)
        sha1 = a.read_bytes()
        supplementary.figures_pdf(p, p / "08_submission/bundle/ESM_3.pdf", 3)
    text = "\n".join(pg.extract_text() or "" for pg in PdfReader(str(a)).pages)
    flat = " ".join(text.split())
    record("E5 the supplementary figures PDF: info page (journal label + n, title, authors, "
           "corresponding author with affiliation, journal), then each figure with its legend; "
           "the same inputs give the same bytes",
           a.read_bytes() == sha1 and len(PdfReader(str(a)).pages) == 3
           and "Electronic Supplementary Material 3" in flat and "Li Wei" in flat
           and "Dept of Hepatology" in flat and "Figure S2. Negative control" in flat
           and "Journal: Eur J" in flat, flat[:400])
    q = _supp_paper(tmp / "blind", {"journal": "J", "double_blind": True})
    with _envset(MEDPAPER_PROJECT=q):
        b, _r = supplementary.figures_pdf(q, q / "08_submission/bundle/ESM_3.pdf", 3,
                                          ["Figure S1"])
        x, _r = supplementary.xlsx_info(p, p / "04_tables/main/Table1.xlsx",
                                        p / "08_submission/bundle/ESM_2.xlsx", 2)
    btxt = " ".join("\n".join(pg.extract_text() or "" for pg in PdfReader(str(b)).pages).split())
    wb = load_workbook(x)
    info = [c.value for c in wb["Info"]["A"] if c.value]
    page = bd._info_page(p, 1)
    record("E5 blinded journals get no author names; the workbook gets an Info sheet first with "
           "the same header block (table sheets unchanged); Word supplementary files use the "
           "same 'Online Resource n' header (--esm-number)",
           "Li Wei" not in btxt and "Online Resource 3" in btxt and len(PdfReader(str(b)).pages) == 2
           and wb.sheetnames == ["Info", "Table 1"] and info[0] == "Electronic Supplementary Material 2"
           and wb["Table 1"]["B2"].value == 61 and page.startswith("# Electronic Supplementary Material 1"),
           f"{btxt[:200]} {wb.sheetnames} {info[:3]} {page[:60]}")
    _w(p, "temp/logs/x.log", "log\n")
    _w(p, "temp/helper.py", "print(1)\n")
    _w(p, ".wf/scripts/esm_helper.py", "print(1)\n")
    cl = _run(p, "tools/wf.py", "clean", "--apply")
    s23 = (ROOT / "pipeline/stages/S23_package.md").read_text(encoding="utf-8")
    record("E5 wf clean keeps temp/logs/ and .wf/scripts/ (durable helpers), deletes other temp "
           "files; the S23 card says never keep needed scripts in temp/",
           (p / "temp/logs/x.log").is_file() and not (p / "temp/helper.py").exists()
           and (p / ".wf/scripts/esm_helper.py").is_file() and ".wf/scripts/" in s23
           and "supplementary.py" in s23, _out(cl)[-300:])


# ------------------------------------------------------------------------------ E6
def run_e6(tmp: Path, record, section) -> None:
    from wfcore import displayfiles, journalspec
    from manuscript import build_docx as bd
    section("1.15.0 E6: separate figure/table files and the abbreviation list")
    p = _supp_paper(tmp / "disp", {"journal": "J", "figure_dpi": 300})
    _w(p, "08_submission/bundle/manifest.json", {"items": [
        {"file": "08_submission/bundle/manuscript.docx", "role": "manuscript", "required_by": "x"}]})
    with _envset(MEDPAPER_PROJECT=p):
        made = displayfiles.sync(p)
        again = displayfiles.sync(p)
    if not (p / "08_submission/bundle/Fig1.tif").is_file():
        record("E6 separate display files were written", False,
               f"{made} {displayfiles.wanted(p)}")
        return
    man = json.loads((p / "08_submission/bundle/manifest.json").read_text(encoding="utf-8"))
    gen = [i["file"] for i in man["items"] if i.get("generated_by") == "separate_display_files"]
    record("E6 separate_display_files (default ON): Fig1.tif and Table1.xlsx byte-identical to "
           "the certified masters, manifest/upload list updated automatically, idempotent",
           (p / "08_submission/bundle/Fig1.tif").read_bytes()
           == (p / "05_figures/out/Figure1.tif").read_bytes()
           and (p / "08_submission/bundle/Table1.xlsx").read_bytes()
           == (p / "04_tables/main/Table1.xlsx").read_bytes()
           and sorted(gen) == ["08_submission/bundle/Fig1.tif", "08_submission/bundle/Table1.xlsx"]
           and len(made) == 3 and again == [] and displayfiles.problems(p) == [], str(made))
    old = _supp_paper(tmp / "old")
    st = json.loads((old / ".wf/state.json").read_text(encoding="utf-8"))
    st["engine_pin"] = "1.14.1"
    (old / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    record("E6 the switch is off when the journal forbids separate files, and for a paper "
           "pinned before 1.15.0 until wf migrate (package contents change); an explicit true "
           "always applies",
           not journalspec.separate_display_files({"separate_display_files": "forbidden"}, p)
           and not journalspec.separate_display_files({}, old)
           and journalspec.separate_display_files({"separate_display_files": True}, old)
           and journalspec.abbreviation_list({}, old) == "keep"
           and journalspec.abbreviation_list({}, p) == "after_keywords", "")
    text = ("# Title\n\n# Abstract\n\nText.\n\nKeywords: a, b, c\n\n# Introduction\n\nBody.\n\n"
            "# Declarations\n\n## Funding\n\nNone.\n\n## Abbreviations\n\nAE, adverse event; "
            "ROR, reporting odds ratio\n\n## Competing interests\n\nNone.\n")
    moved = bd.place_abbreviations(text, "after_keywords")
    decl = moved[moved.index("# Declarations"):]
    gone = bd.place_abbreviations(text, "none")
    record("E6 abbreviation_list after_keywords: the list becomes its own section directly after "
           "the Keywords (never inside Declarations); 'none' removes it; displays keep their own "
           "definitions",
           moved.index("# Abbreviations") > moved.index("Keywords:")
           and moved.index("# Abbreviations") < moved.index("# Introduction")
           and "adverse event" not in decl and "## Competing interests" in decl
           and "adverse event" not in gone and bd.place_abbreviations(text, "keep") == text,
           moved[:300])


# ------------------------------------------------------------------------------ E7
def run_e7(tmp: Path, record, section) -> None:
    import tomllib
    from wfcore import guidelines, registry, submissionqc
    from wfcore.state import State
    section("1.15.0 E7: generated submission_qc.md and a dated guideline snapshot")
    pipe = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
    s23 = next(s for s in pipe["stage"] if s["id"] == "S23_package")
    p = _paper(tmp / "qc", "S23_package")
    ctx = SimpleNamespace(project=p, state=State(p, ".wf").load())
    full = submissionqc.render_full(ctx, [{"item": "x", "status": "PASS", "evidence": "ok"}])
    record("E7 submission_qc.md is fully generated (summary, files measured, measured block); "
           "the hand-written headings gate is gone (relaxation: every paper)",
           not any(g["check"] == "md_sections" and "submission_qc" in str(g.get("path", ""))
                   for g in s23["gate"]) and full.startswith("# Submission QC")
           and submissionqc.BEGIN in full and "Rendered files inspected" not in full, full[:300])
    _w(p, "08_submission/guidelines_extract.md", "# Old extract\n\nWord limit 3000.\n")
    new = _w(tmp, "new_extract.md", "# Guideline extract\n\n## Manuscript\n" + "Rule text. " * 30
             + "\n\n## Use of AI\nAuthors must disclose the use of large language models in "
             "the Methods section.\n")
    r = _run(p, "tools/wf.py", "guidelines", "refresh", "--from", new, "--source",
             "https://journal.example/guide")
    text = (p / guidelines.REL).read_text(encoding="utf-8")
    hist = sorted((p / guidelines.HISTORY).glob("*.md"))
    record("E7 wf guidelines refresh installs the new extract with a dated snapshot line, keeps "
           "the old one in guidelines_history/, and status reports the date (audits read it)",
           r.returncode == 0 and guidelines.snapshot_date(text) is not None
           and "journal.example" in text and len(hist) == 1
           and "Word limit 3000" in hist[0].read_text(encoding="utf-8")
           and "dated" in guidelines.status(p), _out(r)[-300:])


# ------------------------------------------------------------------------------ E8
def _pdf_with_text(path: Path, lines: list[str]) -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    matplotlib.rcParams["pdf.fonttype"] = 42
    path.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(path) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        y = 0.9
        for line in lines:
            fig.text(0.1, y, line, fontsize=10)
            y -= 0.03
        fig.text(0.5, 0.03, "1", fontsize=9)
        pdf.savefig(fig)
        plt.close(fig)
    return path


def run_e8(tmp: Path, record, section) -> None:
    from docx import Document
    from wfcore import finalcopy
    section("1.15.0 E8: the author final copy also covers PDFs and figure files")
    p = _supp_paper(tmp / "fc")
    _w(p, "03_analysis/results/primary.json", {"hr": 1.34, "lo": 1.1, "hi": 1.6})
    lines = ["Supplementary results", "The hazard ratio was 1.34 (95% CI 1.10 to 1.60)."]
    d = Document()
    for ln in lines:
        d.add_paragraph(ln)
    (p / "08_submission/final_copy").mkdir(parents=True, exist_ok=True)
    d.save(p / "08_submission/final_copy/ESM_1.docx")
    good = _pdf_with_text(p / "08_submission/final_copy/ESM_1.pdf", lines)
    ok = finalcopy.compare_pdf(p, good, {})
    _pdf_with_text(good, [lines[0], "The hazard ratio was 1.43 (95% CI 1.10 to 1.60)."])
    bad = finalcopy.compare_pdf(p, good, {})
    fig_ok = p / "08_submission/final_copy/Fig1.tif"
    fig_ok.write_bytes((p / "05_figures/out/Figure1.tif").read_bytes())
    fig_bad = _png(p / "08_submission/final_copy/Fig2.tif", (0, 0, 0), fmt="TIFF")
    old = _supp_paper(tmp / "fc_old")
    st = json.loads((old / ".wf/state.json").read_text(encoding="utf-8"))
    st["engine_pin"] = "1.14.1"
    (old / ".wf/state.json").write_text(json.dumps(st), encoding="utf-8")
    _w(old, "08_submission/final_copy/ESM_1.pdf", "%PDF")
    record("E8 a final-copy PDF is compared with its checked Word copy (numbers, citations, "
           "meaning; page numbers ignored); a TIFF/PNG must be byte-identical to a certified "
           "master; papers pinned before 1.15.0 keep the docx/xlsx-only check until migrate",
           ok == [] and bad and "1.34" in " ".join(bad) and finalcopy.compare_image(p, fig_ok) == []
           and finalcopy.compare_image(p, fig_bad) and finalcopy.copies(old) == []
           and any(c.suffix == ".pdf" for c in finalcopy.copies(p))
           and "export-pdf" in (ROOT / "tools/final_copy.py").read_text(encoding="utf-8"),
           f"{ok} {bad}")


# ------------------------------------------------------------------------------ E9-E12
def run_e9(tmp: Path, record, section) -> None:
    from wfcore import journalspec
    import selftest_v180 as v180
    section("1.15.0 E9-E12: soft page-derived caps, AI fact, checklist order, references")
    est = {"limits": {"main_text_words": 3500, "abstract_words": 250,
                      "source": "8-10 printed pages; approx. 3500 words estimated"}}
    exp = {"limits": {"main_text_words": 3500, "source": "Main text max 3500 words"}}
    p = _paper(tmp / "soft", "S22_polish")
    _w(p, "08_submission/target_journal.json", {"journal": "J", **est})
    _w(p, "07_manuscript/full_manuscript.md", "# T\n\n# Abstract\n\nShort.\n\n# Introduction\n\n"
       + " ".join(["word"] * 3700) + "\n\n# References\n")
    g = v180._gate(p, "journal_limits_met", "S22_polish")
    card = (ROOT / "pipeline/stages/S20_journal.md").read_text(encoding="utf-8")
    record("E9 only explicit numeric limits are hard: a word cap estimated from a page count is "
           "a soft hint (never a reason to cut); limits.estimated marks one explicitly; S20 card "
           "says never derive word caps from page counts",
           journalspec.soft_limits(est) == {"main_text_words"} and not journalspec.soft_limits(exp)
           and journalspec.soft_limits({"limits": {"main_text_words": 1, "estimated":
                                                   ["main_text_words"], "source": "x"}})
           == {"main_text_words"} and g.ok and "soft hint" in g.detail
           and "page count" in card, g.detail[:300])

    from wfcore import audit, declarations
    q = _paper(tmp / "ai", "S21_authors")
    _w(q, declarations.REL, {"questions": [], "answers": {}})
    _w(q, "08_submission/guidelines_extract.md", "# Extract\n\n## Use of AI\nThe use of large "
       "language models (LLMs) must be documented in the Methods section.\n\n## Other\nx\n")
    rule = _run(q, "tools/declarations.py", "ai-rule")
    ai = _run(q, "tools/declarations.py", "ai", "--used", "no", "--words", "没有使用任何AI工具")
    fact = declarations.ai_fact(q)
    _w(q, "08_submission/package_review_freeze.json", {"freeze_id": "c" * 64})
    b, _m = audit.start(q, [])
    record("E10 S21 shows the journal's AI/LLM rule from the guideline snapshot, records the "
           "author's factual answer, generates the statement from it, and the audit bundle "
           "presents it as a recorded fact (not pending)",
           rule.returncode == 0 and "large language models" in _out(rule)
           and ai.returncode == 0 and fact and fact["used"] is False
           and fact["author_words"] == "没有使用任何AI工具" and "No generative" in fact["statement"]
           and "recorded fact" in (b / "MANIFEST.md").read_text(encoding="utf-8")
           and "declarations.py ai" in (ROOT / "pipeline/stages/S21_authors.md").read_text(encoding="utf-8"),
           _out(ai)[-300:])

    agent = (ROOT / ".kiro/agents/medpaper-reviewer.md").read_text(encoding="utf-8")
    record("E11 reporting checklists/forms are numbered after the content supplementary files "
           "and exempt from first-citation ordering (reviewer prompt, package and audit briefs)",
           "exempt from first-citation ordering" in agent
           and "exempt from first-citation ordering" in audit.engine_text(), "")

    from wfcore import finalcopy, nonfacts, refoverrides
    bib = ("@article{smith2010,\n  title = {A trial},\n  pages = {731-40; discussion 741-2},\n"
           "  year = {2010}\n}\n@article{lee2012,\n  title = {B},\n  pages = {12-9},\n  year = {2012}\n}\n")
    r = _paper(tmp / "refs", "S23_package")
    _w(r, "06_refs/refs.bib", bib)
    patched, changes = refoverrides.patch(r, bib, True)
    _w(r, refoverrides.REL, {"lee2012": [{"field": "pages", "value": "12-19",
                                          "reason": "full page range for the journal style"}]})
    good = refoverrides.problems(r)
    patched2, _c = refoverrides.patch(r, bib, True)
    _w(r, refoverrides.REL, {"lee2012": [{"field": "pages", "value": "12-29", "reason": "invented range"}],
                             "nokey": [{"field": "pages", "value": "1", "reason": "missing key x"}]})
    bad = refoverrides.problems(r)
    record("E12 a '; discussion ...' tail is dropped from page fields at render (refs.bib "
           "untouched); 06_refs/ref_overrides.json is validated against the PubMed field (may "
           "only shorten it or write the same range in full) and applied at render",
           "pages = {731-40}" in patched and "discussion" not in patched and len(changes) == 1
           and good == [] and "pages = {12-19}" in patched2 and len(bad) == 2
           and "06_refs/ref_overrides.json" in (ROOT / "tools/manuscript/build_docx.py").read_text(encoding="utf-8")
           and "refoverrides.render_copy" in (ROOT / "tools/manuscript/build_docx.py").read_text(encoding="utf-8"),
           str(bad))
    doc = ("Results\nThe hazard ratio was 1.34.\nReferences\n1. Smith A. Trial. Ann Surg. "
           "2010;251:731-40; discussion 741-2.\nFigure legends\nFigure 1. Flow.\n")
    edited = doc.replace("731-40; discussion 741-2", "731-40")
    _w(r, "03_analysis/results/primary.json", {"hr": 1.34})
    probs = finalcopy.compare_text(r, doc, edited, r / "08_submission/final_copy/m.docx")
    record("E12 digits inside the reference list are not compared as facts (a page-range fix "
           "in the final copy is not 'a reported value changed'); body numbers still are",
           probs == [] and "731" not in nonfacts.strip_reference_list(doc)
           and "1.34" in nonfacts.strip_reference_list(doc)
           and "Figure 1. Flow" in nonfacts.strip_reference_list(doc), str(probs))


# ------------------------------------------------------------------------------ R
def run_r(tmp: Path, record, section) -> None:
    from wfcore import registry
    from wfcore.state import State
    section("1.15.0 release: version, one migration event, pinned paper policy, card length")
    q = tmp / "old"
    q.mkdir(parents=True, exist_ok=True)
    st = State(q, ".wf").create("medpaper", "1.14.1", "S25_submission_audit")
    st.migrate_pipeline("1.15.0", ["S25_submission_audit"])
    st.migrate_pipeline("1.15.0", ["S25_submission_audit"])
    notes = [e for e in st.data["events"] if e["kind"] == "requirements_changed"
             and str(e.get("detail", "")).startswith("1.15.0")]
    pipe = registry.load()
    sizes = {s.id: len((ROOT / "pipeline" / s.card).read_text(encoding="utf-8")) for s in pipe.stages}
    from wfcore import journalspec
    record("R pipeline 1.15.0; a 1.14.1 paper gets exactly one 1.15.0 event and stays pinned "
           "(package changes and stricter checks only after wf migrate --yes); cards <= limit",
           pipe.meta.get("version") == "1.15.0" and len(notes) == 1 and st.engine_pin == "1.14.1"
           and not journalspec.since_1150(q) and max(sizes.values()) <= registry.CARD_MAX_CHARS,
           f"{len(notes)} {max(sizes.values())} {sorted(sizes.items(), key=lambda kv: -kv[1])[:2]}")


GROUPS = [("E1", run_e1), ("E2", run_e2), ("E3", run_e3), ("E4", run_e4), ("E5", run_e5),
          ("E6", run_e6), ("E7", run_e7), ("E8", run_e8), ("E9", run_e9), ("R", run_r)]
